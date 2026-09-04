from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import math
import os
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.request
import zipfile
from collections import Counter
from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait
from decimal import ROUND_CEILING, Decimal, InvalidOperation
from pathlib import PurePath
from statistics import fmean
from typing import Any, Mapping

from docx import Document
from openpyxl import load_workbook  # type: ignore[import-untyped]
import pymupdf  # type: ignore[import-untyped]
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image
from ai_companion_worker.document_parser import ParseResult, count_tokens, parse_document
from ai_companion_worker.pdf_translation_layout import (
    PdfLayout,
    PdfTextBlock,
    extract_pdf_layout,
    render_layout_translation,
)
from ai_companion_worker.task_quality import (
    presentation_audience_content_violations,
    presentation_exhaustive_scope_violations,
    normalize_presentation_title,
    presentation_table_field_semantic_violations,
    presentation_table_language_violations,
    presentation_visible_language_violations,
)
from ai_companion_worker.translation_language import translation_language_policy

MAX_SOURCE_BYTES = 700 * 1024
MAX_PDF_SOURCE_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_SOURCE_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_CONTEXT_TOKENS = 4_000
MAX_DOCUMENT_CONTEXT_ROUNDS = 4
MAX_PDF_TRANSLATION_CHARS = 400_000
MAX_TRANSLATION_BATCH_CHARS = 7_000
MAX_TRANSLATION_BATCH_PAGES = 4
MAX_TRANSLATION_COST_MICROS = 60_000
MAX_MODEL_RESPONSE_BYTES = 4 << 20
MAX_ARCHIVE_BYTES = 25 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_PRESENTATION_SLIDES = 60
MAX_PRESENTATION_SOURCE_BYTES = 20 * 1024 * 1024
MAX_PRESENTATION_SOURCE_DOCUMENTS = 3
MAX_PRESENTATION_SOURCE_VISUALS = 12
MAX_PRESENTATION_VISUAL_BYTES = 2 * 1024 * 1024
MAX_PRESENTATION_IMAGE_RESPONSE_BYTES = 12 * 1024 * 1024
MAX_PRESENTATION_IMAGE_COST_MICROS = 250_000
PRESENTATION_DISPLAY_CELL_CHARS = 42
PRESENTATION_ROWS_PER_SLIDE = 6
PRESENTATION_TABLE_PAGE_CAPACITY = 12
PRESENTATION_DETAIL_CELL_THRESHOLD = 320
PRESENTATION_DETAIL_MAX_CHARS = 2_200
PRESENTATION_BACKGROUND = RGBColor(247, 249, 255)
PRESENTATION_TITLE_COLOR = RGBColor(43, 63, 117)
PRESENTATION_ACCENT_COLOR = RGBColor(72, 104, 183)
PRESENTATION_TEXT_COLOR = RGBColor(31, 41, 55)
PRESENTATION_MUTED_COLOR = RGBColor(91, 100, 116)
CJK_FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/System/Library/Fonts/STHeiti Light.ttc",
    "/System/Library/Fonts/STHeiti Medium.ttc",
    "/System/Library/Fonts/Supplemental/Songti.ttc",
)


class ModelBackedOperationError(ValueError):
    """A post-dispatch failure that must retain the provider usage record."""

    def __init__(self, code: str, model_usage: dict[str, Any]) -> None:
        super().__init__(code)
        self.code = code
        self.model_usage = model_usage


def execute(operation: str, payload: dict[str, Any]) -> dict[str, Any]:
    if operation == "docx_edit":
        return _edit_docx(payload)
    if operation == "pptx_generate":
        return _generate_pptx(payload, include_file=True)
    if operation == "pptx_outline":
        return _generate_pptx(payload, include_file=False)
    if operation == "document_extract":
        return _extract_document(payload)
    if operation == "tabular_profile":
        return _profile_tabular(payload)
    if operation == "pdf_translate":
        return _translate_pdf(payload)
    raise ValueError("unsupported_operation")


def _edit_docx(payload: dict[str, Any]) -> dict[str, Any]:
    source_name, source = _source_file(payload, {".docx"})
    _validate_ooxml(source, reject_macros=True)
    append_text = _required_text(payload, "append_text", 20_000)
    document = Document(io.BytesIO(source))
    paragraphs = [line.strip() for line in append_text.splitlines() if line.strip()]
    if not paragraphs:
        raise ValueError("append_text_is_empty")
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    output = io.BytesIO()
    document.save(output)
    data = output.getvalue()
    name = _output_name(payload.get("output_filename"), source_name, "-edited.docx")
    return {
        "output": {
            "source_filename": source_name,
            "output_filename": name,
            "change_summary": f"在文档副本末尾新增 {len(paragraphs)} 个段落",
            "appended_paragraphs": len(paragraphs),
            "source_overwritten": False,
        },
        "files": [
            _file(
                name,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                data,
            )
        ],
    }


def _generate_pptx(payload: dict[str, Any], *, include_file: bool) -> dict[str, Any]:
    title = normalize_presentation_title(_required_text(payload, "title", 200))
    audience = _required_text(payload, "audience", 200)
    style = _required_text(payload, "style", 100)
    brief = _required_text(payload, "brief", 10_000)
    requested_slide_count = payload.get("slide_count")
    if (
        not isinstance(requested_slide_count, int)
        or isinstance(requested_slide_count, bool)
        or not 3 <= requested_slide_count <= MAX_PRESENTATION_SLIDES
    ):
        raise ValueError("slide_count_must_be_between_3_and_60")
    table = _presentation_table(payload.get("table"))
    task_contract = payload.get("task_contract")
    task_contract = dict(task_contract) if isinstance(task_contract, dict) else {}
    source_coverage = _presentation_source_coverage(payload.get("source_coverage"))
    display_rows = _presentation_display_rows(table["rows"]) if table else []
    if table:
        table_pages = _presentation_table_pages(display_rows)
        slide_count = min(MAX_PRESENTATION_SLIDES, len(table_pages) + 2)
    else:
        table_pages = []
        slide_count = requested_slide_count
    section_specs = _presentation_section_specs(brief) if not table else []
    sections = _presentation_units(brief)
    groups: list[list[str]] = []
    headings: list[str] = []
    if not table:
        content_pages = slide_count - 2
        if section_specs:
            groups = [spec["bullets"] for spec in section_specs]
            headings = [spec["heading"] for spec in section_specs]
        else:
            sections = _ensure_presentation_units(sections, content_pages)
            groups = _balanced_groups(sections, content_pages)
            headings = [
                _presentation_heading(group[0], index)
                for index, group in enumerate(groups, start=1)
            ]
    visual_topics = (
        [
            {
                "title": table["title"] or title,
                "bullets": [
                    *table["columns"],
                    f"共 {len(table['rows'])} 条记录",
                    *list(
                        dict.fromkeys(str(row.get("source_locator") or "") for row in table["rows"])
                    )[:12],
                ],
            }
        ]
        if table
        else [{"title": headings[index], "bullets": group} for index, group in enumerate(groups)]
    )
    visual_assignments, visual_report, model_usage = _presentation_visuals(
        payload,
        visual_topics,
        include_file=include_file,
        # Table and visual are composable capabilities. A structured deck uses
        # its visual on the summary slide, so it must be allowed to fill an
        # empty source slot through the configured image provider as well.
        allow_generated=True,
    )

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    outline: list[dict[str, Any]] = []
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    subtitle = title_slide.placeholders[1]
    subtitle.text = f"面向：{audience}"
    _style_title_slide(title_slide, title, subtitle)
    outline.append({"page": 1, "title": title, "bullets": [f"面向：{audience}"]})

    if table:
        for index, page in enumerate(table_pages, start=1):
            rows = page["rows"]
            heading = table["title"] or f"结构化数据（{index}/{len(table_pages)}）"
            if len(table_pages) > 1 and table["title"]:
                heading = f"{table['title']}（{index}/{len(table_pages)}）"
            if page.get("detail") is True:
                heading = f"{rows[0]['cells'][0]} · {rows[0]['cells'][1]}"
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = heading[:32]
            if page.get("detail") is True:
                _add_course_detail_slide(slide, table["columns"], rows[0])
            else:
                _add_table_slide(slide, table["columns"], rows)
            _style_slide(slide, title_color=RGBColor(72, 104, 183))
            source_locators = list(dict.fromkeys(row["source_locator"] for row in rows))
            outline.append(
                {
                    "page": len(outline) + 1,
                    "title": heading[:32],
                    "table": {
                        "columns": table["columns"],
                        "row_count": sum(not row.get("continuation") for row in rows),
                        "display_row_count": len(rows),
                        "detail": page.get("detail") is True,
                        "source_locators": source_locators,
                    },
                }
            )
    else:
        for index, group in enumerate(groups, start=1):
            heading = headings[index - 1]
            bullets = [value[:100] for value in group[:5]]
            visual = visual_assignments[index - 1] if index - 1 < len(visual_assignments) else None
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = heading
            frame = slide.placeholders[1].text_frame
            frame.clear()
            bullet_font_size = _presentation_bullet_font_size(
                bullets,
                has_visual=visual is not None,
            )
            for bullet_index, bullet in enumerate(bullets):
                paragraph = frame.paragraphs[0] if bullet_index == 0 else frame.add_paragraph()
                paragraph.text = _presentation_render_text(bullet)
                paragraph.level = 0
                paragraph.font.size = Pt(bullet_font_size)
            _style_body_placeholder(slide.placeholders[1], bullets, has_visual=visual is not None)
            _style_slide(slide, title_color=PRESENTATION_ACCENT_COLOR)
            if visual is not None:
                _add_presentation_visual(slide, visual)
            outline_item: dict[str, Any] = {
                "page": len(outline) + 1,
                "title": heading,
                "bullets": bullets,
            }
            if visual is not None:
                outline_item["visual"] = _presentation_visual_outline(visual)
            outline.append(outline_item)

    summary_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    summary_title = "内容概览" if table else "要点回顾"
    summary_slide.shapes.title.text = summary_title
    summary_bullets = (
        [
            f"共整理 {len(table['rows'])} 条记录",
            "全部内容均已完整展示",
            "详情见表格，来源见页脚",
        ]
        if table
        else (
            _representative_presentation_headings(section_specs, maximum=5)
            if section_specs
            else ["核心内容已按主题分组", "请结合实际场景确认后续行动"]
        )
    )
    summary_frame = summary_slide.placeholders[1].text_frame
    summary_frame.clear()
    for index, bullet in enumerate(summary_bullets):
        paragraph = summary_frame.paragraphs[0] if index == 0 else summary_frame.add_paragraph()
        paragraph.text = bullet
        paragraph.font.size = Pt(22)
    summary_visual = visual_assignments[0] if table and visual_assignments else None
    _style_body_placeholder(
        summary_slide.placeholders[1],
        summary_bullets,
        has_visual=summary_visual is not None,
    )
    _style_slide(summary_slide, title_color=PRESENTATION_TITLE_COLOR)
    if summary_visual is not None:
        _add_presentation_visual(summary_slide, summary_visual)
    summary_outline: dict[str, Any] = {
        "page": len(outline) + 1,
        "title": summary_title,
        "bullets": summary_bullets,
    }
    if summary_visual is not None:
        summary_outline["visual"] = _presentation_visual_outline(summary_visual)
    outline.append(summary_outline)

    quality_report = _presentation_quality_report(
        outline=outline,
        style=style,
        table=table,
        task_contract=task_contract,
        source_coverage=source_coverage,
        display_rows=display_rows,
        layout_violations=_presentation_layout_violations(presentation),
        visual_report=visual_report,
        filename=payload.get("filename"),
        brief=brief,
        requested_slide_count=requested_slide_count,
    )

    files: list[dict[str, Any]] = []
    if include_file and quality_report["passed"] is True:
        output = io.BytesIO()
        presentation.save(output)
        files.append(
            _file(
                _pptx_output_name(payload.get("filename"), title),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                output.getvalue(),
            )
        )
    return {
        "output": {
            "title": title,
            "audience": audience,
            "style": style,
            "slide_count": len(outline),
            "requested_slide_count": requested_slide_count,
            "outline": outline,
            "source_coverage": source_coverage,
            "visual_report": visual_report,
            "model_usage": model_usage,
            "quality_report": quality_report,
            "source_overwritten": False,
        },
        "files": files,
    }


def _extract_document(payload: dict[str, Any]) -> dict[str, Any]:
    source_name, source = _source_file(
        payload,
        {".pdf", ".txt", ".md"},
        MAX_DOCUMENT_SOURCE_BYTES,
    )
    media_type = _required_text(payload, "media_type", 200)
    round_start = payload.get("round_start", 1)
    if (
        not isinstance(round_start, int)
        or isinstance(round_start, bool)
        or not 1 <= round_start <= 100
    ):
        raise ValueError("document_round_start_must_be_between_1_and_100")
    try:
        parsed = parse_document(source, media_type)
    except UnicodeDecodeError as exc:
        raise ValueError("text_document_must_be_utf8") from exc
    except ValueError as exc:
        code = str(exc)
        if code == "too_many_pages":
            raise ValueError("pdf_has_too_many_pages") from exc
        if code == "no_extractable_text":
            raise ValueError("document_has_no_extractable_text") from exc
        raise
    extracted, manifest, context_tokens, cleaning_report = _document_context_rounds(
        parsed,
        MAX_DOCUMENT_CONTEXT_TOKENS,
        MAX_DOCUMENT_CONTEXT_ROUNDS,
        round_start,
    )
    selected_chunks = sum(int(item["chunk_count"]) for item in manifest["rounds"])
    original_character_count = sum(len(page.text) for page in parsed.pages)
    low_quality_pages = [page.page_no for page in parsed.pages if page.quality < 0.5]
    source_ir = _source_ir_slice(parsed.source_ir, manifest)
    return {
        "output": {
            "source_filename": source_name,
            "media_type": media_type,
            "format": "markdown",
            "parser_version": parsed.parser_version,
            "page_count": len(parsed.pages),
            "character_count": original_character_count,
            "token_count": context_tokens,
            "selected_chunk_count": selected_chunks,
            "total_chunk_count": len(parsed.chunks),
            "text": extracted,
            "truncated": selected_chunks < len(parsed.chunks),
            "round_count": manifest["round_count"],
            "completed_rounds": manifest["completed_rounds"],
            "round_start": manifest["round_start"],
            "next_round": manifest["next_round"],
            "has_more": manifest["has_more"],
            "coverage_ratio": manifest["coverage_ratio"],
            "rounds": manifest["rounds"],
            "source_ir": source_ir,
            "cleaning_report": cleaning_report,
            "low_quality_pages": low_quality_pages,
            "source_overwritten": False,
        },
        "files": [],
    }


def _document_context(parsed: ParseResult, token_budget: int) -> tuple[str, int, int]:
    chunks = parsed.chunks
    if not chunks:
        raise ValueError("document_has_no_extractable_text")
    content_budget = max(256, int(token_budget * 0.9))
    total_tokens = sum(chunk.token_count for chunk in chunks)
    if total_tokens <= content_budget:
        selected = list(chunks)
    else:
        slots = max(2, min(len(chunks), content_budget // 600))
        indices = [0, len(chunks) - 1]
        indices.extend(
            round(index * (len(chunks) - 1) / (slots - 1)) for index in range(1, slots - 1)
        )
        selected = []
        selected_tokens = 0
        for index in dict.fromkeys(indices):
            chunk = chunks[index]
            if selected and selected_tokens + chunk.token_count > content_budget:
                continue
            selected.append(chunk)
            selected_tokens += chunk.token_count
        for chunk in chunks:
            if chunk in selected or selected_tokens + chunk.token_count > content_budget:
                continue
            selected.append(chunk)
            selected_tokens += chunk.token_count
        selected.sort(key=lambda chunk: chunk.ordinal)
    blocks: list[str] = []
    for chunk in selected:
        marker = f"[[PAGE {chunk.page_start}]]"
        if chunk.page_end > chunk.page_start:
            marker = f"[[PAGES {chunk.page_start}-{chunk.page_end}]]"
        section = f"<!-- section: {chunk.section_path} -->\n" if chunk.section_path else ""
        blocks.append(f"{marker}\n{section}{chunk.content}".strip())
    context = "\n\n".join(blocks)
    while len(selected) > 1 and count_tokens(context) > token_budget:
        selected.pop()
        blocks.pop()
        context = "\n\n".join(blocks)
    return context, len(selected), count_tokens(context)


def _document_context_rounds(
    parsed: ParseResult,
    token_budget_per_round: int,
    max_rounds: int,
    round_start: int = 1,
) -> tuple[str, dict[str, Any], int, dict[str, Any]]:
    if not parsed.chunks:
        raise ValueError("document_has_no_extractable_text")
    if token_budget_per_round <= 0:
        token_budget_per_round = MAX_DOCUMENT_CONTEXT_TOKENS
    if max_rounds <= 0:
        max_rounds = MAX_DOCUMENT_CONTEXT_ROUNDS

    rounds: list[list[Any]] = []
    current: list[Any] = []
    current_tokens = 0
    for chunk in parsed.chunks:
        chunk_tokens = max(1, int(chunk.token_count)) + 16
        if current and current_tokens + chunk_tokens > token_budget_per_round:
            rounds.append(current)
            current = []
            current_tokens = 0
        current.append(chunk)
        current_tokens += chunk_tokens
    if current:
        rounds.append(current)

    start_index = min(max(0, round_start - 1), len(rounds))
    if start_index >= len(rounds):
        raise ValueError("document_round_start_out_of_range")
    end_index = min(len(rounds), start_index + max_rounds)
    completed = rounds[start_index:end_index]
    blocks: list[str] = []
    round_manifest: list[dict[str, Any]] = []
    duplicate_lines_removed = 0
    blank_lines_collapsed = 0
    selected_chunks = 0
    for round_index, chunks in enumerate(completed, start=start_index + 1):
        round_blocks: list[str] = []
        round_tokens = 0
        for chunk in chunks:
            cleaned, duplicates, blanks = _clean_context_text(chunk.content)
            duplicate_lines_removed += duplicates
            blank_lines_collapsed += blanks
            marker = f"[[PAGE {chunk.page_start}]]"
            if chunk.page_end > chunk.page_start:
                marker = f"[[PAGES {chunk.page_start}-{chunk.page_end}]]"
            section = f"<!-- section: {chunk.section_path} -->\n" if chunk.section_path else ""
            round_blocks.append(f"{marker}\n{section}{cleaned}".strip())
            round_tokens += max(1, int(chunk.token_count)) + 16
            selected_chunks += 1
        # Preserve the extraction boundary in the trusted observation.  The
        # Agent uses this marker to process oversized sources one bounded
        # round at a time instead of feeding the re-joined document back to a
        # single model call.  Page markers remain unchanged for citations.
        blocks.append(f"[[DOCUMENT ROUND {round_index}]]\n" + "\n\n".join(round_blocks))
        round_manifest.append(
            {
                "round_no": round_index,
                "chunk_start": int(chunks[0].ordinal),
                "chunk_end": int(chunks[-1].ordinal),
                "chunk_count": len(chunks),
                "token_count": round_tokens,
                "page_start": min(int(chunk.page_start) for chunk in chunks),
                "page_end": max(int(chunk.page_end) for chunk in chunks),
            }
        )

    context = "\n\n".join(blocks)
    manifest = {
        "round_count": len(rounds),
        "completed_rounds": len(completed),
        "round_start": round_start,
        "next_round": end_index + 1,
        "has_more": end_index < len(rounds),
        "coverage_ratio": selected_chunks / len(parsed.chunks),
        "rounds": round_manifest,
    }
    cleaning_report = {
        "policy_version": "document-cleaning-v1",
        "duplicate_lines_removed": duplicate_lines_removed,
        "blank_lines_collapsed": blank_lines_collapsed,
    }
    return context, manifest, count_tokens(context), cleaning_report


def _source_ir_slice(source_ir: dict[str, Any], manifest: dict[str, Any]) -> dict[str, Any]:
    selected_pages: set[int] = set()
    for item in manifest.get("rounds", []):
        if not isinstance(item, dict):
            continue
        start = int(item.get("page_start") or 0)
        end = int(item.get("page_end") or start)
        selected_pages.update(range(start, end + 1))
    tables: list[dict[str, Any]] = []
    for raw_table in source_ir.get("tables", []):
        if not isinstance(raw_table, dict):
            continue
        rows = [
            dict(row)
            for row in raw_table.get("rows", [])
            if isinstance(row, dict) and int(row.get("page") or 0) in selected_pages
        ]
        if not rows:
            continue
        row_ids = {str(row.get("id") or "") for row in rows}
        groups: list[dict[str, Any]] = []
        for raw_group in raw_table.get("row_groups", []):
            if not isinstance(raw_group, dict):
                continue
            selected_ids = [
                str(value) for value in raw_group.get("row_ids", []) if str(value) in row_ids
            ]
            if not selected_ids:
                continue
            group = dict(raw_group)
            group["row_ids"] = selected_ids
            group["partial"] = len(selected_ids) < len(raw_group.get("row_ids", []))
            groups.append(group)
        table = dict(raw_table)
        table["rows"] = rows
        table["row_groups"] = groups
        table["partial"] = len(rows) < len(raw_table.get("rows", []))
        tables.append(table)
    blocks = [
        dict(block)
        for block in source_ir.get("blocks", [])
        if isinstance(block, dict)
        and any(
            page in selected_pages
            for page in range(
                int(block.get("page_start") or 0),
                int(block.get("page_end") or block.get("page_start") or 0) + 1,
            )
        )
    ]
    return {
        "version": str(source_ir.get("version") or "document-source-ir-v1"),
        "structure_preserved": bool(tables),
        "selected_pages": sorted(selected_pages),
        "tables": tables,
        "blocks": blocks,
    }


def _clean_context_text(value: str) -> tuple[str, int, int]:
    lines = value.replace("\r\n", "\n").split("\n")
    cleaned: list[str] = []
    previous = ""
    previous_blank = False
    duplicate_lines_removed = 0
    blank_lines_collapsed = 0
    for raw_line in lines:
        line = raw_line.rstrip()
        blank = not line.strip()
        if blank and previous_blank:
            blank_lines_collapsed += 1
            continue
        if not blank and line == previous:
            duplicate_lines_removed += 1
            continue
        cleaned.append(line)
        previous = line
        previous_blank = blank
    return "\n".join(cleaned).strip(), duplicate_lines_removed, blank_lines_collapsed


def _profile_tabular(payload: dict[str, Any]) -> dict[str, Any]:
    source_name, source = _source_file(payload, {".csv", ".xlsx"})
    suffix = PurePath(source_name).suffix.lower()
    sheet_name = "CSV"
    if suffix == ".csv":
        rows = _read_csv(source)
    else:
        _validate_ooxml(source, reject_macros=True)
        rows, sheet_name = _read_xlsx(source)
    if not rows:
        raise ValueError("tabular_file_is_empty")
    headers = _headers(rows[0])
    body = [
        list(row[: len(headers)]) + [None] * max(0, len(headers) - len(row)) for row in rows[1:]
    ]
    body = [row[: len(headers)] for row in body]
    profiles = [
        _column_profile(header, [row[index] for row in body])
        for index, header in enumerate(headers)
    ]
    duplicate_rows = len(body) - len({_stable_row(row) for row in body})
    report = {
        "analysis_version": "office-tools-tabular-v1",
        "source_filename": source_name,
        "sheet_name": sheet_name,
        "row_count": len(body),
        "column_count": len(headers),
        "duplicate_rows": duplicate_rows,
        "truncated": len(body) >= MAX_ROWS,
        "columns": profiles,
        "source_overwritten": False,
    }
    report_name = _safe_filename(PurePath(source_name).stem + "-analysis", ".json")
    report_data = json.dumps(report, ensure_ascii=False, indent=2, default=str).encode("utf-8")
    return {
        "output": report,
        "files": [_file(report_name, "application/json; charset=utf-8", report_data)],
    }


def _translate_pdf(payload: dict[str, Any]) -> dict[str, Any]:
    source_name, source = _source_file(payload, {".pdf"}, MAX_PDF_SOURCE_BYTES)
    language = translation_language_policy(_required_text(payload, "target_language", 80))
    try:
        layout = extract_pdf_layout(source, max_pages=100)
    except ValueError as exc:
        code = str(exc)
        if code == "too_many_pages":
            raise ValueError("pdf_has_too_many_pages") from exc
        if code == "no_extractable_text":
            raise ValueError("pdf_has_no_extractable_text") from exc
        raise
    source_character_count = sum(len(block.text) for block in layout.blocks)
    if source_character_count > MAX_PDF_TRANSLATION_CHARS:
        raise ValueError("pdf_text_is_too_long_for_translation")
    output_name = _pdf_output_name(
        payload.get("output_filename"), source_name, language.output_label
    )
    translations, model_usage, translation_round_count = _translate_layout_blocks(
        layout,
        language.model_label,
    )
    try:
        output, layout_report = render_layout_translation(
            layout,
            translations,
            font_path=_embedded_cjk_font_path(),
            simplified_chinese=language.simplified_chinese,
        )
    except Exception as exc:
        raise ModelBackedOperationError("translation_render_failed", model_usage) from exc
    return {
        "output": {
            "source_filename": source_name,
            "output_filename": output_name,
            "target_language": language.output_label,
            "page_count": layout.page_count,
            "parser_version": layout.parser_version,
            "source_overwritten": False,
            "model_usage": model_usage,
            "translation_round_count": translation_round_count,
            "visual_context_page_count": len(layout.page_previews),
            **layout_report,
        },
        "files": [_file(output_name, "application/pdf", output)],
    }


def _translate_layout_blocks(
    layout: PdfLayout,
    target_language: str,
) -> tuple[dict[str, str], dict[str, Any], int]:
    batches = _translation_batches(layout.blocks)
    if not batches:
        return {}, _aggregate_translation_usage([], _translation_cost_limit()), 0

    cost_limit = _translation_cost_limit()
    concurrency = min(
        len(batches),
        _bounded_positive_int("MODEL_TRANSLATION_CONCURRENCY", 2, 4),
    )
    started_ns = time.perf_counter_ns()
    results: list[dict[str, str] | None] = [None] * len(batches)
    usages: list[dict[str, Any] | None] = [None] * len(batches)
    used_cost = 0
    reserved_cost = 0
    next_batch = 0
    first_failure: tuple[int, Exception] | None = None
    pending: dict[Future[tuple[dict[str, str], dict[str, Any]]], tuple[int, int]] = {}

    with ThreadPoolExecutor(
        max_workers=concurrency,
        thread_name_prefix="pdf-translation",
    ) as executor:
        while next_batch < len(batches) or pending:
            while (
                first_failure is None and next_batch < len(batches) and len(pending) < concurrency
            ):
                available_cost = cost_limit - used_cost - reserved_cost
                open_slots = concurrency - len(pending)
                if available_cost <= 0 or open_slots <= 0:
                    first_failure = (
                        next_batch,
                        ValueError("translation_model_budget_exhausted"),
                    )
                    break
                # Reserve the entire remaining budget across the currently open
                # execution slots. A completed call releases the difference
                # between its reservation and reported cost before another batch
                # can start, so concurrent calls can never multiply the run cap.
                reservation = max(1, available_cost // open_slots)
                batch_index = next_batch
                future = executor.submit(
                    _translate_layout_batch,
                    layout,
                    batches[batch_index],
                    target_language,
                    reservation,
                )
                pending[future] = (batch_index, reservation)
                reserved_cost += reservation
                next_batch += 1

            if not pending:
                break

            completed, _ = wait(tuple(pending), return_when=FIRST_COMPLETED)
            for future in sorted(completed, key=lambda item: pending[item][0]):
                batch_index, reservation = pending.pop(future)
                reserved_cost -= reservation
                try:
                    translated_batch, usage = future.result()
                except ModelBackedOperationError as exc:
                    usage = exc.model_usage
                    usages[batch_index] = usage
                    used_cost += _non_negative_int(usage.get("cost_micros"))
                    if first_failure is None:
                        first_failure = (batch_index, exc)
                except Exception as exc:
                    if first_failure is None:
                        first_failure = (batch_index, exc)
                else:
                    results[batch_index] = translated_batch
                    usages[batch_index] = usage
                    batch_cost = _non_negative_int(usage.get("cost_micros"))
                    used_cost += batch_cost
                    if batch_cost > reservation or used_cost + reserved_cost > cost_limit:
                        first_failure = first_failure or (
                            batch_index,
                            ModelBackedOperationError(
                                "translation_model_budget_exhausted",
                                usage,
                            ),
                        )

    ordered_usages = [usage for usage in usages if usage is not None]
    wall_latency_ms = max(0, (time.perf_counter_ns() - started_ns) // 1_000_000)
    aggregate = _aggregate_translation_usage(
        ordered_usages,
        cost_limit,
        wall_latency_ms=wall_latency_ms,
        concurrency=concurrency,
    )
    if first_failure is not None:
        _, failure = first_failure
        if isinstance(failure, ModelBackedOperationError):
            raise ModelBackedOperationError(failure.code, aggregate) from failure
        if ordered_usages:
            raise ModelBackedOperationError(str(failure), aggregate) from failure
        raise failure

    translations: dict[str, str] = {}
    for batch_result in results:
        if batch_result is None:
            raise ModelBackedOperationError("translation_model_incomplete", aggregate)
        translations.update(batch_result)
    return translations, aggregate, len(batches)


def _translate_layout_batch(
    layout: PdfLayout,
    batch: list[PdfTextBlock],
    target_language: str,
    reserved_cost_micros: int,
) -> tuple[dict[str, str], dict[str, Any]]:
    text = "\n\n".join(f"{block.marker}\n{block.text}" for block in batch)
    page_numbers = sorted({block.page_no for block in batch})
    previews = [
        (page_no, layout.page_previews[page_no])
        for page_no in page_numbers
        if page_no in layout.page_previews
    ]
    translated, usage = _openrouter_translate(
        text,
        target_language,
        page_previews=previews,
        remaining_cost_micros=reserved_cost_micros,
    )
    try:
        return _parse_block_translations(translated, batch), usage
    except ValueError as exc:
        raise ModelBackedOperationError(str(exc), usage) from exc


def _translation_batches(
    blocks: tuple[PdfTextBlock, ...],
) -> list[list[PdfTextBlock]]:
    batches: list[list[PdfTextBlock]] = []
    current: list[PdfTextBlock] = []
    current_chars = 0
    current_pages: set[int] = set()
    for block in blocks:
        block_chars = len(block.marker) + len(block.text) + 2
        next_pages = current_pages | {block.page_no}
        if current and (
            current_chars + block_chars > MAX_TRANSLATION_BATCH_CHARS
            or len(next_pages) > MAX_TRANSLATION_BATCH_PAGES
        ):
            batches.append(current)
            current = []
            current_chars = 0
            current_pages = set()
        current.append(block)
        current_chars += block_chars
        current_pages.add(block.page_no)
    if current:
        batches.append(current)
    return batches


def _parse_block_translations(
    translated: str,
    blocks: list[PdfTextBlock],
) -> dict[str, str]:
    marker_pattern = re.compile(
        r"^[ \t]*\[\[PAGE\s+(\d+)\s+BLOCK\s+(\d+)\]\][ \t]*\r?$",
        re.IGNORECASE | re.MULTILINE,
    )
    matches = list(marker_pattern.finditer(translated))
    expected = [block.marker for block in blocks]
    actual = [f"[[PAGE {int(match.group(1))} BLOCK {int(match.group(2))}]]" for match in matches]
    if actual != expected:
        raise ValueError("translation_block_markers_invalid")
    output: dict[str, str] = {}
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(translated)
        value = translated[start:end].strip()
        if not value:
            raise ValueError("translation_model_returned_empty_block")
        output[actual[index]] = value
    return output


def _aggregate_translation_usage(
    usages: list[dict[str, Any]],
    max_cost_micros: int,
    *,
    wall_latency_ms: int = 0,
    concurrency: int = 1,
) -> dict[str, Any]:
    first = usages[0] if usages else {}
    returned_models = {
        str(usage.get("returned_model", "")) for usage in usages if usage.get("returned_model")
    }
    providers = {
        str(usage.get("upstream_provider", ""))
        for usage in usages
        if usage.get("upstream_provider")
    }
    return {
        "provider": first.get("provider", "openrouter"),
        "config_version": first.get(
            "config_version",
            os.environ.get("MODEL_CONFIG_VERSION", "2026-08-structured-composer-v4"),
        ),
        "requested_model": first.get("requested_model", ""),
        "returned_model": (
            returned_models.pop()
            if len(returned_models) == 1
            else ("mixed" if returned_models else "")
        ),
        "upstream_provider": (
            providers.pop() if len(providers) == 1 else ("mixed" if providers else "")
        ),
        "prompt_tokens": sum(_non_negative_int(usage.get("prompt_tokens")) for usage in usages),
        "completion_tokens": sum(
            _non_negative_int(usage.get("completion_tokens")) for usage in usages
        ),
        "cost_micros": sum(_non_negative_int(usage.get("cost_micros")) for usage in usages),
        "cost_accounting": (
            "reported"
            if usages and all(usage.get("cost_accounting") == "reported" for usage in usages)
            else "reserved_upper_bound"
        ),
        "latency_ms": sum(_non_negative_int(usage.get("latency_ms")) for usage in usages),
        "wall_latency_ms": max(0, wall_latency_ms),
        "concurrency": max(1, concurrency),
        "max_cost_micros": max_cost_micros,
        "round_count": len(usages),
        "vision_page_count": sum(
            _non_negative_int(usage.get("vision_page_count")) for usage in usages
        ),
    }


def _translation_cost_limit() -> int:
    return min(
        _bounded_positive_int(
            "MODEL_TRANSLATION_MAX_COST_MICROS",
            MAX_TRANSLATION_COST_MICROS,
            MAX_TRANSLATION_COST_MICROS,
        ),
        MAX_TRANSLATION_COST_MICROS,
    )


def _openrouter_translate(
    text: str,
    target_language: str,
    *,
    page_previews: list[tuple[int, bytes]] | None = None,
    remaining_cost_micros: int | None = None,
) -> tuple[str, dict[str, Any]]:
    provider = os.environ.get("MODEL_PROVIDER", "development").strip().lower()
    if provider != "openrouter":
        raise ValueError("translation_model_provider_must_be_openrouter")
    base_url = os.environ.get("MODEL_BASE_URL", "https://openrouter.ai/api/v1").rstrip("/")
    if base_url != "https://openrouter.ai/api/v1":
        raise ValueError("MODEL_BASE_URL_must_use_canonical_openrouter_endpoint")
    api_key = os.environ.get("MODEL_API_KEY", "").strip()
    model = os.environ.get("MODEL_TRANSLATION_NAME", "openai/gpt-5-mini").strip()
    if not api_key or not model:
        raise ValueError("translation_model_is_not_configured")
    if _dynamic_model(model):
        raise ValueError("translation_model_must_use_a_concrete_slug")
    configured_max_tokens = _bounded_positive_int("MODEL_TRANSLATION_MAX_TOKENS", 8_000, 32_768)
    max_cost_micros = _translation_cost_limit()
    if remaining_cost_micros is not None:
        max_cost_micros = min(max_cost_micros, max(0, remaining_cost_micros))
    prompt_price = _positive_decimal("MODEL_TRANSLATION_MAX_PROMPT_PRICE", "0.3")
    completion_price = _positive_decimal("MODEL_TRANSLATION_MAX_COMPLETION_PRICE", "2.5")
    if prompt_price > Decimal("0.3") or completion_price > Decimal("2.5"):
        raise ValueError("translation_model_price_ceiling_is_too_high")
    system_message = {
        "role": "system",
        "content": (
            f"Translate the supplied PDF text into {target_language}. "
            "The page images are visual context for charts, labels and layout only. "
            "Preserve every [[PAGE n]] or [[PAGE n BLOCK m]] marker exactly, on its own line, "
            "and in the original order. Translate every block completely without merging, "
            "omitting or summarizing it. Keep the translation concise enough to fit the same "
            "text box. Return only markers and translated text, without commentary."
            " Use plain ASCII hyphens for list bullets, avoid emoji or decorative symbols, "
            "and preserve mathematical notation as text."
        ),
    }
    previews = page_previews or []
    if previews:
        user_content: str | list[dict[str, Any]] = [{"type": "text", "text": text}]
        for page_no, preview in previews:
            user_content.extend(
                [
                    {"type": "text", "text": f"Visual reference for PDF page {page_no}:"},
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": "data:image/jpeg;base64,"
                            + base64.b64encode(preview).decode("ascii"),
                            "detail": "low",
                        },
                    },
                ]
            )
    else:
        user_content = text
    messages = [system_message, {"role": "user", "content": user_content}]
    text_only_messages = [system_message, {"role": "user", "content": text}]
    prompt_bound = (
        len(
            json.dumps(
                text_only_messages,
                ensure_ascii=False,
                separators=(",", ":"),
            ).encode("utf-8")
        )
        + len(previews) * 1_024
        + 128
    )
    reserved_prompt_cost = _decimal_ceil(Decimal(prompt_bound) * prompt_price)
    affordable_output_tokens = int(
        Decimal(max(0, max_cost_micros - reserved_prompt_cost)) / completion_price
    )
    max_tokens = min(configured_max_tokens, affordable_output_tokens)
    if max_tokens < 256:
        raise ValueError("translation_model_budget_exhausted")
    request_body = json.dumps(
        {
            "model": model,
            "messages": messages,
            "max_tokens": max_tokens,
            "reasoning": {"effort": "minimal", "exclude": True},
            "provider": {
                "data_collection": os.environ.get("MODEL_DATA_COLLECTION", "deny"),
                "zdr": _env_bool("MODEL_ZDR_REQUIRED", False),
                "sort": os.environ.get("MODEL_PROVIDER_SORT", "price"),
                "allow_fallbacks": _env_bool("MODEL_ALLOW_PROVIDER_FALLBACKS", True),
                "require_parameters": True,
                "max_price": {
                    "prompt": float(prompt_price),
                    "completion": float(completion_price),
                },
            },
        },
        ensure_ascii=False,
    ).encode("utf-8")
    request = urllib.request.Request(
        base_url + "/chat/completions",
        data=request_body,
        method="POST",
        headers={
            "Authorization": "Bearer " + api_key,
            "Content-Type": "application/json",
            "HTTP-Referer": os.environ.get("MODEL_HTTP_REFERER", "http://localhost:3000"),
            "X-OpenRouter-Title": os.environ.get("MODEL_APP_TITLE", "伴AI"),
        },
    )
    started_ns = time.perf_counter_ns()
    try:
        with urllib.request.urlopen(request, timeout=90) as response:
            raw = response.read(MAX_MODEL_RESPONSE_BYTES + 1)
    except urllib.error.HTTPError as exc:
        raise ValueError(f"translation_model_http_{exc.code}") from exc
    except Exception as exc:
        raise ValueError("translation_model_request_failed") from exc
    latency_ms = max(0, (time.perf_counter_ns() - started_ns) // 1_000_000)
    model_usage = {
        "provider": "openrouter",
        "config_version": os.environ.get("MODEL_CONFIG_VERSION", "2026-08-structured-composer-v4"),
        "requested_model": model,
        "returned_model": "",
        "upstream_provider": "",
        "prompt_tokens": 0,
        "completion_tokens": 0,
        "cost_micros": max_cost_micros,
        "cost_accounting": "reserved_upper_bound",
        "latency_ms": latency_ms,
        "max_cost_micros": max_cost_micros,
        "vision_page_count": len(previews),
    }
    if len(raw) > MAX_MODEL_RESPONSE_BYTES:
        raise ModelBackedOperationError("translation_model_response_too_large", model_usage)
    try:
        payload = json.loads(raw)
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise ModelBackedOperationError(
            "translation_model_returned_invalid_json", model_usage
        ) from exc
    if not isinstance(payload, dict):
        raise ModelBackedOperationError("translation_model_returned_invalid_json", model_usage)
    usage = payload.get("usage")
    usage = usage if isinstance(usage, dict) else {}
    returned_model = payload.get("model")
    upstream_provider = payload.get("provider")
    reported_cost_micros = _reported_cost_micros(usage.get("cost"))
    model_usage.update(
        {
            "returned_model": (returned_model if isinstance(returned_model, str) else ""),
            "upstream_provider": (upstream_provider if isinstance(upstream_provider, str) else ""),
            "prompt_tokens": _non_negative_int(usage.get("prompt_tokens")),
            "completion_tokens": _non_negative_int(usage.get("completion_tokens")),
            "cost_micros": (
                reported_cost_micros if reported_cost_micros is not None else max_cost_micros
            ),
            "cost_accounting": (
                "reported" if reported_cost_micros is not None else "reserved_upper_bound"
            ),
        }
    )
    if reported_cost_micros is None:
        raise ModelBackedOperationError("translation_model_usage_invalid", model_usage)
    if reported_cost_micros > max_cost_micros:
        raise ModelBackedOperationError("translation_model_budget_exhausted", model_usage)
    try:
        finish_reason = payload["choices"][0].get("finish_reason")
    except (KeyError, IndexError, TypeError, AttributeError) as exc:
        raise ModelBackedOperationError("translation_model_response_invalid", model_usage) from exc
    if finish_reason != "stop":
        code = (
            "translation_model_truncated"
            if finish_reason == "length"
            else "translation_model_incomplete"
        )
        raise ModelBackedOperationError(code, model_usage)
    try:
        translated = _translated_content(payload)
    except ValueError as exc:
        code = str(exc)
        if not code.startswith("translation_model_"):
            code = "translation_model_response_invalid"
        raise ModelBackedOperationError(code, model_usage) from exc
    if _translation_markers(translated) != _translation_markers(text):
        raise ModelBackedOperationError("translation_page_markers_invalid", model_usage)
    return translated, model_usage


def _translated_content(payload: dict[str, Any]) -> str:
    try:
        content = payload["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError) as exc:
        raise ValueError("translation_model_response_invalid") from exc
    if isinstance(content, str):
        translated = content.strip()
    elif isinstance(content, list):
        translated = "\n".join(
            str(item.get("text", "")).strip()
            for item in content
            if isinstance(item, dict) and item.get("type") == "text"
        ).strip()
    else:
        translated = ""
    if not translated:
        raise ValueError("translation_model_returned_empty_text")
    return translated


def _translation_markers(text: str) -> list[tuple[int, int | None]]:
    return [
        (int(page), int(block) if block else None)
        for page, block in re.findall(
            r"^[ \t]*\[\[PAGE\s+(\d+)(?:\s+BLOCK\s+(\d+))?\]\][ \t]*\r?$",
            text,
            re.IGNORECASE | re.MULTILINE,
        )
    ]


def _embedded_cjk_font_path() -> str:
    configured = os.environ.get("PDF_CJK_FONT_PATH", "").strip()
    if configured:
        if os.path.isfile(configured):
            return configured
        raise ValueError("pdf_cjk_font_unavailable")
    for candidate in CJK_FONT_CANDIDATES:
        if os.path.isfile(candidate):
            return candidate
    raise ValueError("pdf_cjk_font_unavailable")


def _read_csv(data: bytes) -> list[list[Any]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("csv_must_be_utf8") from exc
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows: list[list[Any]] = []
    for index, row in enumerate(reader):
        if index > MAX_ROWS:
            break
        if len(row) > MAX_COLUMNS:
            raise ValueError("too_many_columns")
        rows.append(row)
    return rows


def _read_xlsx(data: bytes) -> tuple[list[list[Any]], str]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook.active
        rows: list[list[Any]] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index > MAX_ROWS:
                break
            values = list(row)
            if len(values) > MAX_COLUMNS:
                values = values[:MAX_COLUMNS]
            rows.append(values)
        return rows, sheet.title
    finally:
        workbook.close()


def _column_profile(name: str, values: list[Any]) -> dict[str, Any]:
    missing = sum(value is None or str(value).strip() == "" for value in values)
    present = [value for value in values if value is not None and str(value).strip() != ""]
    numbers = [float(value) for value in present if _is_number(value)]
    inferred = "empty"
    if present:
        inferred = (
            "number"
            if len(numbers) == len(present)
            else "boolean"
            if all(isinstance(value, bool) for value in present)
            else "text"
        )
    profile: dict[str, Any] = {
        "name": name,
        "inferred_type": inferred,
        "missing_count": missing,
        "unique_count": len({_stable_value(value) for value in present}),
    }
    if numbers:
        profile["numeric"] = {"min": min(numbers), "max": max(numbers), "mean": fmean(numbers)}
    return profile


def _headers(row: list[Any]) -> list[str]:
    if len(row) > MAX_COLUMNS:
        raise ValueError("too_many_columns")
    result: list[str] = []
    counts: Counter[str] = Counter()
    for index, value in enumerate(row, start=1):
        base = str(value).strip() if value is not None else ""
        base = base or f"column_{index}"
        counts[base] += 1
        result.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return result


def _source_file(
    payload: dict[str, Any], suffixes: set[str], maximum_bytes: int = MAX_SOURCE_BYTES
) -> tuple[str, bytes]:
    name = str(payload.get("source_filename") or "").strip()
    if PurePath(name).name != name or PurePath(name).suffix.lower() not in suffixes:
        raise ValueError("unsupported_source_filename")
    encoded = payload.get("source_base64")
    if not isinstance(encoded, str) or not encoded:
        raise ValueError("source_base64_is_required")
    try:
        data = base64.b64decode(encoded, validate=True)
    except ValueError as exc:
        raise ValueError("invalid_source_base64") from exc
    if not data or len(data) > maximum_bytes:
        raise ValueError("source_file_size_is_invalid")
    return name, data


def _validate_ooxml(data: bytes, reject_macros: bool) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(item.file_size for item in archive.infolist())
            names = {item.filename.lower() for item in archive.infolist()}
            if total > MAX_ARCHIVE_BYTES:
                raise ValueError("expanded_archive_is_too_large")
            if reject_macros and any(name.endswith("vbaproject.bin") for name in names):
                raise ValueError("macro_enabled_files_are_not_supported")
    except zipfile.BadZipFile as exc:
        raise ValueError("invalid_ooxml_file") from exc


def _required_text(payload: dict[str, Any], field: str, maximum: int) -> str:
    value = payload.get(field)
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field}_is_required")
    value = value.strip()
    if len(value) > maximum:
        raise ValueError(f"{field}_is_too_long")
    return value


def _output_name(value: Any, source_name: str, fallback_suffix: str) -> str:
    if isinstance(value, str) and value.strip():
        return _safe_filename(value, ".docx")
    return _safe_filename(PurePath(source_name).stem + fallback_suffix, ".docx")


def _pdf_output_name(value: Any, source_name: str, target_language: str) -> str:
    if isinstance(value, str) and value.strip():
        return _safe_filename(value, ".pdf")
    language = re.sub(r"[^\w\u4e00-\u9fff]+", "-", target_language, flags=re.UNICODE).strip("-")
    language = language or "translated"
    return _safe_filename(PurePath(source_name).stem + f"-{language}", ".pdf")


def _pptx_output_name(value: Any, title: str) -> str:
    if isinstance(value, str) and value.strip():
        return _safe_filename(value, ".pptx")
    title_filename = re.sub(r"[/\\]+", "-", title)
    title_filename = re.sub(r"\s*-\s*", "-", title_filename)
    return _safe_filename(title_filename, ".pptx")


def _safe_filename(value: str, suffix: str) -> str:
    raw_name = value.strip()
    name = PurePath(raw_name).name.strip()
    if name != raw_name or "/" in raw_name or "\\" in raw_name:
        raise ValueError("invalid_output_filename")
    if name.lower().endswith(suffix):
        name = name[: -len(suffix)]
    name = re.sub(r"[^\w\-.\u4e00-\u9fff]+", "-", name, flags=re.UNICODE).strip(".-")
    if not name:
        name = "output"
    return name[: 240 - len(suffix)] + suffix


def _slide_bullets(sections: list[str], offset: int, brief: str) -> list[str]:
    primary = sections[offset % len(sections)]
    candidates = [part.strip(" -•\t") for part in re.split(r"[。；;]", primary) if part.strip()]
    if len(candidates) < 2:
        candidates.extend([part for part in sections if part != primary])
    if len(candidates) < 2:
        candidates.append(brief[:120])
    return [value[:120] for value in candidates[:4]]


def _presentation_units(brief: str) -> list[str]:
    units: list[str] = []
    for block in re.split(r"[\r\n]+|(?<=[。；;!?！？])\s*", brief):
        block = block.strip(" -•\t")
        if not block:
            continue
        if len(block) > 180:
            parts = [part.strip() for part in re.split(r"[、,，]", block) if part.strip()]
            if len(parts) > 1:
                units.extend(parts)
                continue
        units.append(block)
    return units or [brief.strip()]


def _presentation_section_specs(brief: str) -> list[dict[str, Any]]:
    """Parse the canonical, audience-facing non-table presentation format."""

    sections: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for raw_line in brief.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = re.fullmatch(r"##\s+(.+?)\s*", line)
        if heading:
            current = {"heading": heading.group(1).strip()[:28], "bullets": []}
            sections.append(current)
            continue
        bullet = re.fullmatch(r"-\s+(.+?)\s*", line)
        if bullet and current is not None:
            current["bullets"].append(bullet.group(1).strip()[:100])
            continue
        return []
    if not sections or any(len(section["bullets"]) < 2 for section in sections):
        return []
    return sections


def _ensure_presentation_units(units: list[str], count: int) -> list[str]:
    result = list(units)
    if count > len(result):
        preview = "、".join(value[:24] for value in units[:5])
        result.insert(0, f"主题概览：{preview}"[:100])
    fallbacks = ("关键事实与背景", "需要关注的重点", "后续讨论与行动")
    for value in fallbacks:
        if len(result) >= count:
            break
        if value not in result:
            result.append(value)
    return result


def _representative_presentation_headings(
    sections: list[dict[str, Any]], *, maximum: int
) -> list[str]:
    """Sample the whole deck for the recap instead of only its opening pages."""

    if len(sections) <= maximum:
        return [str(section["heading"]) for section in sections]
    indices = [round(index * (len(sections) - 1) / (maximum - 1)) for index in range(maximum)]
    return [str(sections[index]["heading"]) for index in dict.fromkeys(indices)]


def _balanced_groups(items: list[Any], maximum_groups: int) -> list[list[Any]]:
    if not items:
        return []
    group_count = max(1, min(maximum_groups, len(items)))
    base, remainder = divmod(len(items), group_count)
    result: list[list[Any]] = []
    offset = 0
    for index in range(group_count):
        size = base + (1 if index < remainder else 0)
        result.append(items[offset : offset + size])
        offset += size
    return result


def _presentation_heading(value: str, index: int) -> str:
    cleaned = re.sub(r"\s+", " ", value).strip("：:。；;，,")
    if not cleaned:
        return f"主题 {index}"
    return cleaned[:28]


def _presentation_bullet_font_size(
    bullets: list[str],
    *,
    has_visual: bool = False,
) -> int:
    """Choose a readable size while avoiding punctuation-only orphan lines."""

    longest = max((len(value) for value in bullets), default=0)
    total_characters = sum(len(value) for value in bullets)
    if longest > 82 or len(bullets) >= 4:
        size = 17
    elif longest > 62:
        size = 18
    elif longest > 34 and len(bullets) >= 3:
        size = 17
    elif len(bullets) >= 3:
        size = 22
    elif longest <= 48:
        size = 24
    else:
        size = 20
    if not has_visual:
        return size
    # An image halves the available reading width. Five source-grounded
    # paragraphs can fit the placeholder rectangle while their wrapped lines
    # still spill below it, so reserve a stable dense-page size.
    if len(bullets) >= 5 or total_characters > 260 or longest > 88:
        return min(size, 15)
    if len(bullets) >= 4 or total_characters > 210 or longest > 68:
        return min(size, 16)
    return min(size, 18)


def _style_body_placeholder(
    placeholder: Any,
    bullets: list[str],
    *,
    has_visual: bool,
) -> None:
    """Give narrative text a stable reading column and consistent rhythm."""

    placeholder.left = Inches(0.78 if has_visual else 1.08)
    placeholder.top = Inches(1.42)
    placeholder.width = Inches(5.45 if has_visual else 10.86)
    placeholder.height = Inches(5.28)
    frame = placeholder.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.06)
    frame.margin_bottom = Inches(0.06)
    total_characters = sum(len(value) for value in bullets)
    frame.vertical_anchor = (
        MSO_ANCHOR.MIDDLE
        if not has_visual and len(bullets) <= 3 and total_characters <= 260
        else MSO_ANCHOR.TOP
    )
    if has_visual:
        paragraph_spacing = 4 if len(bullets) >= 5 or total_characters > 260 else 6
        line_spacing = 1.0 if len(bullets) >= 4 or total_characters > 210 else 1.05
    else:
        paragraph_spacing = 16 if len(bullets) <= 2 else 12 if len(bullets) == 3 else 8
        line_spacing = 1.1
    for paragraph in frame.paragraphs:
        paragraph.font.color.rgb = PRESENTATION_TEXT_COLOR
        paragraph.space_after = Pt(paragraph_spacing)
        paragraph.line_spacing = line_spacing


def _presentation_render_text(value: str) -> str:
    """Prevent semantic identifiers from being split across rendered lines."""

    def protect(match: re.Match[str]) -> str:
        return "\u2060".join(match.group(0))

    text = re.sub(
        r"第\s*\d{1,4}\s*(?:页(?:\s*[-–—~～至到]\s*(?:第\s*)?\d{1,4}\s*页?)?"
        r"|[-–—~～至到]\s*(?:第\s*)?\d{1,4}\s*页)",
        protect,
        value,
    )
    text = re.sub(r"\b[A-Za-z]{2,12}\+[A-Za-z]{2,12}\b", protect, text)
    text = re.sub(r"\b[A-Za-z][A-Za-z0-9]*(?:_[A-Za-z0-9]+)+\b", protect, text)
    text = re.sub(r"\b(?:[A-Z][a-z0-9]+){2,}\b", protect, text)
    text = re.sub(r"\b[A-Za-z]{2,8}\d{2,8}[A-Za-z]?\b", protect, text)
    text = re.sub(r"(?<![\d.])\d+(?:\.\d+)?%", protect, text)
    return re.sub(r"《[^》\r\n]{1,30}》", protect, text)


def _presentation_topic_source_pages(topic: Mapping[str, Any]) -> set[int]:
    """Return source PDF pages explicitly cited by one audience topic."""

    text = "\n".join(
        [
            str(topic.get("title") or ""),
            *[str(value or "") for value in topic.get("bullets", [])],
        ]
    )
    pages: set[int] = set()
    for match in re.finditer(
        r"第\s*(\d{1,4})\s*(?:页(?:\s*[-–—~～至到]\s*(?:第\s*)?(\d{1,4})\s*页?)?"
        r"|[-–—~～至到]\s*(?:第\s*)?(\d{1,4})\s*页)",
        text,
    ):
        start = int(match.group(1))
        end = int(match.group(2) or match.group(3) or start)
        if start <= 0 or end <= 0 or abs(end - start) > 100:
            continue
        pages.update(range(min(start, end), max(start, end) + 1))
    return pages


def _presentation_visuals(
    payload: dict[str, Any],
    topics: list[dict[str, Any]],
    *,
    include_file: bool,
    allow_generated: bool,
) -> tuple[list[dict[str, Any] | None], dict[str, Any], dict[str, Any]]:
    mode = str(payload.get("visual_mode") or "auto").strip().lower()
    if mode not in {"auto", "source_only", "none"}:
        raise ValueError("presentation_visual_mode_is_invalid")
    cost_limit = _presentation_image_cost_limit()
    model = os.environ.get("MODEL_PRESENTATION_IMAGE_NAME", "").strip()
    empty_usage = _presentation_visual_model_usage([], cost_limit, model)
    if not include_file:
        return (
            [None] * len(topics),
            {
                "policy_version": "presentation-visuals-v1",
                "mode": "outline",
                "requested_slot_count": len(topics),
                "source_document_count": 0,
                "extracted_visual_count": 0,
                "generated_visual_count": 0,
                "placed_visual_count": 0,
                "source_visual_count": 0,
                "fallback_text_only_count": 0,
                "generation_attempt_count": 0,
                "generation_errors": [],
            },
            empty_usage,
        )

    source_documents = _presentation_source_documents(payload.get("source_documents"))
    extracted: list[dict[str, Any]] = []
    if mode != "none":
        for document in source_documents:
            if document["media_type"] == "application/pdf":
                extracted.extend(_extract_pdf_presentation_visuals(document))
            if len(extracted) >= MAX_PRESENTATION_SOURCE_VISUALS:
                break
    extracted = extracted[:MAX_PRESENTATION_SOURCE_VISUALS]
    assignments: list[dict[str, Any] | None] = [None] * len(topics)
    unused_visuals = list(extracted)
    for index, topic in enumerate(topics):
        cited_pages = _presentation_topic_source_pages(topic)
        matching_index = next(
            (
                visual_index
                for visual_index, visual in enumerate(unused_visuals)
                if int(visual.get("source_page") or 0) in cited_pages
            ),
            None,
        )
        if matching_index is not None:
            assignments[index] = unused_visuals.pop(matching_index)

    usages: list[dict[str, Any]] = []
    generation_errors: list[str] = []
    generation_attempts = 0
    generated_count = 0
    configured_maximum_generated = _bounded_positive_int(
        "MODEL_PRESENTATION_IMAGE_MAX_COUNT",
        2,
        4,
    )
    maximum_generated = min(configured_maximum_generated, len(topics))
    can_generate = (
        mode == "auto"
        and allow_generated
        and bool(model)
        and os.environ.get("MODEL_PROVIDER", "development").strip().lower() == "openrouter"
        and bool(os.environ.get("MODEL_API_KEY", "").strip())
    )
    if can_generate:
        for index, topic in enumerate(topics):
            if assignments[index] is not None:
                continue
            if generated_count >= maximum_generated or generation_attempts >= maximum_generated:
                break
            used_cost = sum(_non_negative_int(item.get("cost_micros")) for item in usages)
            remaining_cost = cost_limit - used_cost
            if remaining_cost <= 0:
                generation_errors.append("presentation_image_budget_exhausted")
                break
            generation_attempts += 1
            try:
                visual, usage = _openrouter_generate_presentation_visual(
                    topic,
                    remaining_cost_micros=remaining_cost,
                )
            except Exception as exc:
                code = str(exc).strip()
                generation_errors.append(
                    code if code.startswith("presentation_image_") else "presentation_image_failed"
                )
                continue
            assignments[index] = visual
            usages.append(usage)
            generated_count += 1

    source_visual_count = sum(
        visual is not None and visual.get("kind") == "source" for visual in assignments
    )
    placed_count = sum(visual is not None for visual in assignments)
    fallback_count = sum(visual is None for visual in assignments) if mode != "none" else 0
    report = {
        "policy_version": "presentation-visuals-v1",
        "mode": mode,
        "requested_slot_count": len(topics),
        "source_document_count": len(source_documents),
        "extracted_visual_count": len(extracted),
        "generated_visual_count": generated_count,
        "placed_visual_count": placed_count,
        "source_visual_count": source_visual_count,
        "fallback_text_only_count": fallback_count,
        "generation_attempt_count": generation_attempts,
        "generation_errors": list(dict.fromkeys(generation_errors))[:8],
    }
    return assignments, report, _presentation_visual_model_usage(usages, cost_limit, model)


def _presentation_source_documents(value: Any) -> list[dict[str, Any]]:
    if value is None:
        return []
    if not isinstance(value, list) or len(value) > MAX_PRESENTATION_SOURCE_DOCUMENTS:
        raise ValueError("presentation_source_documents_are_invalid")
    documents: list[dict[str, Any]] = []
    total_bytes = 0
    for raw in value:
        if not isinstance(raw, dict):
            raise ValueError("presentation_source_document_is_invalid")
        filename = str(raw.get("filename") or "").strip()
        media_type = str(raw.get("media_type") or "").strip().lower()
        encoded = raw.get("data_base64")
        if (
            PurePath(filename).name != filename
            or media_type != "application/pdf"
            or not isinstance(encoded, str)
            or not encoded
        ):
            raise ValueError("presentation_source_document_is_invalid")
        try:
            data = base64.b64decode(encoded, validate=True)
        except ValueError as exc:
            raise ValueError("presentation_source_document_is_invalid") from exc
        total_bytes += len(data)
        if not data or total_bytes > MAX_PRESENTATION_SOURCE_BYTES:
            raise ValueError("presentation_source_documents_are_too_large")
        documents.append({"filename": filename, "media_type": media_type, "data": data})
    return documents


def _extract_pdf_presentation_visuals(document: dict[str, Any]) -> list[dict[str, Any]]:
    pdf = pymupdf.open(stream=document["data"], filetype="pdf")
    candidates: list[dict[str, Any]] = []
    seen: set[str] = set()
    try:
        if pdf.needs_pass and not pdf.authenticate(""):
            return []
        for page_index, page in enumerate(pdf, start=1):
            page_area = max(1.0, float(page.rect.get_area()))
            page_candidates: list[dict[str, Any]] = []
            for image_info in page.get_image_info(xrefs=True):
                bbox = image_info.get("bbox")
                if not isinstance(bbox, (list, tuple)) or len(bbox) != 4:
                    continue
                rect = pymupdf.Rect(tuple(float(value) for value in bbox)) & page.rect
                width = int(image_info.get("width") or 0)
                height = int(image_info.get("height") or 0)
                area_ratio = float(rect.get_area()) / page_area
                aspect_ratio = float(rect.width) / max(1.0, float(rect.height))
                if (
                    width < 180
                    or height < 120
                    or rect.width < 90
                    or rect.height < 60
                    or area_ratio < 0.035
                    or area_ratio > 0.96
                    or not 0.25 <= aspect_ratio <= 4.0
                ):
                    continue
                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(1.5, 1.5),
                    clip=rect,
                    alpha=False,
                )
                data = pixmap.tobytes("jpeg", jpg_quality=84)
                if not data or len(data) > MAX_PRESENTATION_VISUAL_BYTES:
                    continue
                digest = hashlib.sha256(data).hexdigest()
                if digest in seen:
                    continue
                seen.add(digest)
                page_candidates.append(
                    {
                        "kind": "source",
                        "data": data,
                        "media_type": "image/jpeg",
                        "width": pixmap.width,
                        "height": pixmap.height,
                        "source_filename": document["filename"],
                        "source_page": page_index,
                        "source_locator": f"{document['filename']} · 第 {page_index} 页",
                        "score": area_ratio * min(width, 1600) * min(height, 1200),
                    }
                )
            if page_candidates:
                candidates.append(max(page_candidates, key=lambda item: float(item["score"])))
            if len(candidates) >= MAX_PRESENTATION_SOURCE_VISUALS:
                break
    finally:
        pdf.close()
    for candidate in candidates:
        candidate.pop("score", None)
    return candidates


def _openrouter_generate_presentation_visual(
    topic: dict[str, Any],
    *,
    remaining_cost_micros: int,
) -> tuple[dict[str, Any], dict[str, Any]]:
    base_url = os.environ.get("MODEL_BASE_URL", "https://openrouter.ai/api/v1").rstrip("/")
    if base_url != "https://openrouter.ai/api/v1":
        raise ValueError("presentation_image_endpoint_is_invalid")
    api_key = os.environ.get("MODEL_API_KEY", "").strip()
    model = os.environ.get("MODEL_PRESENTATION_IMAGE_NAME", "").strip()
    if not api_key or not model or _dynamic_model(model):
        raise ValueError("presentation_image_model_is_not_configured")
    title = re.sub(r"\s+", " ", str(topic.get("title") or "")).strip()[:160]
    bullets = [
        re.sub(r"\s+", " ", str(value or "")).strip()[:220]
        for value in topic.get("bullets", [])
        if str(value or "").strip()
    ][:4]
    prompt = (
        "Create one clean 16:9 editorial presentation illustration. "
        "Use a coherent professional visual metaphor, realistic or polished 3D style, "
        "ample negative space, and no words, letters, numerals, logos, watermarks, UI, "
        f"or decorative borders. Topic: {title}. Context: {'; '.join(bullets)}"
    )
    body = json.dumps(
        {
            "model": model,
            "prompt": prompt,
            "n": 1,
            "aspect_ratio": "16:9",
            "quality": "medium",
            "output_format": "jpeg",
            "output_compression": 84,
        },
        ensure_ascii=False,
    ).encode("utf-8")
    request = urllib.request.Request(
        base_url + "/images",
        data=body,
        method="POST",
        headers={
            "Authorization": "Bearer " + api_key,
            "Content-Type": "application/json",
            "HTTP-Referer": os.environ.get("MODEL_HTTP_REFERER", "http://localhost:3000"),
            "X-OpenRouter-Title": os.environ.get("MODEL_APP_TITLE", "伴AI"),
        },
    )
    started_ns = time.perf_counter_ns()
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            raw = response.read(MAX_PRESENTATION_IMAGE_RESPONSE_BYTES + 1)
    except urllib.error.HTTPError as exc:
        raise ValueError(f"presentation_image_http_{exc.code}") from exc
    except Exception as exc:
        raise ValueError("presentation_image_request_failed") from exc
    latency_ms = max(0, (time.perf_counter_ns() - started_ns) // 1_000_000)
    if len(raw) > MAX_PRESENTATION_IMAGE_RESPONSE_BYTES:
        raise ValueError("presentation_image_response_too_large")
    try:
        response_payload = json.loads(raw)
        image_item = response_payload["data"][0]
        image_data = base64.b64decode(image_item["b64_json"], validate=True)
    except (KeyError, IndexError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise ValueError("presentation_image_response_invalid") from exc
    image_data, media_type, width, height = _normalize_presentation_visual_image(image_data)
    usage = response_payload.get("usage")
    usage = usage if isinstance(usage, dict) else {}
    reported_cost = _reported_cost_micros(usage.get("cost"))
    cost_micros = reported_cost if reported_cost is not None else remaining_cost_micros
    return {
        "kind": "generated",
        "data": image_data,
        "media_type": media_type,
        "width": width,
        "height": height,
        "model": model,
        "prompt": prompt,
        "source_locator": "AI 生成配图",
    }, {
        "provider": "openrouter",
        "requested_model": model,
        "returned_model": str(response_payload.get("model") or model),
        "upstream_provider": str(response_payload.get("provider") or ""),
        "prompt_tokens": _non_negative_int(usage.get("prompt_tokens")),
        "completion_tokens": _non_negative_int(usage.get("completion_tokens")),
        "cost_micros": cost_micros,
        "cost_accounting": "reported" if reported_cost is not None else "reserved_upper_bound",
        "latency_ms": latency_ms,
    }


def _normalize_presentation_visual_image(data: bytes) -> tuple[bytes, str, int, int]:
    try:
        with Image.open(io.BytesIO(data)) as image:
            width, height = image.size
            if width < 320 or height < 180 or width * height > 25_000_000:
                raise ValueError("presentation_image_dimensions_are_invalid")
            image.load()
            image_format = str(image.format or "").upper()
            if image_format in {"JPEG", "PNG"} and len(data) <= MAX_PRESENTATION_VISUAL_BYTES:
                return data, "image/jpeg" if image_format == "JPEG" else "image/png", width, height
            converted = image.convert("RGB")
            output = io.BytesIO()
            converted.save(output, format="JPEG", quality=84, optimize=True)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("presentation_image_is_invalid") from exc
    normalized = output.getvalue()
    if len(normalized) > MAX_PRESENTATION_VISUAL_BYTES:
        raise ValueError("presentation_image_is_too_large")
    return normalized, "image/jpeg", width, height


def _presentation_image_cost_limit() -> int:
    return min(
        _bounded_positive_int(
            "MODEL_PRESENTATION_IMAGE_MAX_COST_MICROS",
            MAX_PRESENTATION_IMAGE_COST_MICROS,
            MAX_PRESENTATION_IMAGE_COST_MICROS,
        ),
        MAX_PRESENTATION_IMAGE_COST_MICROS,
    )


def _presentation_visual_model_usage(
    usages: list[dict[str, Any]],
    cost_limit: int,
    model: str,
) -> dict[str, Any]:
    returned = {
        str(usage.get("returned_model") or "") for usage in usages if usage.get("returned_model")
    }
    providers = {
        str(usage.get("upstream_provider") or "")
        for usage in usages
        if usage.get("upstream_provider")
    }
    return {
        "provider": "openrouter" if usages else "none",
        "config_version": os.environ.get(
            "MODEL_CONFIG_VERSION",
            "2026-08-structured-composer-v4",
        ),
        "requested_model": model,
        "returned_model": returned.pop() if len(returned) == 1 else ("mixed" if returned else ""),
        "upstream_provider": (
            providers.pop() if len(providers) == 1 else ("mixed" if providers else "")
        ),
        "prompt_tokens": sum(_non_negative_int(item.get("prompt_tokens")) for item in usages),
        "completion_tokens": sum(
            _non_negative_int(item.get("completion_tokens")) for item in usages
        ),
        "cost_micros": sum(_non_negative_int(item.get("cost_micros")) for item in usages),
        "cost_accounting": (
            "reported"
            if usages and all(item.get("cost_accounting") == "reported" for item in usages)
            else "reserved_upper_bound"
            if usages
            else "none"
        ),
        "latency_ms": sum(_non_negative_int(item.get("latency_ms")) for item in usages),
        "max_cost_micros": cost_limit,
        "round_count": len(usages),
    }


def _add_presentation_visual(
    slide: Any,
    visual: dict[str, Any],
) -> None:
    frame_left = int(Inches(6.63))
    frame_top = int(Inches(1.38))
    frame_width = int(Inches(6.08))
    frame_height = int(Inches(5.08))
    image_width = max(1, int(visual.get("width") or 1))
    image_height = max(1, int(visual.get("height") or 1))
    image_ratio = image_width / image_height
    frame_ratio = frame_width / frame_height
    if image_ratio >= frame_ratio:
        picture_width = frame_width
        picture_height = int(frame_width / image_ratio)
    else:
        picture_height = frame_height
        picture_width = int(frame_height * image_ratio)
    picture_left = frame_left + (frame_width - picture_width) // 2
    picture_top = frame_top + (frame_height - picture_height) // 2
    slide.shapes.add_picture(
        io.BytesIO(visual["data"]),
        picture_left,
        picture_top,
        width=picture_width,
        height=picture_height,
    )
    caption = slide.shapes.add_textbox(Inches(6.63), Inches(6.58), Inches(6.08), Inches(0.28))
    caption_frame = caption.text_frame
    caption_frame.clear()
    caption_frame.paragraphs[0].text = str(visual.get("source_locator") or "")[:100]
    caption_frame.paragraphs[0].font.size = Pt(10)
    caption_frame.paragraphs[0].font.color.rgb = PRESENTATION_MUTED_COLOR
    caption_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    _add_presentation_visual_notes(slide, visual)


def _add_presentation_visual_notes(slide: Any, visual: dict[str, Any]) -> None:
    if visual.get("kind") == "source":
        source = str(visual.get("source_locator") or "原文件")
    else:
        source = (
            f"AI-generated visual; model={visual.get('model', '')}; "
            f"prompt={visual.get('prompt', '')}"
        )
    try:
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = f"[Sources]\n- {source}"
    except Exception:
        return


def _presentation_visual_outline(visual: dict[str, Any]) -> dict[str, Any]:
    outline = {
        "kind": str(visual.get("kind") or ""),
        "source_locator": str(visual.get("source_locator") or ""),
    }
    if visual.get("kind") == "source":
        outline["source_page"] = int(visual.get("source_page") or 0)
    return outline


def _presentation_table(value: Any) -> dict[str, Any] | None:
    if value is None:
        return None
    if not isinstance(value, dict):
        raise ValueError("presentation_table_must_be_an_object")
    columns = value.get("columns")
    rows = value.get("rows")
    if not isinstance(columns, list) or not 2 <= len(columns) <= 6:
        raise ValueError("presentation_table_columns_must_be_between_2_and_6")
    normalized_columns = []
    for column in columns:
        if not isinstance(column, str) or not column.strip():
            raise ValueError("presentation_table_column_is_invalid")
        normalized_columns.append(column.strip()[:40])
    if not isinstance(rows, list) or not rows or len(rows) > 500:
        raise ValueError("presentation_table_rows_must_be_between_1_and_500")
    normalized_rows: list[dict[str, Any]] = []
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError("presentation_table_row_must_be_an_object")
        cells = row.get("cells")
        if not isinstance(cells, list) or len(cells) != len(normalized_columns):
            raise ValueError("presentation_table_row_cell_count_mismatch")
        normalized_cells = []
        for cell in cells:
            text = str(cell).strip()
            if not text:
                raise ValueError("presentation_table_cell_is_empty")
            if len(text) > 4000:
                raise ValueError("presentation_table_cell_is_too_long")
            normalized_cells.append(text)
        source_locator = str(row.get("source_locator") or "").strip()
        if not source_locator:
            raise ValueError("presentation_table_source_locator_is_empty")
        if len(source_locator) > 160:
            raise ValueError("presentation_table_source_locator_is_too_long")
        normalized_row: dict[str, Any] = {
            "cells": normalized_cells,
            "source_locator": source_locator,
        }
        entity_id = str(row.get("entity_id") or "").strip()
        if entity_id:
            normalized_row["entity_id"] = entity_id[:160]
        source_refs = row.get("source_refs")
        if isinstance(source_refs, list):
            normalized_row["source_refs"] = list(
                dict.fromkeys(str(value).strip() for value in source_refs if str(value).strip())
            )[:500]
        normalized_rows.append(normalized_row)
    return {
        "title": str(value.get("title") or "").strip()[:60],
        "columns": normalized_columns,
        "rows": normalized_rows,
    }


def _presentation_source_coverage(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        return {
            "coverage_ratio": 1.0,
            "truncated": False,
            "completed_rounds": 0,
            "round_count": 0,
        }
    ratio = value.get("coverage_ratio", 0)
    if isinstance(ratio, bool) or not isinstance(ratio, (int, float)):
        ratio = 0
    return {
        "coverage_ratio": max(0.0, min(1.0, float(ratio))),
        "truncated": value.get("truncated") is True,
        "completed_rounds": _non_negative_int(value.get("completed_rounds")),
        "round_count": _non_negative_int(value.get("round_count")),
        "selected_chunk_count": _non_negative_int(value.get("selected_chunk_count")),
        "total_chunk_count": _non_negative_int(value.get("total_chunk_count")),
    }


def _presentation_display_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Keep exactly one visual record for every logical source entity."""

    display: list[dict[str, Any]] = []
    for source_index, row in enumerate(rows):
        entity_id = str(row.get("entity_id") or f"record:{source_index + 1}")
        display.append(
            {
                "cells": list(row["cells"]),
                "source_locator": row["source_locator"],
                "entity_id": entity_id,
                "continuation": False,
            }
        )
    return display


def _presentation_row_units(row: dict[str, Any]) -> int:
    """Estimate the vertical space a logical table row needs."""

    cells = [str(value or "") for value in row.get("cells", [])]
    if not cells:
        return 1
    line_estimates = [
        math.ceil(len(value) / (20 if index == 0 else 32 if index == 1 else 68))
        for index, value in enumerate(cells)
    ]
    return max(1, min(PRESENTATION_TABLE_PAGE_CAPACITY, max(line_estimates)))


def _presentation_table_pages(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Paginate whole entities; never turn one logical record into table rows."""

    pages: list[dict[str, Any]] = []
    current: list[dict[str, Any]] = []
    current_units = 0

    def flush() -> None:
        nonlocal current, current_units
        if current:
            pages.append({"detail": False, "rows": current})
        current = []
        current_units = 0

    for row in rows:
        maximum_cell = max((len(str(value or "")) for value in row.get("cells", [])), default=0)
        units = _presentation_row_units(row)
        if maximum_cell >= PRESENTATION_DETAIL_CELL_THRESHOLD or units > 6:
            flush()
            pages.append({"detail": True, "rows": [row]})
            continue
        if current and (
            current_units + units > PRESENTATION_TABLE_PAGE_CAPACITY
            or len(current) >= PRESENTATION_ROWS_PER_SLIDE
        ):
            flush()
        current.append(row)
        current_units += units
    flush()
    return pages


def _split_visible_cell(value: Any) -> list[str]:
    remaining = re.sub(r"\s+", " ", str(value or "")).strip()
    if not remaining:
        return [""]
    result: list[str] = []
    while len(remaining) > PRESENTATION_DISPLAY_CELL_CHARS:
        window = remaining[:PRESENTATION_DISPLAY_CELL_CHARS]
        positions = [window.rfind(separator) for separator in ("；", ";", "，", ",", " ")]
        boundary = max(positions)
        cut = (
            boundary + 1
            if boundary >= PRESENTATION_DISPLAY_CELL_CHARS // 2
            else PRESENTATION_DISPLAY_CELL_CHARS
        )
        result.append(remaining[:cut].strip())
        remaining = remaining[cut:].strip()
    if remaining:
        result.append(remaining)
    return result or [""]


def _audience_source_locators(values: list[str]) -> list[str]:
    pages: list[str] = []
    for value in values:
        matches = re.findall(r"(?i)\bpage\s+(\d+)(?:\s*[-–]\s*(\d+))?", value)
        for start, end in matches:
            pages.append(f"第 {start}-{end} 页" if end else f"第 {start} 页")
    return list(dict.fromkeys(pages)) or ["附件原文"]


def _add_table_slide(slide: Any, columns: list[str], rows: list[dict[str, Any]]) -> None:
    row_units = [_presentation_row_units(row) for row in rows]
    header_height = 0.56
    minimum_body_height = min(4.59, 0.9 * len(rows) + 0.45)
    ideal_body_height = sum(0.78 + 0.28 * (units - 1) for units in row_units)
    body_height = min(4.59, max(minimum_body_height, ideal_body_height))
    table_height = header_height + body_height
    table_top = 1.42 + (5.15 - table_height) * 0.22
    shape = slide.shapes.add_table(
        len(rows) + 1,
        len(columns),
        Inches(0.45),
        Inches(table_top),
        Inches(12.43),
        Inches(table_height),
    )
    table = shape.table
    weights = []
    for column_index, column in enumerate(columns):
        maximum = max(
            [len(column), *[len(str(row["cells"][column_index])) for row in rows]],
            default=len(column),
        )
        weights.append(max(1.0, min(3.0, math.sqrt(maximum) / 3)))
    total_weight = sum(weights)
    for column_index, weight in enumerate(weights):
        table.columns[column_index].width = Inches(12.43 * weight / total_weight)
    table.rows[0].height = Inches(header_height)
    row_height_weights = [0.78 + 0.28 * (units - 1) for units in row_units]
    total_height_weight = max(1.0, sum(row_height_weights))
    for row_index, height_weight in enumerate(row_height_weights, start=1):
        table.rows[row_index].height = Inches(body_height * height_weight / total_height_weight)
    for column_index, column in enumerate(columns):
        cell = table.cell(0, column_index)
        cell.text = column
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.03)
        cell.margin_bottom = Inches(0.03)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRESENTATION_TITLE_COLOR
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(255, 255, 255)
            paragraph.font.size = Pt(17 if len(columns) <= 4 else 15)
            paragraph.line_spacing = 1.0
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row["cells"]):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (
                RGBColor(238, 243, 255) if row_index % 2 == 0 else RGBColor(255, 255, 255)
            )
            base_font_size = 16 if len(columns) <= 4 and len(rows) <= 4 else 15
            if row_units[row_index - 1] > 3:
                base_font_size = 14
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(base_font_size)
                paragraph.font.color.rgb = PRESENTATION_TEXT_COLOR
                paragraph.line_spacing = 1.0
    source_locators = _audience_source_locators(
        list(dict.fromkeys(row["source_locator"] for row in rows))
    )
    footer = slide.shapes.add_textbox(Inches(0.48), Inches(7.02), Inches(12.35), Inches(0.25))
    footer_frame = footer.text_frame
    footer_frame.clear()
    footer_frame.paragraphs[0].text = ("来源：" + "；".join(source_locators))[:160]
    footer_frame.paragraphs[0].font.size = Pt(10)
    footer_frame.paragraphs[0].font.color.rgb = PRESENTATION_MUTED_COLOR


def _add_course_detail_slide(
    slide: Any,
    columns: list[str],
    row: dict[str, Any],
) -> None:
    """Render an oversized logical entity with one unsplit value cell."""

    cells = [str(value or "") for value in row["cells"]]
    detail_index = max(range(len(cells)), key=lambda index: len(cells[index]))
    metadata = "    ".join(
        f"{columns[index]}：{value}" for index, value in enumerate(cells) if index != detail_index
    )
    metadata_box = slide.shapes.add_textbox(Inches(0.55), Inches(1.34), Inches(12.2), Inches(0.48))
    metadata_frame = metadata_box.text_frame
    metadata_frame.clear()
    metadata_frame.word_wrap = True
    metadata_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    metadata_frame.paragraphs[0].text = metadata
    metadata_frame.paragraphs[0].font.size = Pt(17)
    metadata_frame.paragraphs[0].font.bold = True
    metadata_frame.paragraphs[0].font.color.rgb = PRESENTATION_TITLE_COLOR

    shape = slide.shapes.add_table(
        1,
        1,
        Inches(0.55),
        Inches(1.94),
        Inches(12.2),
        Inches(4.62),
    )
    cell = shape.table.cell(0, 0)
    cell.text = f"{columns[detail_index]}：\n{cells[detail_index]}"
    cell.margin_left = Inches(0.16)
    cell.margin_right = Inches(0.16)
    cell.margin_top = Inches(0.12)
    cell.margin_bottom = Inches(0.12)
    cell.vertical_anchor = MSO_ANCHOR.TOP
    cell.fill.solid()
    cell.fill.fore_color.rgb = PRESENTATION_BACKGROUND
    detail_length = len(cells[detail_index])
    if detail_length <= 900:
        font_size = 16
    elif detail_length <= 1_300:
        font_size = 14
    elif detail_length <= 1_800:
        font_size = 12
    else:
        font_size = 10
    detail_frame = cell.text_frame
    detail_frame.clear()
    label_paragraph = detail_frame.paragraphs[0]
    label_paragraph.text = columns[detail_index]
    label_paragraph.font.size = Pt(17)
    label_paragraph.font.bold = True
    label_paragraph.font.color.rgb = PRESENTATION_TITLE_COLOR
    value_paragraph = detail_frame.add_paragraph()
    value_paragraph.text = cells[detail_index]
    value_paragraph.font.size = Pt(font_size)
    value_paragraph.font.color.rgb = PRESENTATION_TEXT_COLOR
    value_paragraph.line_spacing = 1.0

    source_locators = _audience_source_locators([row["source_locator"]])
    footer = slide.shapes.add_textbox(Inches(0.58), Inches(6.89), Inches(12.1), Inches(0.26))
    footer_frame = footer.text_frame
    footer_frame.clear()
    footer_frame.paragraphs[0].text = ("来源：" + "；".join(source_locators))[:160]
    footer_frame.paragraphs[0].font.size = Pt(10)
    footer_frame.paragraphs[0].font.color.rgb = PRESENTATION_MUTED_COLOR


def _presentation_quality_report(
    *,
    outline: list[dict[str, Any]],
    style: str,
    table: dict[str, Any] | None,
    task_contract: dict[str, Any],
    source_coverage: dict[str, Any],
    display_rows: list[dict[str, Any]],
    layout_violations: list[dict[str, Any]],
    visual_report: dict[str, Any],
    filename: Any,
    brief: str,
    requested_slide_count: int,
) -> dict[str, Any]:
    violations: list[dict[str, Any]] = list(layout_violations)
    violations.extend(
        presentation_exhaustive_scope_violations(
            {
                "title": outline[0].get("title") if outline else "",
                "filename": filename,
                "table": table,
            },
            task_contract,
        )
    )
    violations.extend(
        presentation_visible_language_violations(
            {
                "title": outline[0].get("title") if outline else "",
                "table": table,
                "brief": brief,
            },
            task_contract,
        )
    )
    violations.extend(
        presentation_audience_content_violations(
            {
                "title": outline[0].get("title") if outline else "",
                "brief": brief,
                "slide_count": requested_slide_count,
                "table": table,
            },
            task_contract,
        )
    )
    presentation_capabilities = task_contract.get("presentation_capabilities")
    table_required = (
        isinstance(presentation_capabilities, list) and "table" in presentation_capabilities
    )
    if (table_required or "表格" in style) and not table:
        violations.append(
            {"code": "requested_table_missing", "message": "用户要求表格，但未提供结构化表格数据"}
        )
    signatures = []
    for slide in outline[1:-1]:
        signature = json.dumps(slide, ensure_ascii=False, sort_keys=True)
        if signature in signatures:
            violations.append({"code": "duplicate_slide_content", "message": "存在重复内容页"})
            break
        signatures.append(signature)
        if len(str(slide.get("title") or "")) > 36:
            violations.append({"code": "slide_title_too_long", "message": "页面标题过长"})
            break
        visual = slide.get("visual")
        if isinstance(visual, Mapping) and visual.get("kind") == "source":
            cited_pages = _presentation_topic_source_pages(slide)
            source_page = int(visual.get("source_page") or 0)
            if not cited_pages or source_page not in cited_pages:
                violations.append(
                    {
                        "code": "source_visual_topic_mismatch",
                        "message": "原文图片与当前主题引用的页码不一致",
                        "slide": int(slide.get("page") or 0),
                        "source_page": source_page,
                    }
                )
                break
    if task_contract.get("exhaustive") is True:
        if source_coverage["truncated"] or source_coverage["coverage_ratio"] < 1:
            violations.append({"code": "source_coverage_incomplete", "message": "来源尚未完整处理"})
    requested_fields = task_contract.get("requested_fields")
    if isinstance(requested_fields, list) and requested_fields and not table:
        violations.append(
            {
                "code": "structured_table_missing",
                "message": "结构化字段任务必须提供表格数据，不能只生成通用内容页",
            }
        )
    if (
        isinstance(presentation_capabilities, list)
        and "visual" in presentation_capabilities
        and int(visual_report.get("placed_visual_count") or 0) < 1
    ):
        violations.append(
            {
                "code": "requested_visual_missing",
                "message": "用户明确要求插图，但视觉子能力没有生成或提取到可用图片",
            }
        )
    if table and isinstance(requested_fields, list):
        header_text = " ".join(table["columns"]).casefold()
        aliases = {
            "code": ("代码", "编号", "code"),
            "name": ("名称", "课程", "name"),
            "time": ("时间", "日期", "星期", "时段", "time", "date", "schedule"),
            "venue": ("地点", "教室", "场地", "venue", "location"),
        }
        for field in requested_fields:
            candidates = aliases.get(str(field), (str(field),))
            if not any(candidate in header_text for candidate in candidates):
                violations.append(
                    {
                        "code": "requested_field_missing",
                        "field": field,
                        "message": f"缺少字段：{field}",
                    }
                )
    if table:
        entity_rows: dict[str, list[int]] = {}
        delimiter_noise_rows: list[int] = []
        repeated_delimiter = re.compile(r"(?:[；;]\s*){2,}")
        punctuation_only = re.compile(r"^[\s；;,，、|/\\:：.。·•↳\-–—]+$")
        for index, row in enumerate(table["rows"], start=1):
            entity_id = str(row.get("entity_id") or "").strip()
            if entity_id:
                entity_rows.setdefault(entity_id, []).append(index)
            if any(
                repeated_delimiter.search(str(value or ""))
                or punctuation_only.fullmatch(str(value or "").strip())
                for value in row.get("cells", [])
                if str(value or "").strip()
            ):
                delimiter_noise_rows.append(index)
        duplicates = {
            entity_id: positions
            for entity_id, positions in entity_rows.items()
            if len(positions) > 1
        }
        if duplicates:
            violations.append(
                {
                    "code": "duplicate_logical_entities",
                    "message": "同一来源实体在最终表格中出现了多条逻辑记录",
                    "entity_ids": list(duplicates)[:20],
                }
            )
        if delimiter_noise_rows:
            violations.append(
                {
                    "code": "presentation_cell_delimiter_noise",
                    "message": "最终表格包含重复分隔符或纯标点碎片",
                    "affected_rows": delimiter_noise_rows[:20],
                }
            )
        contextless_rows = [
            index
            for index, row in enumerate(display_rows, start=1)
            if row.get("continuation") is True
            and any(
                not str(value or "").strip() or str(value).strip() == "↳" for value in row["cells"]
            )
        ]
        if contextless_rows:
            violations.append(
                {
                    "code": "continuation_row_context_missing",
                    "message": "跨行展示必须重复关键字段，不能使用空白或箭头代替上下文",
                    "affected_rows": contextless_rows[:20],
                }
            )
        split_rows = [
            index
            for index, row in enumerate(display_rows, start=1)
            if row.get("continuation") is True
        ]
        if split_rows:
            violations.append(
                {
                    "code": "presentation_logical_record_split",
                    "message": "同一逻辑记录不得被渲染为多个表格行",
                    "affected_rows": split_rows[:20],
                }
            )
        violations.extend(
            presentation_table_field_semantic_violations(
                table.get("columns", []),
                table.get("rows", []),
                requested_fields,
            )
        )
        violations.extend(
            presentation_table_language_violations(
                table.get("columns", []),
                table.get("rows", []),
                task_contract,
            )
        )
        oversized_detail_rows = [
            index
            for index, row in enumerate(table.get("rows", []), start=1)
            if max((len(str(value or "")) for value in row.get("cells", [])), default=0)
            > PRESENTATION_DETAIL_MAX_CHARS
        ]
        if oversized_detail_rows:
            violations.append(
                {
                    "code": "presentation_detail_cell_capacity_exceeded",
                    "message": "单个逻辑记录超过一页单元格的可读容量",
                    "affected_rows": oversized_detail_rows[:20],
                }
            )
    visible_text = " ".join(
        [
            str(outline[0].get("title") or "") if outline else "",
            *[str(column) for column in (table.get("columns", []) if table else [])],
            *[
                str(cell)
                for row in (table.get("rows", []) if table else [])
                for cell in row.get("cells", [])
            ],
        ]
    )
    if re.search(
        r"(?:详?见(?:下列(?:来源|内容)?|下方|后文|原文|来源|课程表)|"
        r"see\s+(?:below|source|original))",
        visible_text,
        re.I,
    ):
        violations.append(
            {
                "code": "presentation_reference_guidance_visible",
                "message": "最终演示文稿不得用来源指引或“见下方”占位替代结构化字段值",
            }
        )
    if re.search(
        r"(?:\battachment\s*:\s*\d+|\bround\s*:\s*\d+|\bsegment\s*:\s*\d+|"
        r"source\s*ir|source_locator|本轮观测|当前轮次|harness)",
        visible_text,
        re.I,
    ):
        violations.append(
            {"code": "internal_workflow_text_visible", "message": "可见内容泄露了内部处理术语"}
        )
    content_capacity = (MAX_PRESENTATION_SLIDES - 2) * PRESENTATION_ROWS_PER_SLIDE
    if table and len(display_rows) > content_capacity:
        violations.append(
            {
                "code": "presentation_capacity_exceeded",
                "message": "完整内容超过当前可读字号下的确定性分页容量",
                "display_row_count": len(display_rows),
                "capacity": content_capacity,
            }
        )
    return {
        "policy_version": "pptx-quality-v2",
        "passed": not violations,
        "hard_gate": True,
        "violations": violations,
        "slide_count": len(outline),
        "table_row_count": len(table["rows"]) if table else 0,
        "display_row_count": len(display_rows),
        "source_coverage": source_coverage,
        "visual_report": visual_report,
    }


def _presentation_layout_violations(presentation: Presentation) -> list[dict[str, Any]]:
    violations: list[dict[str, Any]] = []
    width = int(presentation.slide_width)
    height = int(presentation.slide_height)
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            left = int(getattr(shape, "left", 0))
            top = int(getattr(shape, "top", 0))
            shape_width = int(getattr(shape, "width", 0))
            shape_height = int(getattr(shape, "height", 0))
            if left < 0 or top < 0 or left + shape_width > width or top + shape_height > height:
                violations.append(
                    {
                        "code": "shape_outside_slide_canvas",
                        "message": "页面元素超出演示文稿画布",
                        "slide": slide_index,
                    }
                )
                break
            if _presentation_body_text_overflow_risk(shape):
                violations.append(
                    {
                        "code": "body_text_overflow_risk",
                        "message": "正文按实际文本宽度估算将超出占位区域",
                        "slide": slide_index,
                    }
                )
                break
    return violations


def _presentation_body_text_overflow_risk(shape: Any) -> bool:
    """Conservatively detect wrapped body text that cannot fit its box.

    OOXML records the placeholder rectangle but not the final rendered text
    extent. Estimate wide-character line usage so the hard quality gate sees
    the same overflow class as an independent slide renderer.
    """

    if (
        not getattr(shape, "is_placeholder", False)
        or shape.placeholder_format.type != PP_PLACEHOLDER.OBJECT
        or not getattr(shape, "has_text_frame", False)
    ):
        return False
    frame = shape.text_frame
    usable_width = (
        float(shape.width - frame.margin_left - frame.margin_right) / 12_700.0 - 36.0
    )
    usable_height = float(shape.height - frame.margin_top - frame.margin_bottom) / 12_700.0
    if usable_width <= 0 or usable_height <= 0:
        return True
    required_height = 0.0
    for paragraph in frame.paragraphs:
        if not paragraph.text.strip():
            continue
        font_size = paragraph.font.size.pt if paragraph.font.size is not None else 18.0
        display_units = sum(
            1.0 if unicodedata.east_asian_width(character) in {"W", "F"} else 0.55
            for character in paragraph.text
        )
        characters_per_line = max(8.0, usable_width / max(font_size, 1.0))
        line_count = max(1, math.ceil(display_units / characters_per_line))
        line_spacing = paragraph.line_spacing
        spacing_multiple = float(line_spacing) if isinstance(line_spacing, float) else 1.0
        spacing_after = paragraph.space_after.pt if paragraph.space_after is not None else 0.0
        required_height += (line_count + 0.6) * font_size * spacing_multiple + spacing_after
    return required_height > usable_height


def _style_slide(slide: Any, title_color: RGBColor) -> None:
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = PRESENTATION_BACKGROUND  # type: ignore[no-untyped-call]
    if slide.shapes.title is not None:
        slide.shapes.title.left = Inches(0.62)
        slide.shapes.title.top = Inches(0.32)
        slide.shapes.title.width = Inches(12.05)
        slide.shapes.title.height = Inches(0.78)
        slide.shapes.title.text_frame.word_wrap = True
        title_length = len(slide.shapes.title.text)
        title_size = 28 if title_length <= 24 else 26 if title_length <= 36 else 24
        for paragraph in slide.shapes.title.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.color.rgb = title_color
            paragraph.font.size = Pt(title_size)
            paragraph.alignment = PP_ALIGN.LEFT


def _style_title_slide(slide: Any, title: str, subtitle: Any) -> None:
    """Keep the cover minimal while grouping its title and audience clearly."""

    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = PRESENTATION_BACKGROUND
    title_shape = slide.shapes.title
    title_shape.left = Inches(0.88)
    title_shape.top = Inches(2.18)
    title_shape.width = Inches(11.55)
    title_shape.height = Inches(1.38)
    title_shape.text_frame.word_wrap = True
    title_length = len(title)
    title_size = (
        40 if title_length <= 22 else 35 if title_length <= 40 else 30 if title_length <= 68 else 26
    )
    for paragraph in title_shape.text_frame.paragraphs:
        paragraph.font.bold = True
        paragraph.font.color.rgb = PRESENTATION_TITLE_COLOR
        paragraph.font.size = Pt(title_size)
        paragraph.alignment = PP_ALIGN.LEFT
    subtitle.left = Inches(0.91)
    subtitle.top = Inches(3.72)
    subtitle.width = Inches(11.1)
    subtitle.height = Inches(0.58)
    subtitle.text_frame.word_wrap = True
    for paragraph in subtitle.text_frame.paragraphs:
        paragraph.font.size = Pt(18)
        paragraph.font.color.rgb = PRESENTATION_MUTED_COLOR
        paragraph.alignment = PP_ALIGN.LEFT


def _file(name: str, media_type: str, data: bytes) -> dict[str, str]:
    return {
        "name": name,
        "media_type": media_type,
        "data_base64": base64.b64encode(data).decode("ascii"),
    }


def _is_number(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    try:
        number = float(value)
    except (TypeError, ValueError):
        return False
    return math.isfinite(number)


def _stable_value(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)


def _stable_row(row: list[Any]) -> str:
    return json.dumps(row, ensure_ascii=False, sort_keys=True, default=str)


def _dynamic_model(model: str) -> bool:
    normalized = model.strip().lower()
    return (
        normalized in ("openrouter/free", "openrouter/auto")
        or normalized.startswith("~")
        or normalized.endswith((":free", ":nitro", ":floor"))
        or normalized.endswith("-latest")
    )


def _bounded_positive_int(name: str, fallback: int, maximum: int) -> int:
    raw = os.environ.get(name, str(fallback)).strip()
    try:
        value = int(raw)
    except ValueError as exc:
        raise ValueError(f"{name}_must_be_a_positive_integer") from exc
    if value <= 0 or value > maximum:
        raise ValueError(f"{name}_is_out_of_range")
    return value


def _positive_decimal(name: str, fallback: str) -> Decimal:
    try:
        value = Decimal(os.environ.get(name, fallback).strip())
    except InvalidOperation as exc:
        raise ValueError(f"{name}_must_be_positive") from exc
    if not value.is_finite() or value <= 0:
        raise ValueError(f"{name}_must_be_positive")
    return value


def _decimal_ceil(value: Decimal) -> int:
    return int(value.to_integral_value(rounding=ROUND_CEILING))


def _env_bool(name: str, fallback: bool) -> bool:
    raw = os.environ.get(name)
    if raw is None or not raw.strip():
        return fallback
    normalized = raw.strip().lower()
    if normalized in ("1", "true", "yes", "on"):
        return True
    if normalized in ("0", "false", "no", "off"):
        return False
    raise ValueError(f"{name}_must_be_true_or_false")


def _non_negative_int(value: Any) -> int:
    if isinstance(value, int) and not isinstance(value, bool) and value >= 0:
        return value
    return 0


def _reported_cost_micros(value: Any) -> int | None:
    try:
        cost = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    if not cost.is_finite() or cost < 0:
        return None
    return _decimal_ceil(cost * Decimal(1_000_000))


def main() -> None:
    try:
        request = json.load(sys.stdin)
        operation = request.get("operation")
        payload = request.get("input")
        if not isinstance(operation, str) or not isinstance(payload, dict):
            raise ValueError("invalid_request")
        result = execute(operation, payload)
    except ModelBackedOperationError as exc:
        result = {
            "output": {
                "model_usage": exc.model_usage,
                "operation_error": {"code": exc.code},
            },
            "files": [],
        }
    except Exception as exc:
        print(
            json.dumps(_tool_failure_for_exception(exc), ensure_ascii=False),
            file=sys.stderr,
        )
        raise SystemExit(2) from exc
    print(json.dumps(result, ensure_ascii=False, separators=(",", ":")))


def _tool_failure_for_exception(exc: Exception) -> dict[str, Any]:
    code = str(exc).strip() or "office_worker_failed"
    if code == "invalid_output_filename":
        return {
            "contract_version": "tool-failure-v1",
            "code": code,
            "category": "argument_validation",
            "phase": "pre_execution",
            "message": "输出文件名必须是文件名，不能包含目录分隔符。",
            "retry_same_input": False,
            "repairable": True,
            "side_effect_state": "none",
            "field_paths": ["/filename"],
            "allowed_repairs": [
                "remove_optional_filename",
                "filename.safe_basename",
            ],
            "safe_details": {"constraint": "output_basename"},
        }
    return {
        "contract_version": "tool-failure-v1",
        "code": code if len(code) <= 128 else "office_worker_failed",
        "category": "argument_validation" if isinstance(exc, ValueError) else "execution",
        "phase": "pre_execution" if isinstance(exc, ValueError) else "execution",
        "message": "Office 工具参数无效。"
        if isinstance(exc, ValueError)
        else "Office 工具执行失败。",
        "retry_same_input": False,
        "repairable": False,
        "side_effect_state": "none" if isinstance(exc, ValueError) else "unknown",
    }


if __name__ == "__main__":
    main()
