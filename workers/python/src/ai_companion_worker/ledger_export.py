"""Generate the monthly ledger workbook used by the Go ledger exporter."""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


def _parse_time(value: str, timezone: str) -> datetime:
    instant = datetime.fromisoformat(value.replace("Z", "+00:00"))
    return instant.astimezone(ZoneInfo(timezone)).replace(tzinfo=None)


def export_workbook(payload: dict[str, Any], output_path: Path) -> None:
    month = str(payload["month"])
    currency = str(payload["currency"])
    timezone = str(payload["timezone"])
    entries = list(payload.get("entries", []))
    generated_at = _parse_time(str(payload["generated_at"]), timezone)

    workbook = Workbook()
    detail = workbook.active
    detail.title = "账单明细"
    summary = workbook.create_sheet("月度汇总")
    detail.sheet_view.showGridLines = False
    summary.sheet_view.showGridLines = False

    orange = "C65D2E"
    green = "6B7C5A"
    cream = "FFF3E8"
    white = "FFFFFF"
    brown = "6B4A36"
    thin = Side(style="thin", color="E7DDD4")

    detail.merge_cells("A1:H1")
    detail["A1"] = f"伴AI 账单明细 · {month}"
    detail["A1"].fill = PatternFill("solid", fgColor=orange)
    detail["A1"].font = Font(bold=True, color=white, size=16)
    detail["A1"].alignment = Alignment(vertical="center")
    detail.row_dimensions[1].height = 32
    detail.merge_cells("A2:H2")
    detail["A2"] = f"生成时间：{generated_at:%Y-%m-%d %H:%M} · 报表时区：{timezone}"
    detail["A2"].fill = PatternFill("solid", fgColor=cream)
    detail["A2"].font = Font(color=brown, size=10)
    detail.row_dimensions[2].height = 22

    headers = ["发生时间", "收支方向", "分类", "商户", "备注", "币种", "金额", "记录 ID"]
    for column, value in enumerate(headers, 1):
        cell = detail.cell(row=5, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=green)
        cell.font = Font(bold=True, color=white)
    detail.row_dimensions[5].height = 24

    for row_number, entry in enumerate(entries, 6):
        values = [
            _parse_time(str(entry["occurred_at"]), timezone),
            entry.get("direction", ""),
            entry.get("category", ""),
            entry.get("merchant", ""),
            entry.get("note", ""),
            entry.get("currency", ""),
            int(entry.get("amount_minor", 0)) / 100,
            entry.get("id", ""),
        ]
        for column, value in enumerate(values, 1):
            cell = detail.cell(row=row_number, column=column, value=value)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        detail.cell(row=row_number, column=1).number_format = "yyyy-mm-dd hh:mm"
        detail.cell(row=row_number, column=7).number_format = "#,##0.00"

    if entries:
        table = Table(displayName="LedgerEntriesTable", ref=f"A5:H{5 + len(entries)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
        detail.add_table(table)
    detail.freeze_panes = "A6"
    widths = {"A": 20, "B": 12, "C": 14, "D": 18, "E": 22, "F": 10, "G": 14, "H": 38}
    for column_letter, column_width in widths.items():
        detail.column_dimensions[column_letter].width = column_width

    summary.merge_cells("A1:F1")
    summary["A1"] = f"伴AI 月度账本 · {month}"
    summary["A1"].fill = PatternFill("solid", fgColor=orange)
    summary["A1"].font = Font(bold=True, color=white, size=16)
    summary.row_dimensions[1].height = 32
    summary["A2"] = "统计币种"
    summary["B2"] = currency
    labels = ["收入", "支出", "结余", "记录数"]
    for row, label in enumerate(labels, 4):
        summary.cell(row=row, column=1, value=label)
        summary.cell(row=row, column=1).fill = PatternFill("solid", fgColor=cream)
        summary.cell(row=row, column=1).font = Font(bold=True, color=brown)

    data_end = max(105, 5 + len(entries))
    summary["B4"] = f'=SUMIFS(\'账单明细\'!$G$6:$G${data_end},\'账单明细\'!$B$6:$B${data_end},"income",\'账单明细\'!$F$6:$F${data_end},$B$2)'
    summary["B5"] = f'=SUMIFS(\'账单明细\'!$G$6:$G${data_end},\'账单明细\'!$B$6:$B${data_end},"expense",\'账单明细\'!$F$6:$F${data_end},$B$2)'
    summary["B6"] = "=B4-B5"
    summary["B7"] = f'=COUNTIFS(\'账单明细\'!$F$6:$F${data_end},$B$2)'
    for row in range(4, 7):
        summary.cell(row=row, column=2).number_format = "#,##0.00"

    categories = sorted(
        {
            str(entry.get("category", ""))
            for entry in entries
            if entry.get("direction") == "expense" and entry.get("currency") == currency
        }
    )
    summary["A10"] = "支出分类"
    summary["B10"] = "金额"
    for cell in summary[10][:2]:
        cell.fill = PatternFill("solid", fgColor=green)
        cell.font = Font(bold=True, color=white)
    for index, category in enumerate(categories, 11):
        summary.cell(row=index, column=1, value=category)
        summary.cell(
            row=index,
            column=2,
            value=f'=SUMIFS(\'账单明细\'!$G$6:$G${data_end},\'账单明细\'!$B$6:$B${data_end},"expense",\'账单明细\'!$C$6:$C${data_end},A{index},\'账单明细\'!$F$6:$F${data_end},$B$2)',
        ).number_format = "#,##0.00"
    if categories:
        chart = BarChart()
        chart.title = "支出分类分布"
        chart.legend = None
        chart.add_data(Reference(summary, min_col=2, min_row=10, max_row=10 + len(categories)), titles_from_data=True)
        chart.set_categories(Reference(summary, min_col=1, min_row=11, max_row=10 + len(categories)))
        chart.height = 8
        chart.width = 14
        summary.add_chart(chart, "D3")
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 16

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("output path is required")
    export_workbook(json.load(sys.stdin), Path(sys.argv[1]))


if __name__ == "__main__":
    main()
