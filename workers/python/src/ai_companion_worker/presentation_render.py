"""Native Office export and independent render verification for design v2."""
# Office and PyMuPDF expose untyped constructors; keep the design boundary typed.
# mypy: disable-error-code="no-untyped-call"
from __future__ import annotations

import base64
import io
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import pymupdf
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from ai_companion_worker.presentation_design import theme_for
from ai_companion_worker.presentation_images import normalize_visual_image


def normalize_assets(deck: dict[str, Any]) -> None:
    for asset in deck["assets"].values():
        try:
            raw = base64.b64decode(asset["data_base64"], validate=True)
        except (ValueError, TypeError) as error:
            raise ValueError("presentation_asset_base64_invalid") from error
        normalized, media_type, width, height = normalize_visual_image(raw)
        asset.update(data_base64=base64.b64encode(normalized).decode(), media_type=media_type, width=width, height=height)


def _xml(parent: Any, tag: str, **attrs: str) -> Any:
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(key, value)
    parent.append(element)
    return element


def _font(paragraph: Any, name: str, size: float, color: str, bold: bool = False) -> None:
    paragraph.font.name = name
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    properties = paragraph._p.get_or_add_pPr().get_or_add_defRPr()
    for tag in ("a:ea", "a:cs"):
        existing = properties.find(tag, properties.nsmap)
        if existing is not None:
            properties.remove(existing)
        _xml(properties, tag, typeface=name)
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)


def _text(slide: Any, node: dict[str, Any]) -> Any:
    shape = slide.shapes.add_textbox(Pt(node["x"]), Pt(node["y"]), Pt(node["w"]), Pt(node["h"]))
    shape.name = node["id"]
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(node["lines"]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        _font(paragraph, node["font"], node["font_size"], node["color"], node.get("bold", False))
        paragraph.line_spacing = Pt(node["line_height"])
    return shape


def _chart(slide: Any, node: dict[str, Any], theme: dict[str, str]) -> Any:
    spec = node["chart"]
    data = CategoryChartData()
    data.categories = spec["labels"]
    data.add_series(spec["series"], spec["values"])
    kind = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS}[node["chart_type"]]
    shape = slide.shapes.add_chart(kind, Pt(node["x"]), Pt(node["y"]), Pt(node["w"]), Pt(node["h"]), data)
    shape.name = node["id"]
    chart = shape.chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name = node["font"]
    chart.font.size = Pt(12)
    chart.font.color.rgb = RGBColor.from_string(theme["ink"])
    for series in chart.series:
        if node["chart_type"] != "line":
            series.invert_if_negative = False
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(node["color"])
        series.format.line.color.rgb = RGBColor.from_string(node["color"])
    chart.value_axis.has_major_gridlines = False
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    if node["chart_type"] == "bar":
        chart.category_axis.reverse_order = True
    if all(v >= 0 for v in spec["values"]):
        chart.value_axis.minimum_scale = 0
    chart.plots[0].has_data_labels = True
    labels = chart.plots[0].data_labels
    labels.font.name = node["font"]
    labels.font.size = Pt(12)
    labels.font.color.rgb = RGBColor.from_string(theme["ink"])
    suffix = '"' + spec.get("unit", "").replace('"', '') + '"'
    # WPS retains the decimal separator for integers with Excel's 0.## format.
    precision = max(len(f"{value:.2f}".rstrip("0").split(".")[1]) for value in spec["values"])
    number = "0" + ("." + "0" * precision if precision else "")
    labels.number_format = f'{number}{suffix};-{number}{suffix};0{suffix}'
    if node["chart_type"] != "line":
        labels.position = XL_LABEL_POSITION.OUTSIDE_END
    return shape


def _table(slide: Any, node: dict[str, Any]) -> Any:
    spec = node["table"]
    rows = [spec["columns"], *spec["rows"]]
    count = len(spec["columns"])
    size, heights = node["font_size"], node["row_heights"]
    if not node["fits"]:
        raise ValueError("presentation_table_capacity_exceeded")
    shape = slide.shapes.add_table(len(rows), count, Pt(node["x"]), Pt(node["y"]), Pt(node["w"]), Pt(sum(heights)))
    shape.name = node["id"]
    for ri, row in enumerate(rows):
        shape.table.rows[ri].height = Pt(heights[ri])
        for ci, value in enumerate(row):
            cell = shape.table.cell(ri, ci)
            cell.margin_left = cell.margin_right = Pt(9)
            cell.margin_top = cell.margin_bottom = Pt(7)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(node["accent"] if ri == 0 else node["fill"])
            cell.text = "\n".join(node["cell_lines"][ri][ci])
            for paragraph in cell.text_frame.paragraphs:
                _font(paragraph, node["font"], size, node["header_color"] if ri == 0 else node["color"], ri == 0)
                paragraph.line_spacing = Pt(size * 1.3)
    return shape


def add_motion(slide: Any, scene: dict[str, Any], groups: dict[int, list[int]]) -> None:
    if scene["animation"] == "none":
        return
    root = slide._element
    transition = _xml(root, "p:transition", spd="med", advClick="1")
    _xml(transition, "p:push", dir="l") if scene["animation"] == "showcase" else _xml(transition, "p:fade")
    if scene["animation"] == "subtle" or not groups:
        return
    timing = _xml(root, "p:timing")
    timing_list = _xml(timing, "p:tnLst")
    main = _xml(_xml(timing_list, "p:par"), "p:cTn", id="1", dur="indefinite", restart="never", nodeType="tmRoot")
    sequence = _xml(_xml(main, "p:childTnLst"), "p:seq", concurrent="1", nextAc="seek")
    main_sequence = _xml(sequence, "p:cTn", id="2", dur="indefinite", nodeType="mainSeq")
    children = _xml(main_sequence, "p:childTnLst")
    builds = _xml(timing, "p:bldLst")
    timing_id = 3
    automatic_children = None
    if scene["advance"] == "auto":
        automatic = _xml(_xml(children, "p:par"), "p:cTn", id=str(timing_id), fill="hold")
        timing_id += 1
        _xml(_xml(automatic, "p:stCondLst"), "p:cond", delay="0")
        automatic_children = _xml(automatic, "p:childTnLst")
    for group_index, (_, targets) in enumerate(sorted(groups.items()), 1):
        if automatic_children is not None:
            phase_parent = automatic_children
        else:
            click = _xml(_xml(children, "p:par"), "p:cTn", id=str(timing_id), fill="hold")
            timing_id += 1
            _xml(_xml(click, "p:stCondLst"), "p:cond", delay="indefinite")
            phase_parent = _xml(click, "p:childTnLst")
        # Office expects separate click, phase and effect containers. Omitting
        # the phase or build records makes WPS discard otherwise valid effects.
        group = _xml(_xml(phase_parent, "p:par"), "p:cTn", id=str(timing_id), fill="hold")
        timing_id += 1
        _xml(_xml(group, "p:stCondLst"), "p:cond", delay="0" if scene["advance"] == "click" else str(group_index * max(700, scene["duration_ms"] + 250)))
        group_children = _xml(group, "p:childTnLst")
        for target_index, target in enumerate(targets):
            trigger = "withEffect" if target_index else "clickEffect" if scene["advance"] == "click" else "afterEffect"
            effect = _xml(_xml(group_children, "p:par"), "p:cTn", id=str(timing_id), presetID="10", presetClass="entr", presetSubtype="0", fill="hold", grpId="0", nodeType=trigger)
            _xml(builds, "p:bldP", spid=str(target), grpId="0")
            timing_id += 1
            _xml(_xml(effect, "p:stCondLst"), "p:cond", delay="0")
            behavior_list = _xml(effect, "p:childTnLst")
            visibility = _xml(behavior_list, "p:set")
            behavior = _xml(visibility, "p:cBhvr")
            visibility_time = _xml(behavior, "p:cTn", id=str(timing_id), dur="1", fill="hold")
            timing_id += 1
            _xml(_xml(visibility_time, "p:stCondLst"), "p:cond", delay="0")
            _xml(_xml(behavior, "p:tgtEl"), "p:spTgt", spid=str(target))
            _xml(_xml(behavior, "p:attrNameLst"), "p:attrName").text = "style.visibility"
            _xml(_xml(visibility, "p:to"), "p:strVal", val="visible")
            fade = _xml(behavior_list, "p:animEffect", transition="in", filter="fade")
            fade_behavior = _xml(fade, "p:cBhvr")
            _xml(fade_behavior, "p:cTn", id=str(timing_id), dur=str(scene["duration_ms"]))
            timing_id += 1
            _xml(_xml(fade_behavior, "p:tgtEl"), "p:spTgt", spid=str(target))
    for tag, event in (("p:prevCondLst", "onPrev"), ("p:nextCondLst", "onNext")):
        cond = _xml(_xml(sequence, tag), "p:cond", evt=event, delay="0")
        _xml(_xml(cond, "p:tgtEl"), "p:sldTgt")


def export_pptx(deck: dict[str, Any], scenes: list[dict[str, Any]]) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Pt(960)
    presentation.slide_height = Pt(540)
    presentation.core_properties.title = deck["title"]
    presentation.core_properties.subject = deck["audience"]
    presentation.core_properties.author = "伴AI"
    theme = theme_for(deck)
    for slide_spec, scene in zip(deck["slides"], scenes, strict=True):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(scene["background"])
        groups: dict[int, list[int]] = {}
        for node in scene["nodes"]:
            kind = node["type"]
            if kind == "text":
                shape = _text(slide, node)
            elif kind == "shape":
                shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if node["kind"] == "ellipse" else MSO_SHAPE.ROUNDED_RECTANGLE if node["kind"] == "roundrect" else MSO_SHAPE.RECTANGLE, Pt(node["x"]), Pt(node["y"]), Pt(node["w"]), Pt(node["h"]))
                shape.name = node["id"]
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(node["fill"])
                shape.line.fill.background()
                if node["kind"] == "roundrect":
                    shape.adjustments[0] = min(.5, node.get("radius", 12) / max(1, min(node["w"], node["h"])))
                    style = shape._element.find(qn('p:style'))
                    if style is not None:
                        shape._element.remove(style)
            elif kind == "connector":
                x, y, w, h = (node[k] for k in ("x", "y", "w", "h"))
                start, end = (x, y), (x+w, y+h)
                if node.get("direction") == "up":
                    start, end = end, start
                shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(start[0]), Pt(start[1]), Pt(end[0]), Pt(end[1]))
                style = shape._element.find(qn('p:style'))
                if style is not None:
                    shape._element.remove(style)
                shape.name = node["id"]
                shape.line.color.rgb = RGBColor.from_string(node["color"])
                shape.line.width = Pt(1.3)
                _xml(shape.line._get_or_add_ln(), "a:tailEnd", type="triangle", w="sm", len="sm")
            elif kind == "image":
                asset = deck["assets"][node["asset_id"]]
                ratio = asset["width"] / asset["height"]
                x, y, width, height = (node[k] for k in ("x", "y", "w", "h"))
                if node["fit"] == "contain":
                    if ratio > width / height:
                        y += (height - width / ratio) / 2
                        height = width / ratio
                    else:
                        x += (width - height * ratio) / 2
                        width = height * ratio
                shape = slide.shapes.add_picture(io.BytesIO(base64.b64decode(asset["data_base64"])), Pt(x), Pt(y), Pt(width), Pt(height))
                shape.name = node["id"]
                if node["fit"] == "cover":
                    if ratio > width / height:
                        shape.crop_left = shape.crop_right = (1 - width / height / ratio) / 2
                    else:
                        shape.crop_top = shape.crop_bottom = (1 - ratio / (width / height)) / 2
            elif kind == "chart":
                shape = _chart(slide, node, theme)
            elif kind == "table":
                shape = _table(slide, node)
            else:
                raise ValueError("presentation_node_type_invalid")
            if node.get("group", 0) > 0:
                groups.setdefault(node["group"], []).append(shape.shape_id)
        notes = [slide_spec.get("notes", ""), slide_spec.get("source", "")]
        notes.extend(i["source"] for i in slide_spec["items"] if i.get("source"))
        if slide_spec.get("image_id"):
            notes.append(deck["assets"][slide_spec["image_id"]].get("source", ""))
        slide.notes_slide.notes_text_frame.text = "\n".join(dict.fromkeys(n for n in notes if n))
        add_motion(slide, scene, groups)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def render_pptx(data: bytes, page_count: int) -> tuple[dict[str, Any], bytes | None, list[str]]:
    """Use the actual exported deck, not the authoring scene, for independent QA."""
    executable = os.environ.get("PRESENTATION_SOFFICE_PATH") or shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        return {"status": "unavailable", "message": "服务端未安装 LibreOffice，尚未进行导出渲染验证", "passed": False}, None, []
    with tempfile.TemporaryDirectory(prefix="presentation-render-") as directory:
        root = Path(directory)
        (root / "deck.pptx").write_bytes(data)
        environment = os.environ.copy()
        font_dir = Path(environment.get("PRESENTATION_FONT_DIR", ".data/presentation-fonts"))
        if font_dir.is_dir():
            # Isolate font configuration as well as the Office profile. This also
            # supports headless macOS builds which cannot see native font catalogs.
            config = root / "fonts.conf"
            config.write_text('<?xml version="1.0"?><fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include><dir>' + escape(str(font_dir.resolve())) + '</dir><cachedir>' + escape(str(root / "font-cache")) + '</cachedir></fontconfig>')
            environment["FONTCONFIG_FILE"] = str(config)
        try:
            result = subprocess.run([executable, "-env:UserInstallation=" + (root / "profile").as_uri(), "--headless", "--convert-to", "pdf", "--outdir", directory, str(root / "deck.pptx")], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=90, check=False, env=environment)
        except (subprocess.TimeoutExpired, OSError):
            return {"status": "failed", "message": "导出渲染超时或渲染器无法启动", "passed": False}, None, []
        if result.returncode != 0 or not (root / "deck.pdf").is_file():
            return {"status": "failed", "message": "无法渲染导出的 PPTX", "passed": False}, None, []
        pdf = (root / "deck.pdf").read_bytes()
        previews = []
        issues: list[dict[str, Any]] = []
        with pymupdf.open(stream=pdf, filetype="pdf") as document:
            if len(document) != page_count:
                issues.append({"code": "render_page_count", "message": "实际渲染页数与设计稿不同"})
            for index, page in enumerate(document):
                if not page.get_text().strip():
                    issues.append({"code": "render_blank_page", "slide": index + 1, "message": "渲染页面没有可见文字"})
                if any(glyph[1] == 0 and not chr(glyph[0]).isspace() for span in page.get_texttrace() for glyph in span["chars"]):
                    issues.append({"code": "render_missing_glyph", "slide": index + 1, "message": "渲染缺少字形，请安装所选中文字体后重试"})
                for word in page.get_text("words"):
                    if word[0] < -1 or word[1] < -1 or word[2] > page.rect.width + 1 or word[3] > page.rect.height + 1:
                        issues.append({"code": "render_text_bounds", "slide": index + 1, "message": "实际渲染文字超出页面"})
                        break
                scale = 2 if page_count == 1 else .6
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                previews.append("data:image/png;base64," + base64.b64encode(pixmap.tobytes("png")).decode())
        return {"status": "verified", "renderer": "LibreOffice", "passed": not issues, "page_count": page_count, "violations": issues, "animation_playback_verified": False}, pdf, previews


def design_file(deck: dict[str, Any]) -> bytes:
    return json.dumps(deck, ensure_ascii=False, indent=2, allow_nan=False).encode()
