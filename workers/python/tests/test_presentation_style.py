from __future__ import annotations

import copy
import io
import json
import zipfile
from unittest.mock import patch

import pytest
from pptx import Presentation

from ai_companion_worker import office_tools as office
from ai_companion_worker.presentation_design import layout_deck, plan_deck
from ai_companion_worker.presentation_projects import project_deck
from ai_companion_worker.presentation_render import export_pptx, normalize_assets
from ai_companion_worker.presentation_style import apply_style, contrast, design_style
from ai_companion_worker.task_quality import presentation_audience_content_violations


def design(deck):
    return {
        'name': '课堂清晰图解', 'rationale': '蓝白配色与分步讲解', 'theme_id': 'business',
        'palette': {'background': 'FFFFFF', 'surface': 'EFF5FF', 'ink': '142B45',
                    'muted': '415571', 'accent': '125CA4',
                    'heading_font': 'Noto Sans CJK SC', 'body_font': 'Noto Sans CJK SC'},
        'typography': {'title_size': 36, 'body_size': 22, 'radius': 5},
        'mode': 'teach', 'animation': 'explain',
        'layouts': [{'slide_id': s['id'], 'layout': s['layout'], 'regions': s.get('regions')} for s in deck['slides']],
    }


def test_builtin_project_is_editable_and_uses_own_product_facts():
    deck = plan_deck({'project_name': '伴AI', 'slide_count': 1})
    assert deck['theme_id'] == 'companion'
    assert len(deck['slides']) == 1
    normalize_assets(deck)
    scenes, problems = layout_deck(deck)
    assert not problems
    data = export_pptx(deck, scenes)
    ppt = Presentation(io.BytesIO(data))
    visible = '\n'.join(s.text for s in ppt.slides[0].shapes if s.has_text_frame)
    for value in ('情感陪伴', '生活助手', '工作伙伴', '长期记忆', '知识 Wiki', '可扩展技能'):
        assert value in visible
    for unrelated in ('财经', 'GaussDB', 'Eureka', '交易系统'):
        assert unrelated not in visible
    assert len(ppt.slides[0].shapes) > 25
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read('ppt/slides/slide1.xml').decode()
    assert xml.count('<p:cxnSp>') == 6
    assert '<p:timing>' in xml
    assert 'README.zh-CN.md' in deck['slides'][0]['notes']


@pytest.mark.parametrize('payload', [{'project_name': '未知项目'}])
def test_builtin_rejects_unknown_project(payload):
    with pytest.raises(ValueError):
        project_deck(payload)


def test_style_preserves_facts_assets_and_editable_structure():
    source = project_deck({'project_name': '伴AI'})
    before = copy.deepcopy(source)
    styled = apply_style(source, design(source), '适合课堂讲解的清晰图解')
    assert source == before
    assert styled['slides'] == source['slides']
    assert styled['assets'] == source['assets']
    assert styled['typography']['title_size'] == 36
    assert styled['style_profile']['prompt'] == '适合课堂讲解的清晰图解'
    normalize_assets(styled)
    scenes, problems = layout_deck(styled)
    assert not problems
    assert any(n['type'] == 'connector' for n in scenes[0]['nodes'])


@pytest.mark.parametrize('change', ['unknown_page', 'duplicate', 'wrong_layout', 'unsafe_font'])
def test_rejects_model_structure_or_font_errors(change):
    deck = project_deck({'project_name': '伴AI'})
    spec = design(deck)
    if change == 'unknown_page':
        spec['layouts'][0]['slide_id'] = 'forged'
    elif change == 'duplicate':
        spec['layouts'].append(copy.deepcopy(spec['layouts'][0]))
    elif change == 'wrong_layout':
        spec['layouts'][0]['layout'] = 'bullets'
    else:
        spec['palette']['body_font'] = 'Missing font'
    with pytest.raises(ValueError):
        apply_style(deck, spec, '课堂图解')


def test_low_contrast_is_repaired_and_animation_can_be_removed():
    deck = project_deck({'project_name': '伴AI'})
    spec = design(deck)
    spec['palette'].update(ink='EEEEEE', muted='EEEEEE', accent='FFFFFF')
    spec['animation'] = 'none'
    styled = apply_style(deck, spec, '黑白极简，关闭动画')
    for color in ('ink', 'muted', 'accent'):
        assert contrast(styled['brand'][color], styled['brand']['surface']) >= 4.5
    normalize_assets(styled)
    scenes, _ = layout_deck(styled)
    with zipfile.ZipFile(io.BytesIO(export_pptx(styled, scenes))) as archive:
        assert b'<p:timing>' not in archive.read('ppt/slides/slide1.xml')


def test_style_model_request_is_structured_and_usage_is_accounted(monkeypatch):
    deck = project_deck({'project_name': '伴AI'})
    monkeypatch.setenv('MODEL_PROVIDER', 'openrouter')
    monkeypatch.setenv('MODEL_API_KEY', 'test-only')
    monkeypatch.setenv('MODEL_BASE_URL', 'https://openrouter.ai/api/v1')
    response = {'choices': [{'finish_reason': 'stop', 'message': {'content': json.dumps(design(deck))}}],
                'usage': {'cost': 0.001, 'prompt_tokens': 1000, 'completion_tokens': 500}, 'model': 'openai/gpt-5-mini'}
    with patch('urllib.request.urlopen', return_value=io.BytesIO(json.dumps(response).encode())) as call:
        styled, usage = design_style(deck, '适合课堂讲解的清晰图解')
    request = call.call_args.args[0]
    # Real urllib transports headers as latin-1; catch accidental Chinese headers.
    for value in request.headers.values():
        value.encode('latin-1')
    payload = json.loads(request.data)
    assert payload['response_format']['json_schema']['strict']
    assert payload['provider']['require_parameters']
    assert usage['cost_micros'] == 1000
    assert styled['slides'] == deck['slides']


def test_style_failures_keep_usage_and_locks_prevent_model_calls(monkeypatch):
    deck = project_deck({'project_name': '伴AI'})
    monkeypatch.setenv('MODEL_PROVIDER', 'openrouter')
    monkeypatch.setenv('MODEL_API_KEY', 'test-only')
    monkeypatch.setenv('MODEL_BASE_URL', 'https://openrouter.ai/api/v1')
    response = {'choices': [{'finish_reason': 'length'}], 'usage': {'cost': 0.002}}
    with patch('urllib.request.urlopen', return_value=io.BytesIO(json.dumps(response).encode())):
        with pytest.raises(office.ModelBackedOperationError) as caught:
            design_style(deck, '课堂图解')
    assert caught.value.model_usage['cost_micros'] == 2000
    deck['slides'][0]['locked'] = True
    with patch('urllib.request.urlopen') as call:
        with pytest.raises(ValueError, match='解锁'):
            design_style(deck, '换配色')
        call.assert_not_called()


def test_audience_gate_accepts_only_known_builtin_or_authenticated_revision_marker():
    contract = {'artifact_types': ['pptx']}
    assert not presentation_audience_content_violations({'project_name': '伴AI', 'slide_count': 1}, contract)
    assert not presentation_audience_content_violations({'revision_of': 'latest'}, contract)
    assert presentation_audience_content_violations({'project_name': '未知项目', 'slide_count': 1}, contract)
    assert presentation_audience_content_violations({'brief': '', 'slide_count': 8}, contract)


@pytest.mark.parametrize('count', range(2, 11))
def test_arbitrary_project_content_uses_variable_diagram_regions(count):
    brief = '\n'.join(f'## 模块{i}\n- 用户提供的事实{i}' for i in range(count))
    payload = {'title': '智慧农业方案', 'brief': brief, 'slide_count': 1}
    deck = plan_deck(payload)
    normalize_assets(deck)
    scenes, problems = layout_deck(deck)
    assert not problems
    assert len(deck['slides']) == 1
    text = '\n'.join(n.get('text', '') for n in scenes[0]['nodes'])
    assert '伴AI' not in text
    for i in range(count):
        assert f'用户提供的事实{i}' in text
    assert not presentation_audience_content_violations(payload, {'artifact_types': ['pptx']})


def test_semantic_regions_cannot_drop_or_duplicate_content():
    deck = plan_deck({'title': '任意项目', 'slide_count': 1, 'brief': '## 目标\n- 完成课程\n## 教学\n- 解释概念\n## 练习\n- 检验掌握'})
    spec = design(deck)
    spec['layouts'][0]['regions'] = {'lead': 'block-0', 'main': ['block-1'], 'aside': '', 'foundation': ['block-2'], 'footer': '', 'flow': 'up'}
    styled = apply_style(deck, spec, '目标在上，练习支撑教学')
    assert styled['slides'][0]['regions']['lead'] == 'block-0'
    spec['layouts'][0]['regions']['foundation'] = []
    with pytest.raises(ValueError, match='完整'):
        apply_style(deck, spec, '课堂讲解')


def test_explicit_project_material_is_never_replaced_by_builtin_default():
    deck = plan_deck({'project_name': '伴AI', 'title': '温室方案', 'slide_count': 1,
                      'brief': '## 环境监测\n- 温湿度采集\n## 水肥控制\n- 分区灌溉'})
    text = '\n'.join(item['text'] for item in deck['slides'][0]['items'])
    assert '温湿度采集' in text and '分区灌溉' in text
    assert '情感陪伴' not in text


@pytest.mark.parametrize('repair_succeeds', [True, False])
def test_semantic_repair_is_bounded_and_accounts_for_every_attempt(monkeypatch, repair_succeeds):
    deck = project_deck({'project_name': '伴AI'})
    monkeypatch.setenv('MODEL_PROVIDER', 'openrouter')
    monkeypatch.setenv('MODEL_API_KEY', 'test-only')
    monkeypatch.setenv('MODEL_BASE_URL', 'https://openrouter.ai/api/v1')
    invalid = design(deck)
    invalid['layouts'][0]['regions'] = copy.deepcopy(invalid['layouts'][0]['regions'])
    invalid['layouts'][0]['regions']['main'].pop()
    def response(spec):
        return io.BytesIO(json.dumps({'choices': [{'finish_reason': 'stop', 'message': {'content': json.dumps(spec)}}],
                         'usage': {'cost': 0.001, 'prompt_tokens': 1000, 'completion_tokens': 500}}).encode())
    with patch('urllib.request.urlopen', side_effect=[response(invalid), response(design(deck) if repair_succeeds else invalid)]) as call:
        if repair_succeeds:
            styled, usage = design_style(deck, '温暖科技感')
            assert styled['slides'] == deck['slides']
        else:
            with pytest.raises(office.ModelBackedOperationError) as caught:
                design_style(deck, '温暖科技感')
            usage = caught.value.model_usage
    assert call.call_count == 2
    assert usage['cost_micros'] == 2000
    assert usage['prompt_tokens'] == 2000
    assert usage['attempts'] == 2
    retry = json.loads(call.call_args.args[0].data)
    assert '设计检查未通过' in retry['messages'][-1]['content']
    schema = retry['response_format']['json_schema']['schema']
    refs = schema['properties']['layouts']['items']['anyOf'][0]['properties']['regions']['properties']['main']['items']['enum']
    assert refs == [item['id'] for item in deck['slides'][0]['items']]


@pytest.mark.parametrize('builtin', [False, True])
def test_chat_exports_validated_design_without_requiring_a_new_brief(builtin):
    source = plan_deck({'title': '温室方案', 'slide_count': 1,
                        'brief': '## 管理目标\n- 统一查看环境与作业\n## 环境监测\n- 温湿度采集'})
    payload = {'title': '温室方案', 'brief': '', 'slide_count': 1,
               'task_contract': {'artifact_types': ['pptx', 'pdf'], 'exhaustive': True},
               'source_coverage': {'coverage_ratio': 1, 'truncated': False}}
    payload.update({'project_name': '伴AI'} if builtin else {'deck_spec': source})
    with patch('ai_companion_worker.presentation_studio.render_pptx', return_value=({'status': 'unavailable', 'passed': False}, None, [])):
        result = office.execute('pptx_generate', payload)
    assert result['output']['quality_report']['passed'], result['output']['quality_report']
    assert any(f['name'].endswith('.pptx') for f in result['files'])
    if not builtin:
        assert result['output']['deck_spec']['slides'][0]['items'] == source['slides'][0]['items']


def test_restored_design_still_enforces_source_coverage():
    source = project_deck({'project_name': '伴AI'})
    result = office.execute('pptx_generate', {'title': source['title'], 'brief': '', 'slide_count': 1,
                            'deck_spec': source, 'task_contract': {'artifact_types': ['pptx'], 'exhaustive': True},
                            'source_coverage': {'coverage_ratio': .5, 'truncated': True}})
    assert not result['output']['quality_report']['passed']
    assert not result['files']
