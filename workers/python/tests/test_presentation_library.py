from __future__ import annotations

import copy
import io
import json
import os
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from ai_companion_worker.presentation_design import THEMES, layout_deck, validate_deck
from ai_companion_worker.presentation_library import build_library
from ai_companion_worker.presentation_render import export_pptx, render_pptx


def test_browser_library_matches_the_production_layout_engine() -> None:
    catalog = Path(__file__).resolve().parents[3] / "web/public/presentation-library.json"
    assert json.loads(catalog.read_text()) == build_library()


@pytest.mark.parametrize("template", build_library()["templates"], ids=lambda item: item["id"])
def test_templates_preserve_complete_editable_content(template: dict[str, Any]) -> None:
    deck = validate_deck(template["deck"])
    scenes, violations = layout_deck(deck)
    assert not violations
    assert len(deck["slides"]) == 7
    ppt = Presentation(io.BytesIO(export_pptx(deck, scenes)))
    assert len(ppt.slides) == 7
    for source, slide in zip(deck["slides"], ppt.slides, strict=True):
        text = "".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert source["title"].replace("\n", "") in text.replace("\n", "")
        if source.get("table"):
            assert any(shape.has_table for shape in slide.shapes)
    assert not any(s.get("chart") for s in deck["slides"]), "Starter decks must not imply fabricated results"


@pytest.mark.parametrize("theme", list(THEMES))
def test_all_page_presets_export_without_layout_violations(theme: str) -> None:
    pages = build_library()["pages"]
    assert len(pages) == 18
    for preset in pages:
        design = preset["variants"][theme]
        scenes, violations = layout_deck(design["deck"])
        assert not violations, (theme, preset["id"], violations)
        ppt = Presentation(io.BytesIO(export_pptx(design["deck"], scenes)))
        assert len(ppt.slides) == 1
        source = design["deck"]["slides"][0]
        if source.get("chart"):
            chart = next(shape.chart for shape in ppt.slides[0].shapes if shape.has_chart)
            assert list(chart.series[0].values) == source["chart"]["values"]
            assert "示例" in source["title"]


@pytest.mark.parametrize("theme", list(THEMES))
def test_library_through_actual_office_renderer(theme: str) -> None:
    audit = os.environ.get("PRESENTATION_RENDER_AUDIT_DIR")
    if not audit:
        pytest.skip("Set PRESENTATION_RENDER_AUDIT_DIR for native rendering")
    library = build_library()
    template = next((t for t in library["templates"] if t["deck"]["theme_id"] == theme), library['templates'][0])
    deck = copy.deepcopy(template["deck"])
    deck['theme_id'] = theme
    deck["slides"] += [copy.deepcopy(p["variants"][theme]["deck"]["slides"][0]) for p in library["pages"]]
    for i, slide in enumerate(deck["slides"]):
        slide["id"] = f"page-{i}"
    scenes, violations = layout_deck(deck)
    assert not violations
    data = export_pptx(deck, scenes)
    report, pdf, previews = render_pptx(data, 25)
    assert report["passed"], report
    assert pdf and len(previews) == 25
    target = Path(audit)
    target.mkdir(parents=True, exist_ok=True)
    (target / f"library-{theme}.pptx").write_bytes(data)
    (target / f"library-{theme}.pdf").write_bytes(pdf)
    (target / f"library-{theme}-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
