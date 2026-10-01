"""Whole-deck regressions: diagrams are one content page, not the whole artifact."""
from __future__ import annotations

import copy
import io
import zipfile
from unittest.mock import patch

import pytest
from pptx import Presentation

from ai_companion_worker import office_tools as office
from ai_companion_worker.presentation_design import layout_deck, plan_deck, repair_layouts, validate_deck
from ai_companion_worker.presentation_render import export_pptx, normalize_assets
from ai_companion_worker.presentation_style import apply_style
from ai_companion_worker.task_quality import apply_planned_task_intent, compile_task_contract, presentation_audience_content_violations
from test_presentation_style import design

BRIEF = '''## 背景与目标
- 环境记录分散在不同设备
- 本次方案统一查看环境与作业
## 温室能力全景
### 环境监测
- 温湿度采集
- 异常提醒
### 水肥控制
- 分区灌溉
- 作业记录
### 设备与数据
- 传感器接入
- 历史记录查询
## 运行流程
- 采集环境数据
- 提醒值班人员
- 人员确认作业
## 方案对比
- 现场操作：由现场人员执行
- 远程查看：管理人员查看记录
## 试点计划
- 准备期：核对设备清单
- 试点期：验证环境数据
- 复盘期：整理使用反馈
## 下一步
- 确认试点范围
- 明确值班与反馈负责人'''
PAYLOAD = {'title': '智慧温室试点方案', 'audience': '项目评审团队', 'brief': BRIEF,
           'style': '商务', 'slide_count': 8, 'visual_mode': 'none'}


@pytest.mark.parametrize('message', [
    '做8页活动介绍PPT，介绍志愿者如何准备场地和发布活动。',
    '生成8页项目介绍PPT。资料如下：\n## 项目计划\n- 确认场地\n- 讨论活动日期和课程名称',
])
def test_project_facts_do_not_force_the_entire_deck_into_a_table(message):
    contract = apply_planned_task_intent(compile_task_contract(message, 'work'), {})
    assert contract['presentation_capabilities'] == ['narrative']
    assert contract['requested_fields'] == []


@pytest.mark.parametrize('message', ['用表格展示活动地点并生成PPT', '按时间字段整理所有活动并生成PPT'])
def test_explicit_table_and_field_requests_still_require_table_capability(message):
    contract = apply_planned_task_intent(compile_task_contract(message, 'work'), {})
    assert 'table' in contract['presentation_capabilities']


def test_diagram_is_one_page_of_an_editable_eight_page_story():
    before = copy.deepcopy(PAYLOAD)
    deck = plan_deck(PAYLOAD)
    assert PAYLOAD == before
    assert len(deck['slides']) == 8
    assert [s['layout'] for s in deck['slides']].count('capability-map') == 1
    assert deck['slides'][3]['layout'] == 'steps'
    assert deck['slides'][4]['layout'] == 'comparison'
    assert deck['slides'][5]['layout'] == 'timeline'
    assert [s['title'] for s in deck['slides'][1:-1]] == [line[3:] for line in BRIEF.splitlines() if line.startswith('## ')]
    deck, _ = repair_layouts(deck)
    scenes, problems = layout_deck(deck)
    assert not problems
    pptx = Presentation(io.BytesIO(export_pptx(deck, scenes)))
    assert len(pptx.slides) == 8
    visible = '\n'.join(shape.text for page in pptx.slides for shape in page.shapes if shape.has_text_frame)
    for line in BRIEF.splitlines():
        if line.startswith('- '):
            assert line[2:] in visible
    assert len(pptx.slides[2].shapes) > 10
    assert not presentation_audience_content_violations(PAYLOAD, {'artifact_types': ['pptx']})


def test_nested_brief_routes_chat_to_design_engine_without_explicit_style():
    output = office.execute('pptx_outline', PAYLOAD)['output']
    assert output['quality_report']['passed'], output['quality_report']
    assert len(output['deck_spec']['slides']) == 8


@pytest.mark.parametrize('brief', [
    '## 图解\n### 空模块\n### 另一个模块\n- 事实',
    '## 图解\n- 未归组事实\n### 模块甲\n- 事实甲\n### 模块乙\n- 事实乙',
])
def test_diagram_contract_rejects_empty_or_ungrouped_facts(brief):
    violations = presentation_audience_content_violations({**PAYLOAD, 'slide_count': 3, 'brief': brief}, {'artifact_types': ['pptx']})
    assert any(v['code'] == 'presentation_diagram_modules_invalid' for v in violations)


def test_whole_deck_direction_survives_export_and_preserves_facts_order_and_evidence():
    deck = plan_deck(PAYLOAD)
    deck['slides'][4].update(layout='chart-column', chart={'labels': ['甲', '乙'], 'values': [12, 18], 'series': '样本', 'unit': '项'})
    deck['slides'][6].update(layout='table', table={'columns': ['工作', '状态'], 'rows': [['核对清单', '待确认']]})
    deck = validate_deck(deck)
    spec = design(deck)
    spec['storyline'] = '说明现状与目标，解释系统构成与运行，再讨论试点安排。'
    for i, page in enumerate(spec['layouts']):
        page.update(purpose=f'说明第{i+1}页提供的信息', animation='explain' if i in (2, 3) else 'none', duration_ms=600)
    # The model response order must not overwrite the user's page order.
    spec['layouts'].reverse()
    styled = apply_style(deck, spec, '适合现场评审，流程逐步讲解')
    assert [s['id'] for s in styled['slides']] == [s['id'] for s in deck['slides']]
    for old, new in zip(deck['slides'], styled['slides'], strict=True):
        for key in ('title', 'items', 'notes', 'source', 'table', 'chart'):
            assert old.get(key) == new.get(key)
    styled, _ = repair_layouts(styled)
    assert styled['style_profile']['storyline'] == spec['storyline']
    scenes, problems = layout_deck(styled)
    assert not problems
    data = export_pptx(styled, scenes)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for i in range(8):
            xml = archive.read(f'ppt/slides/slide{i+1}.xml')
            assert (b'<p:timing>' in xml) == (i in (2, 3))
            assert '说明第'.encode() not in xml  # Editing guidance never leaks onto slides.
    native = Presentation(io.BytesIO(data))
    assert any(s.has_chart for s in native.slides[4].shapes)
    assert any(s.has_table for s in native.slides[6].shapes)


def test_builtin_complete_introduction_reuses_generic_planning_and_reviewed_sources():
    deck = plan_deck({'project_name': '伴AI', 'slide_count': 8})
    assert len(deck['slides']) == 8
    assert sum(s['layout'] == 'capability-map' for s in deck['slides']) == 1
    assert all('内置产品知识' in s['source'] for s in deck['slides'])
    normalize_assets(deck)
    deck, _ = repair_layouts(deck)
    _, problems = layout_deck(deck)
    assert not problems
    assert not presentation_audience_content_violations({'project_name': '伴AI', 'slide_count': 8}, {'artifact_types': ['pptx']})


def test_diagram_does_not_disable_images_on_other_pages():
    deck = plan_deck(PAYLOAD)
    visual = {'media_type': 'image/png', 'data': b'not-read-in-outline', 'source_locator': '测试资料'}
    with patch.object(office, '_presentation_visuals', return_value=([None, visual, None, None, None], {'placed_visual_count': 0}, {'provider': 'none', 'cost_micros': 0})) as visuals, \
         patch('ai_companion_worker.presentation_studio.normalize_assets'), \
         patch('ai_companion_worker.presentation_studio.repair_layouts', side_effect=lambda value: (value, [])):
        output = office.execute('pptx_outline', PAYLOAD)['output']['deck_spec']
    topics = visuals.call_args.args[1]
    assert len(topics) == 5 and all(s['title'] != '温室能力全景' for s in topics)
    assert output['slides'][2]['layout'] == 'capability-map'
    assert output['slides'][2]['items'] == deck['slides'][2]['items']
    assert output['slides'][3]['layout'] == 'image-right'
    assert output['slides'][3]['image_id'] == 'visual-1'
