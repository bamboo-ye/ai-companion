from __future__ import annotations

import base64
import io
import json
import os
import shutil
import tempfile
import threading
import time
import unittest
from dataclasses import replace
from unittest.mock import Mock, patch
from zipfile import ZipFile

import pymupdf  # type: ignore[import-untyped]
from PIL import Image
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from ai_companion_worker.office_tools import (
    CJK_FONT_CANDIDATES,
    ModelBackedOperationError,
    PRESENTATION_BODY_FONT_FAMILY,
    _openrouter_translate,
    _openrouter_generate_presentation_visual,
    _presentation_body_text_overflow_risk,
    _translate_layout_blocks,
    _translation_batches,
    _presentation_display_rows,
    _presentation_render_text,
    _presentation_topic_source_pages,
    _presentation_visuals,
    _translated_content,
    execute,
    main,
)
from ai_companion_worker.pdf_translation_layout import (
    PdfLayout,
    PdfTextBlock,
    extract_pdf_layout,
    render_layout_translation,
    _normalize_translation,
    _validate_cross_renderer_page_coverage,
)
from ai_companion_worker.translation_language import translation_language_policy


class OfficeToolsTest(unittest.TestCase):
    def test_pptx_visual_limit_is_capped_by_available_topics(self) -> None:
        with patch.dict(
            os.environ,
            {"MODEL_PRESENTATION_IMAGE_MAX_COUNT": "2"},
        ):
            assignments, report, usage = _presentation_visuals(
                {},
                [],
                include_file=True,
                allow_generated=False,
            )
        self.assertEqual(assignments, [])
        self.assertEqual(report["generation_attempt_count"], 0)
        self.assertEqual(usage["cost_micros"], 0)

    def test_presentation_source_visuals_follow_cited_pages_not_extraction_order(self) -> None:
        extracted = [
            {"kind": "source", "source_page": 1, "source_locator": "source.pdf · 第 1 页"},
            {"kind": "source", "source_page": 2, "source_locator": "source.pdf · 第 2 页"},
        ]
        payload = {
            "visual_mode": "source_only",
            "source_documents": [
                {
                    "filename": "source.pdf",
                    "media_type": "application/pdf",
                    "data_base64": base64.b64encode(b"pdf-placeholder").decode(),
                }
            ],
        }
        topics = [
            {"title": "第二主题", "bullets": ["结论（来源：第2页）"]},
            {"title": "第一主题", "bullets": ["事实（来源：第1页）"]},
        ]

        with patch(
            "ai_companion_worker.office_tools._extract_pdf_presentation_visuals",
            return_value=extracted,
        ):
            assignments, report, _ = _presentation_visuals(
                payload,
                topics,
                include_file=True,
                allow_generated=False,
            )

        self.assertEqual([item["source_page"] for item in assignments if item], [2, 1])
        self.assertEqual(report["source_visual_count"], 2)
        self.assertEqual(
            _presentation_topic_source_pages({"bullets": ["区间证据（来源：第17-18页）"]}),
            {17, 18},
        )

    def test_source_rich_presentation_skips_generated_images(self) -> None:
        extracted = [
            {"kind": "source", "source_page": page, "source_locator": f"第 {page} 页"}
            for page in range(1, 5)
        ]
        payload = {
            "visual_mode": "auto",
            "source_documents": [
                {
                    "filename": "source.pdf",
                    "media_type": "application/pdf",
                    "data_base64": base64.b64encode(b"pdf-placeholder").decode(),
                }
            ],
        }
        topics = [
            {"title": f"主题 {page}", "bullets": [f"事实（来源：第{page}页）"]}
            for page in range(1, 5)
        ]

        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_PRESENTATION_IMAGE_NAME": "openai/gpt-image-2",
                "MODEL_PRESENTATION_SOURCE_SUFFICIENT_COUNT": "4",
            },
        ), patch(
            "ai_companion_worker.office_tools._extract_pdf_presentation_visuals",
            return_value=extracted,
        ), patch(
            "ai_companion_worker.office_tools._openrouter_generate_presentation_visual",
            side_effect=AssertionError("source-rich deck must not generate an image"),
        ):
            assignments, report, usage = _presentation_visuals(
                payload,
                topics,
                include_file=True,
                allow_generated=True,
            )

        self.assertEqual(len([item for item in assignments if item]), 4)
        self.assertTrue(report["source_material_sufficient"])
        self.assertEqual(report["generation_skipped_reason"], "source_material_sufficient")
        self.assertEqual(report["generation_attempt_count"], 0)
        self.assertEqual(usage["cost_micros"], 0)

    def test_generated_images_run_as_independent_bounded_parallel_calls(self) -> None:
        active = 0
        peak_active = 0
        lock = threading.Lock()
        barrier = threading.Barrier(2)

        def generate(topic: dict[str, object], *, remaining_cost_micros: int):
            nonlocal active, peak_active
            with lock:
                active += 1
                peak_active = max(peak_active, active)
            try:
                barrier.wait(timeout=1)
                time.sleep(0.01)
                return {
                    "kind": "generated",
                    "source_locator": "AI 生成配图",
                    "title": topic["title"],
                }, {
                    "requested_model": "openai/gpt-image-2",
                    "returned_model": "openai/gpt-image-2",
                    "upstream_provider": "OpenAI",
                    "cost_micros": min(80_000, remaining_cost_micros),
                    "cost_accounting": "reported",
                    "latency_ms": 10,
                    "cache_hit": False,
                }
            finally:
                with lock:
                    active -= 1

        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_PRESENTATION_IMAGE_NAME": "openai/gpt-image-2",
                "MODEL_PRESENTATION_IMAGE_MAX_COUNT": "2",
                "MODEL_PRESENTATION_IMAGE_CONCURRENCY": "2",
            },
        ), patch(
            "ai_companion_worker.office_tools._openrouter_generate_presentation_visual",
            side_effect=generate,
        ):
            assignments, report, _usage = _presentation_visuals(
                {"visual_mode": "auto"},
                [
                    {"title": "主题一", "bullets": ["事实一"]},
                    {"title": "主题二", "bullets": ["事实二"]},
                ],
                include_file=True,
                allow_generated=True,
            )

        self.assertEqual(peak_active, 2)
        self.assertEqual(len([item for item in assignments if item]), 2)
        self.assertEqual(report["generation_attempt_count"], 2)
        self.assertEqual(report["generated_visual_count"], 2)

    def test_translation_normalization_removes_unrenderable_inline_bullets(self) -> None:
        normalized = _normalize_translation("• 第一项 •\n• 第二项")

        self.assertEqual(normalized, "- 第一项\n- 第二项")
        self.assertNotIn("•", normalized)

    @unittest.skipUnless(shutil.which("pdftoppm"), "Poppler is not available")
    def test_cross_renderer_validation_rejects_visually_blank_page(self) -> None:
        visible = io.BytesIO()
        document = canvas.Canvas(visible, pagesize=(420, 300))
        document.setFillColorRGB(0, 0, 0)
        document.rect(40, 40, 340, 220, fill=1, stroke=0)
        document.save()
        blank = io.BytesIO()
        document = canvas.Canvas(blank, pagesize=(420, 300))
        document.showPage()
        document.save()

        with self.assertRaisesRegex(ValueError, "visual_content_missing"):
            _validate_cross_renderer_page_coverage(visible.getvalue(), blank.getvalue())

    def test_generic_chinese_target_uses_simplified_chinese_policy(self) -> None:
        for target in ("中文", "Chinese", "zh", "普通話"):
            policy = translation_language_policy(target)
            self.assertTrue(policy.simplified_chinese)
            self.assertEqual(policy.output_label, "简体中文")
            self.assertIn("仅使用简体字", policy.model_label)

    @unittest.skipUnless(
        any(os.path.isfile(path) for path in CJK_FONT_CANDIDATES),
        "no local CJK font available",
    )
    def test_translated_pdf_preserves_page_and_image_layout_while_replacing_text(self) -> None:
        source = io.BytesIO()
        document = canvas.Canvas(source, pagesize=(420, 300))
        image_bytes = io.BytesIO()
        Image.new("RGB", (8, 8), (30, 120, 220)).save(image_bytes, format="PNG")
        image = ImageReader(io.BytesIO(image_bytes.getvalue()))
        document.drawImage(image, 230, 90, width=120, height=90)
        document.setFont("Helvetica", 16)
        document.drawString(40, 235, "Original heading")
        document.setFont("Helvetica", 10)
        document.drawString(40, 205, "Original body text")
        document.save()

        layout = extract_pdf_layout(source.getvalue())
        translations = {
            block.marker: ("中文标题" if block.block_no == 1 else "中文正文")
            for block in layout.blocks
        }
        rendered, report = render_layout_translation(
            layout,
            translations,
            font_path=next(path for path in CJK_FONT_CANDIDATES if os.path.isfile(path)),
        )

        self.assertTrue(rendered.startswith(b"%PDF"))
        self.assertTrue(report["layout_preserved"])
        self.assertEqual(report["image_count"], 1)
        self.assertEqual(report["text_block_count"], 2)
        translated = pymupdf.open(stream=rendered, filetype="pdf")
        try:
            self.assertEqual(translated.page_count, 1)
            self.assertEqual((translated[0].rect.width, translated[0].rect.height), (420, 300))
            self.assertEqual(len(translated[0].get_image_info(xrefs=True)), 1)
            text = translated[0].get_text("text")
            self.assertNotIn("Original heading", text)
            self.assertNotIn("Original body text", text)
            self.assertIn("中文标题", text)
            self.assertIn("中文正文", text)
        finally:
            translated.close()

    @unittest.skipUnless(
        any(os.path.isfile(path) for path in CJK_FONT_CANDIDATES),
        "no local CJK font available",
    )
    def test_translated_pdf_preserves_original_when_translation_is_not_legible(self) -> None:
        source = io.BytesIO()
        document = canvas.Canvas(source, pagesize=(420, 300))
        document.setFont("Helvetica", 10)
        document.drawString(40, 235, "Keep this original content")
        document.save()

        layout = extract_pdf_layout(source.getvalue())
        translations = {
            layout.blocks[0].marker: "\n".join(f"无法容纳的翻译行 {index}" for index in range(20))
        }
        rendered, report = render_layout_translation(
            layout,
            translations,
            font_path=next(path for path in CJK_FONT_CANDIDATES if os.path.isfile(path)),
            simplified_chinese=True,
        )

        self.assertEqual(report["untranslated_block_count"], 1)
        translated = pymupdf.open(stream=rendered, filetype="pdf")
        try:
            text = translated[0].get_text("text")
            self.assertIn("Keep this original content", text)
            self.assertNotIn("无法容纳", text)
        finally:
            translated.close()

    @unittest.skipUnless(
        any(os.path.isfile(path) for path in CJK_FONT_CANDIDATES),
        "no local CJK font available",
    )
    def test_translated_pdf_normalizes_script_symbols_and_duplicate_layout_boxes(self) -> None:
        source = io.BytesIO()
        document = canvas.Canvas(source, pagesize=(420, 300))
        document.setFont("Helvetica", 12)
        document.drawString(40, 235, "First source label")
        document.save()

        extracted = extract_pdf_layout(source.getvalue())
        first = extracted.blocks[0]
        second = replace(first, block_no=2, text="Second source label")
        layout = replace(extracted, blocks=(first, second))
        rendered, report = render_layout_translation(
            layout,
            {
                first.marker: "• 異或問題",
                second.marker: "🙂 • 表徵學習",
            },
            font_path=next(path for path in CJK_FONT_CANDIDATES if os.path.isfile(path)),
            simplified_chinese=True,
        )

        self.assertEqual(report["text_block_count"], 2)
        translated = pymupdf.open(stream=rendered, filetype="pdf")
        try:
            text = translated[0].get_text("text")
            searchable = text.replace("\u2011", "-").replace("\xa0", " ")
            self.assertIn("- 异或问题", searchable)
            self.assertIn("- 表征学习", searchable)
            self.assertNotIn("異", text)
            self.assertNotIn("徵", text)
            self.assertNotIn("\x00", text)
            blocks = [block for block in translated[0].get_text("blocks") if int(block[6]) == 0]
            self.assertEqual(len(blocks), 1)
        finally:
            translated.close()

    def test_presentation_display_rows_keep_one_logical_record(self) -> None:
        schedule = "；".join(f"T{index:02d} 周一 09:00-09:50" for index in range(12))
        rows = _presentation_display_rows(
            [
                {
                    "cells": [
                        "PED1305",
                        "体能训练",
                        schedule,
                    ],
                    "source_locator": "page:2",
                    "entity_id": "course:PED1305",
                }
            ]
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["cells"], ["PED1305", "体能训练", schedule])
        self.assertFalse(rows[0]["continuation"])

    def test_presentation_render_text_protects_codes_and_percentages(self) -> None:
        source = (
            "CS5491《AI 安全专题》平时70%、考试30%；"
            "plot_ccd；imshow+contour；StratifiedShuffleSplit；第19页"
        )
        rendered = _presentation_render_text(source)

        self.assertEqual(rendered.replace("\u2060", ""), source)
        self.assertIn("C\u2060S\u20605\u20604\u20609\u20601", rendered)
        self.assertIn("《\u2060A\u2060I\u2060 \u2060安\u2060全\u2060专\u2060题\u2060》", rendered)
        self.assertIn("7\u20600\u2060%", rendered)
        self.assertIn("\u2060".join("plot_ccd"), rendered)
        self.assertIn("\u2060".join("imshow+contour"), rendered)
        self.assertIn("\u2060".join("StratifiedShuffleSplit"), rendered)
        self.assertIn("\u2060".join("第19页"), rendered)

    def test_pptx_quality_rejects_english_visible_content_for_chinese_deck(self) -> None:
        result = execute(
            "pptx_generate",
            {
                "title": "Important Dates",
                "audience": "新生",
                "style": "简洁表格",
                "brief": "展示重要日期",
                "slide_count": 3,
                "table": {
                    "title": "Semester A 2026/27",
                    "columns": ["Month", "Event"],
                    "rows": [
                        {
                            "cells": ["July", "Release of Class Schedule"],
                            "source_locator": "Page 4",
                        }
                    ],
                },
                "task_contract": {
                    "requested_fields": [],
                    "output_language": "zh-CN",
                },
                "source_coverage": {"coverage_ratio": 1.0, "truncated": False},
            },
        )

        report = result["output"]["quality_report"]
        self.assertFalse(report["passed"])
        self.assertEqual(result["files"], [])
        self.assertIn(
            "presentation_visible_language_mismatch",
            {item["code"] for item in report["violations"]},
        )

    def test_translation_response_rejects_null_content(self) -> None:
        with self.assertRaisesRegex(ValueError, "translation_model_returned_empty_text"):
            _translated_content({"choices": [{"message": {"content": None}}]})
        self.assertEqual(
            _translated_content(
                {"choices": [{"message": {"content": [{"type": "text", "text": "翻译结果"}]}}]}
            ),
            "翻译结果",
        )

    def test_layout_translation_batches_large_documents_without_merging_blocks(self) -> None:
        blocks = tuple(
            PdfTextBlock(
                page_no=page_no,
                block_no=1,
                bbox=(10, 10, 200, 40),
                text=f"Page {page_no} source text",
                font_size=10,
                color=(0, 0, 0),
                rotation=0,
            )
            for page_no in range(1, 6)
        )
        layout = PdfLayout(
            source=b"%PDF-layout-test",
            parser_version="pymupdf-test-layout-v1",
            page_count=5,
            blocks=blocks,
            page_previews={page_no: b"jpeg" for page_no in range(1, 6)},
            page_sizes=tuple((300, 200) for _ in range(5)),
            page_rotations=(0, 0, 0, 0, 0),
            image_placements=(),
        )
        self.assertEqual([len(batch) for batch in _translation_batches(blocks)], [4, 1])

        def translated(text: str, _target: str, **kwargs: object) -> tuple[str, dict[str, object]]:
            previews = kwargs.get("page_previews")
            return text.replace("source text", "中文内容"), {
                "provider": "openrouter",
                "requested_model": "openai/gpt-5-mini",
                "returned_model": "openai/gpt-5-mini",
                "upstream_provider": "OpenAI",
                "prompt_tokens": 100,
                "completion_tokens": 20,
                "cost_micros": 10,
                "cost_accounting": "reported",
                "latency_ms": 5,
                "vision_page_count": len(previews) if isinstance(previews, list) else 0,
            }

        with patch(
            "ai_companion_worker.office_tools._openrouter_translate",
            side_effect=translated,
        ) as translate:
            translations, usage, rounds = _translate_layout_blocks(layout, "Chinese")

        self.assertEqual(translate.call_count, 2)
        self.assertEqual(rounds, 2)
        self.assertEqual(len(translations), 5)
        self.assertTrue(all(value.endswith("中文内容") for value in translations.values()))
        self.assertEqual(usage["cost_micros"], 20)
        self.assertEqual(usage["vision_page_count"], 5)
        self.assertEqual(usage["round_count"], 2)

    def test_layout_translation_uses_bounded_parallelism_and_reduces_wall_latency(
        self,
    ) -> None:
        blocks = tuple(
            PdfTextBlock(
                page_no=page_no,
                block_no=1,
                bbox=(10, 10, 200, 40),
                text=f"Page {page_no} source text",
                font_size=10,
                color=(0, 0, 0),
                rotation=0,
            )
            for page_no in range(1, 17)
        )
        layout = PdfLayout(
            source=b"%PDF-layout-concurrency-test",
            parser_version="pymupdf-test-layout-v1",
            page_count=16,
            blocks=blocks,
            page_previews={},
            page_sizes=tuple((300, 200) for _ in range(16)),
            page_rotations=tuple(0 for _ in range(16)),
            image_placements=(),
        )
        active = 0
        peak_active = 0
        lock = threading.Lock()

        def translated(text: str, _target: str, **_kwargs: object) -> tuple[str, dict[str, object]]:
            nonlocal active, peak_active
            with lock:
                active += 1
                peak_active = max(peak_active, active)
            try:
                time.sleep(0.05)
                return text.replace("source text", "中文内容"), {
                    "provider": "openrouter",
                    "requested_model": "openai/gpt-5-mini",
                    "returned_model": "openai/gpt-5-mini",
                    "upstream_provider": "OpenAI",
                    "prompt_tokens": 100,
                    "completion_tokens": 20,
                    "cost_micros": 10,
                    "cost_accounting": "reported",
                    "latency_ms": 50,
                    "vision_page_count": 0,
                }
            finally:
                with lock:
                    active -= 1

        with (
            patch(
                "ai_companion_worker.office_tools._openrouter_translate",
                side_effect=translated,
            ),
            patch.dict(os.environ, {"MODEL_TRANSLATION_CONCURRENCY": "1"}),
        ):
            sequential_started = time.perf_counter()
            _translate_layout_blocks(layout, "Chinese")
            sequential_elapsed = time.perf_counter() - sequential_started

        peak_active = 0
        with (
            patch(
                "ai_companion_worker.office_tools._openrouter_translate",
                side_effect=translated,
            ),
            patch.dict(os.environ, {"MODEL_TRANSLATION_CONCURRENCY": "2"}),
        ):
            parallel_started = time.perf_counter()
            translations, usage, rounds = _translate_layout_blocks(layout, "Chinese")
            parallel_elapsed = time.perf_counter() - parallel_started

        self.assertEqual(rounds, 4)
        self.assertEqual(len(translations), 16)
        self.assertEqual(peak_active, 2)
        self.assertEqual(usage["concurrency"], 2)
        self.assertEqual(usage["latency_ms"], 200)
        self.assertLess(usage["wall_latency_ms"], usage["latency_ms"])
        self.assertLess(parallel_elapsed, sequential_elapsed * 0.75)

    def test_parallel_translation_stops_dispatch_after_failure_and_retains_usage(
        self,
    ) -> None:
        blocks = tuple(
            PdfTextBlock(
                page_no=page_no,
                block_no=1,
                bbox=(10, 10, 200, 40),
                text=f"Page {page_no} source text",
                font_size=10,
                color=(0, 0, 0),
                rotation=0,
            )
            for page_no in range(1, 13)
        )
        layout = PdfLayout(
            source=b"%PDF-layout-failure-test",
            parser_version="pymupdf-test-layout-v1",
            page_count=12,
            blocks=blocks,
            page_previews={},
            page_sizes=tuple((300, 200) for _ in range(12)),
            page_rotations=tuple(0 for _ in range(12)),
            image_placements=(),
        )

        def translated(text: str, _target: str, **_kwargs: object) -> tuple[str, dict[str, object]]:
            if "[[PAGE 1 BLOCK 1]]" in text:
                raise ModelBackedOperationError(
                    "translation_model_truncated",
                    {
                        "provider": "openrouter",
                        "cost_micros": 7,
                        "cost_accounting": "reported",
                        "latency_ms": 2,
                    },
                )
            time.sleep(0.03)
            return text.replace("source text", "中文内容"), {
                "provider": "openrouter",
                "cost_micros": 5,
                "cost_accounting": "reported",
                "latency_ms": 30,
            }

        with (
            patch(
                "ai_companion_worker.office_tools._openrouter_translate",
                side_effect=translated,
            ) as translate,
            patch.dict(os.environ, {"MODEL_TRANSLATION_CONCURRENCY": "2"}),
        ):
            with self.assertRaises(ModelBackedOperationError) as caught:
                _translate_layout_blocks(layout, "Chinese")

        self.assertEqual(caught.exception.code, "translation_model_truncated")
        self.assertEqual(translate.call_count, 2)
        self.assertEqual(caught.exception.model_usage["round_count"], 2)
        self.assertEqual(caught.exception.model_usage["cost_micros"], 12)
        self.assertEqual(caught.exception.model_usage["concurrency"], 2)

    @patch("urllib.request.urlopen")
    def test_pdf_translation_uses_budgeted_openrouter_policy_and_usage(self, urlopen: Mock) -> None:
        response = urlopen.return_value.__enter__.return_value
        response.read.return_value = json.dumps(
            {
                "model": "openai/gpt-5-mini",
                "provider": "OpenAI",
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": "Translated text"},
                    }
                ],
                "usage": {
                    "prompt_tokens": 30,
                    "completion_tokens": 8,
                    "cost": 0.000004,
                },
            }
        ).encode()
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_NAME": "deepseek/deepseek-v4-flash-0731",
                "MODEL_TRANSLATION_MAX_TOKENS": "8000",
                "MODEL_TRANSLATION_MAX_COST_MICROS": "60000",
                "MODEL_TRANSLATION_MAX_PROMPT_PRICE": "0.3",
                "MODEL_TRANSLATION_MAX_COMPLETION_PRICE": "2.5",
                "MODEL_DATA_COLLECTION": "deny",
                "MODEL_ZDR_REQUIRED": "false",
                "MODEL_PROVIDER_SORT": "price",
                "MODEL_ALLOW_PROVIDER_FALLBACKS": "true",
            },
            clear=True,
        ):
            translated, usage = _openrouter_translate("Source text", "English")
        self.assertEqual(translated, "Translated text")
        self.assertEqual(usage["cost_micros"], 4)
        self.assertEqual(usage["cost_accounting"], "reported")
        self.assertEqual(usage["returned_model"], "openai/gpt-5-mini")
        request = urlopen.call_args.args[0]
        body = json.loads(request.data)
        self.assertNotIn("temperature", body)
        self.assertEqual(body["model"], "openai/gpt-5-mini")
        self.assertEqual(body["max_tokens"], 8000)
        self.assertEqual(body["provider"]["data_collection"], "deny")
        self.assertTrue(body["provider"]["require_parameters"])
        self.assertEqual(
            body["provider"]["max_price"],
            {"prompt": 0.3, "completion": 2.5},
        )

    @patch("urllib.request.urlopen")
    def test_pdf_translation_sends_page_preview_as_low_detail_visual_context(
        self,
        urlopen: Mock,
    ) -> None:
        response = urlopen.return_value.__enter__.return_value
        response.read.return_value = json.dumps(
            {
                "model": "openai/gpt-5-mini",
                "provider": "OpenAI",
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {
                            "content": "[[PAGE 1 BLOCK 1]]\n中文标题",
                        },
                    }
                ],
                "usage": {
                    "prompt_tokens": 200,
                    "completion_tokens": 20,
                    "cost": 0.00001,
                },
            }
        ).encode()
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_TRANSLATION_NAME": "openai/gpt-5-mini",
            },
        ):
            translated, usage = _openrouter_translate(
                "[[PAGE 1 BLOCK 1]]\nHeading",
                "Chinese",
                page_previews=[(1, b"jpeg-preview")],
            )

        self.assertEqual(translated, "[[PAGE 1 BLOCK 1]]\n中文标题")
        self.assertEqual(usage["vision_page_count"], 1)
        body = json.loads(urlopen.call_args.args[0].data)
        content = body["messages"][1]["content"]
        image_parts = [item for item in content if item["type"] == "image_url"]
        self.assertEqual(len(image_parts), 1)
        self.assertEqual(image_parts[0]["image_url"]["detail"], "low")
        self.assertTrue(image_parts[0]["image_url"]["url"].startswith("data:image/jpeg;base64,"))

    @patch("urllib.request.urlopen")
    def test_pdf_translation_rejects_truncated_response_and_retains_usage(
        self, urlopen: Mock
    ) -> None:
        response = urlopen.return_value.__enter__.return_value
        response.read.return_value = json.dumps(
            {
                "model": "openai/gpt-5-mini",
                "choices": [
                    {
                        "finish_reason": "length",
                        "message": {"content": "Partial translation"},
                    }
                ],
                "usage": {
                    "prompt_tokens": 30,
                    "completion_tokens": 8,
                    "cost": 0.000004,
                },
            }
        ).encode()
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_TRANSLATION_NAME": "openai/gpt-5-mini",
            },
        ):
            with self.assertRaises(ModelBackedOperationError) as caught:
                _openrouter_translate("Source text", "English")
        self.assertEqual(caught.exception.code, "translation_model_truncated")
        self.assertEqual(caught.exception.model_usage["cost_micros"], 4)

    @patch("urllib.request.urlopen")
    def test_pdf_translation_rejects_reported_cost_above_reserved_batch_budget(
        self,
        urlopen: Mock,
    ) -> None:
        response = urlopen.return_value.__enter__.return_value
        response.read.return_value = json.dumps(
            {
                "model": "openai/gpt-5-mini",
                "provider": "OpenAI",
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": "Translated text"},
                    }
                ],
                "usage": {
                    "prompt_tokens": 30,
                    "completion_tokens": 8,
                    "cost": 0.006,
                },
            }
        ).encode()
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_TRANSLATION_NAME": "openai/gpt-5-mini",
            },
        ):
            with self.assertRaises(ModelBackedOperationError) as caught:
                _openrouter_translate(
                    "Source text",
                    "English",
                    remaining_cost_micros=5_000,
                )
        self.assertEqual(caught.exception.code, "translation_model_budget_exhausted")
        self.assertEqual(caught.exception.model_usage["cost_micros"], 6_000)
        self.assertEqual(caught.exception.model_usage["max_cost_micros"], 5_000)

    @patch("urllib.request.urlopen")
    def test_pdf_translation_requires_the_complete_page_marker_sequence(
        self, urlopen: Mock
    ) -> None:
        response = urlopen.return_value.__enter__.return_value
        response.read.return_value = json.dumps(
            {
                "model": "openai/gpt-5-mini",
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": "[[PAGE 1]]\nFirst page translated"},
                    }
                ],
                "usage": {
                    "prompt_tokens": 40,
                    "completion_tokens": 10,
                    "cost": 0.000005,
                },
            }
        ).encode()
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_TRANSLATION_NAME": "openai/gpt-5-mini",
            },
        ):
            with self.assertRaises(ModelBackedOperationError) as caught:
                _openrouter_translate("[[PAGE 1]]\nFirst\n\n[[PAGE 2]]\nSecond", "English")
        self.assertEqual(caught.exception.code, "translation_page_markers_invalid")
        self.assertEqual(caught.exception.model_usage["cost_micros"], 5)

    @patch("urllib.request.urlopen")
    def test_pdf_translation_invalid_content_retains_usage(self, urlopen: Mock) -> None:
        response = urlopen.return_value.__enter__.return_value
        response.read.return_value = json.dumps(
            {
                "model": "openai/gpt-5-mini",
                "choices": [{"finish_reason": "stop", "message": {"content": None}}],
                "usage": {
                    "prompt_tokens": 20,
                    "completion_tokens": 1,
                    "cost": 0.000003,
                },
            }
        ).encode()
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_TRANSLATION_NAME": "openai/gpt-5-mini",
            },
        ):
            with self.assertRaises(ModelBackedOperationError) as caught:
                _openrouter_translate("Source text", "English")
        self.assertEqual(caught.exception.code, "translation_model_returned_empty_text")
        self.assertEqual(caught.exception.model_usage["cost_micros"], 3)

    @patch("urllib.request.urlopen")
    def test_pdf_translation_invalid_json_retains_reserved_usage_bound(self, urlopen: Mock) -> None:
        response = urlopen.return_value.__enter__.return_value
        response.read.return_value = b"not-json"
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_TRANSLATION_NAME": "openai/gpt-5-mini",
                "MODEL_TRANSLATION_MAX_COST_MICROS": "60000",
            },
        ):
            with self.assertRaises(ModelBackedOperationError) as caught:
                _openrouter_translate("Source text", "English")
        self.assertEqual(caught.exception.code, "translation_model_returned_invalid_json")
        self.assertEqual(caught.exception.model_usage["cost_micros"], 60_000)
        self.assertEqual(
            caught.exception.model_usage["cost_accounting"],
            "reserved_upper_bound",
        )

    def test_pdf_translation_render_failure_retains_model_usage(self) -> None:
        source = io.BytesIO()
        document = canvas.Canvas(source)
        document.drawString(72, 760, "Source page")
        document.save()
        usage = {"provider": "openrouter", "cost_micros": 7}
        with (
            patch(
                "ai_companion_worker.office_tools._openrouter_translate",
                return_value=("[[PAGE 1 BLOCK 1]]\nTranslated page", usage),
            ),
            patch(
                "ai_companion_worker.office_tools.render_layout_translation",
                side_effect=RuntimeError("render exploded"),
            ),
        ):
            with self.assertRaises(ModelBackedOperationError) as caught:
                execute(
                    "pdf_translate",
                    {
                        "source_filename": "source.pdf",
                        "source_base64": base64.b64encode(source.getvalue()).decode(),
                        "target_language": "English",
                    },
                )
        self.assertEqual(caught.exception.code, "translation_render_failed")
        self.assertEqual(caught.exception.model_usage["provider"], usage["provider"])
        self.assertEqual(caught.exception.model_usage["cost_micros"], usage["cost_micros"])
        self.assertEqual(caught.exception.model_usage["round_count"], 1)

    def test_main_transports_billed_operation_failure_as_structured_output(
        self,
    ) -> None:
        usage = {"provider": "openrouter", "cost_micros": 9}
        stdin = io.StringIO('{"operation":"pdf_translate","input":{}}')
        with (
            patch("sys.stdin", stdin),
            patch("sys.stdout", new_callable=io.StringIO) as stdout,
            patch(
                "ai_companion_worker.office_tools.execute",
                side_effect=ModelBackedOperationError("translation_render_failed", usage),
            ),
        ):
            main()
        envelope = json.loads(stdout.getvalue())
        self.assertEqual(envelope["files"], [])
        self.assertEqual(envelope["output"]["model_usage"], usage)
        self.assertEqual(
            envelope["output"]["operation_error"]["code"],
            "translation_render_failed",
        )

    @patch("urllib.request.urlopen")
    def test_pdf_translation_budget_can_block_before_dispatch(self, urlopen: Mock) -> None:
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_TRANSLATION_NAME": "openai/gpt-5-mini",
                "MODEL_TRANSLATION_MAX_COST_MICROS": "100",
            },
        ):
            with self.assertRaisesRegex(ValueError, "budget_exhausted"):
                _openrouter_translate("Source text", "English")
        urlopen.assert_not_called()

    @patch("urllib.request.urlopen")
    def test_pdf_translation_rejects_relaxed_price_ceiling(self, urlopen: Mock) -> None:
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_API_KEY": "test-key",
                "MODEL_TRANSLATION_NAME": "openai/gpt-5-mini",
                "MODEL_TRANSLATION_MAX_PROMPT_PRICE": "0.31",
            },
        ):
            with self.assertRaisesRegex(ValueError, "price_ceiling"):
                _openrouter_translate("Source text", "English")
        urlopen.assert_not_called()

    def test_docx_edit_creates_copy_and_keeps_source_unchanged(self) -> None:
        source = io.BytesIO()
        document = Document()
        document.add_paragraph("原始段落")
        document.save(source)
        original = source.getvalue()
        result = execute(
            "docx_edit",
            {
                "source_filename": "source.docx",
                "source_base64": base64.b64encode(original).decode(),
                "append_text": "新增第一段\n新增第二段",
            },
        )
        generated = base64.b64decode(result["files"][0]["data_base64"])
        edited = Document(io.BytesIO(generated))
        self.assertEqual(
            [item.text for item in edited.paragraphs], ["原始段落", "新增第一段", "新增第二段"]
        )
        self.assertEqual(source.getvalue(), original)
        self.assertFalse(result["output"]["source_overwritten"])

    def test_pptx_generate_has_requested_slide_count_and_outline(self) -> None:
        result = execute(
            "pptx_generate",
            {
                "title": "季度复盘",
                "audience": "管理层",
                "style": "简洁",
                "brief": "业绩亮点\n风险与机会",
                "slide_count": 5,
            },
        )
        generated = base64.b64decode(result["files"][0]["data_base64"])
        with ZipFile(io.BytesIO(generated)) as archive:
            slides = [
                name
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ]
        self.assertEqual(len(slides), 5)
        self.assertEqual(len(result["output"]["outline"]), 5)
        self.assertTrue(result["output"]["quality_report"]["passed"])

        outline = execute(
            "pptx_outline",
            {
                "title": "季度复盘",
                "audience": "管理层",
                "style": "简洁",
                "brief": "业绩亮点\n风险与机会",
                "slide_count": 5,
            },
        )
        self.assertEqual(outline["files"], [])
        self.assertEqual(outline["output"]["outline"], result["output"]["outline"])

    def test_pptx_reuses_meaningful_pdf_image_with_source_provenance(self) -> None:
        image_bytes = io.BytesIO()
        Image.new("RGB", (800, 450), (45, 105, 190)).save(image_bytes, format="PNG")
        source = io.BytesIO()
        document = canvas.Canvas(source, pagesize=(420, 300))
        document.drawImage(
            ImageReader(io.BytesIO(image_bytes.getvalue())),
            165,
            70,
            width=225,
            height=127,
        )
        document.setFont("Helvetica", 15)
        document.drawString(35, 245, "Source visual")
        document.save()

        result = execute(
            "pptx_generate",
            {
                "title": "多模态课程介绍",
                "audience": "选课学生",
                "style": "图文简洁",
                "brief": "## 课程重点\n- 理解核心内容（来源：第1页）\n- 识别学习目标（来源：第1页）\n## 选课建议\n- 结合个人方向选择（来源：第1页）\n- 对照课程要求评估（来源：第1页）",
                "slide_count": 4,
                "visual_mode": "source_only",
                "source_documents": [
                    {
                        "filename": "课程原文.pdf",
                        "media_type": "application/pdf",
                        "data_base64": base64.b64encode(source.getvalue()).decode(),
                    }
                ],
            },
        )

        report = result["output"]["visual_report"]
        self.assertEqual(report["source_document_count"], 1)
        self.assertEqual(report["extracted_visual_count"], 1)
        self.assertEqual(report["source_visual_count"], 1)
        self.assertEqual(report["generated_visual_count"], 0)
        self.assertEqual(result["output"]["model_usage"]["cost_micros"], 0)
        self.assertEqual(
            result["output"]["outline"][1]["visual"]["source_locator"],
            "课程原文.pdf · 第 1 页",
        )
        generated = base64.b64decode(result["files"][0]["data_base64"])
        deck = Presentation(io.BytesIO(generated))
        self.assertEqual(
            deck.slides[1].placeholders[1].text_frame.vertical_anchor,
            MSO_ANCHOR.TOP,
        )
        with ZipFile(io.BytesIO(generated)) as archive:
            self.assertTrue(any(name.startswith("ppt/media/") for name in archive.namelist()))

    def test_pptx_dense_visual_body_uses_capacity_safe_typography(self) -> None:
        image_bytes = io.BytesIO()
        Image.new("RGB", (800, 450), (45, 105, 190)).save(image_bytes, format="PNG")
        source = io.BytesIO()
        document = canvas.Canvas(source, pagesize=(420, 300))
        document.drawImage(
            ImageReader(io.BytesIO(image_bytes.getvalue())),
            165,
            70,
            width=225,
            height=127,
        )
        document.save()
        bullets = [
            f"第{index}项内容说明分类器如何依据训练样本计算条件概率并保持数值稳定"
            f"，同时记录完整来源位置以便复核（来源：第1页）"
            for index in range(1, 6)
        ]

        result = execute(
            "pptx_generate",
            {
                "title": "密集图文排版",
                "audience": "学习者",
                "style": "图文简洁",
                "brief": "## 核心内容\n" + "\n".join(f"- {bullet}" for bullet in bullets),
                "slide_count": 3,
                "visual_mode": "source_only",
                "source_documents": [
                    {
                        "filename": "课程原文.pdf",
                        "media_type": "application/pdf",
                        "data_base64": base64.b64encode(source.getvalue()).decode(),
                    }
                ],
            },
        )

        self.assertTrue(result["output"]["quality_report"]["passed"])
        generated = base64.b64decode(result["files"][0]["data_base64"])
        deck = Presentation(io.BytesIO(generated))
        body = deck.slides[1].placeholders[1]
        bullet_markers = [
            shape
            for shape in deck.slides[1].shapes
            if shape.name.startswith("Native Bullet Marker ")
        ]
        bullet_texts = [
            shape
            for shape in deck.slides[1].shapes
            if shape.name.startswith("Bullet Text ")
        ]
        self.assertEqual(len(bullet_markers), 5)
        self.assertEqual(len(bullet_texts), 5)
        self.assertEqual(
            {shape.text_frame.paragraphs[0].font.size.pt for shape in bullet_markers},
            {15},
        )
        self.assertEqual(
            {shape.text_frame.paragraphs[0].space_after.pt for shape in bullet_markers},
            {0},
        )
        for current, following in zip(bullet_markers, bullet_markers[1:]):
            self.assertLessEqual(current.top + current.height, following.top)
        for marker, text_shape, expected_text in zip(
            bullet_markers,
            bullet_texts,
            bullets,
            strict=True,
        ):
            self.assertEqual(marker.top, text_shape.top)
            self.assertEqual(marker.height, text_shape.height)
            self.assertEqual(marker.left + marker.width, text_shape.left)
            self.assertEqual(text_shape.text.replace("\u2060", ""), expected_text)
            self.assertEqual(text_shape.text_frame.margin_left, 0)
            self.assertEqual(text_shape.text_frame.paragraphs[0].font.size.pt, 15)
            self.assertFalse(_presentation_body_text_overflow_risk(text_shape))
        for shape in bullet_markers:
            self.assertEqual(len(shape.text_frame.paragraphs), 1)
            paragraph = shape.text_frame.paragraphs[0]
            self.assertEqual(paragraph.text, "\u2060")
            properties = paragraph._p.pPr
            self.assertEqual(properties.get("marL"), str(int(Pt(18))))
            self.assertEqual(properties.get("indent"), str(-int(Pt(9))))
            bullet_elements = {
                child.tag.rsplit("}", 1)[-1]: child
                for child in properties
                if "}bu" in child.tag
            }
            self.assertEqual(bullet_elements["buChar"].get("char"), "•")
            self.assertEqual(bullet_elements["buSzPct"].get("val"), "100000")
            self.assertEqual(
                bullet_elements["buFont"].get("typeface"),
                PRESENTATION_BODY_FONT_FAMILY,
            )
            self.assertFalse(_presentation_body_text_overflow_risk(shape))
        self.assertEqual(body.text_frame.vertical_anchor, MSO_ANCHOR.TOP)

    def test_presentation_gate_detects_wrapped_body_overflow_risk(self) -> None:
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        body = slide.placeholders[1]
        body.left = Inches(0.78)
        body.top = Inches(1.42)
        body.width = Inches(5.45)
        body.height = Inches(5.28)
        frame = body.text_frame
        frame.clear()
        for index in range(5):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = "这是需要换行显示的长段落，用于验证正文虽然位于画布内但渲染后仍会越过占位区域。" * 2
            paragraph.font.size = Pt(17)
            paragraph.line_spacing = 1.1
            paragraph.space_after = Pt(8)

        self.assertTrue(_presentation_body_text_overflow_risk(body))

    @patch("urllib.request.urlopen")
    def test_presentation_image_result_is_reused_across_retry(self, urlopen: Mock) -> None:
        image_bytes = io.BytesIO()
        Image.new("RGB", (1600, 900), (31, 95, 152)).save(image_bytes, format="JPEG")
        response = urlopen.return_value.__enter__.return_value
        response.read.return_value = json.dumps(
            {
                "data": [{"b64_json": base64.b64encode(image_bytes.getvalue()).decode()}],
                "model": "openai/gpt-image-2",
                "provider": "OpenAI",
                "usage": {"cost": 0.08},
            }
        ).encode()
        topic = {"title": "并行生成", "bullets": ["结果应跨重试复用"]}

        with tempfile.TemporaryDirectory() as directory, patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_BASE_URL": "https://openrouter.ai/api/v1",
                "MODEL_API_KEY": "test-key",
                "MODEL_PRESENTATION_IMAGE_NAME": "openai/gpt-image-2",
                "WORKER_RESULT_CACHE_ENABLED": "true",
                "WORKER_RESULT_CACHE_DIR": directory,
            },
        ):
            first_visual, first_usage = _openrouter_generate_presentation_visual(
                topic,
                remaining_cost_micros=100_000,
            )
            second_visual, second_usage = _openrouter_generate_presentation_visual(
                topic,
                remaining_cost_micros=100_000,
            )

        urlopen.assert_called_once()
        self.assertEqual(first_visual["data"], second_visual["data"])
        self.assertEqual(first_usage["cost_micros"], 80_000)
        self.assertTrue(second_usage["cache_hit"])
        self.assertEqual(second_usage["cost_micros"], 0)

    @patch("urllib.request.urlopen", side_effect=OSError("network unavailable"))
    def test_pptx_image_generation_failure_safely_falls_back_to_text(self, urlopen: Mock) -> None:
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_BASE_URL": "https://openrouter.ai/api/v1",
                "MODEL_API_KEY": "test-key",
                "MODEL_PRESENTATION_IMAGE_NAME": "openai/gpt-image-1",
                "MODEL_PRESENTATION_IMAGE_MAX_COUNT": "1",
                "WORKER_RESULT_CACHE_ENABLED": "false",
            },
        ):
            result = execute(
                "pptx_generate",
                {
                    "title": "安全降级演示",
                    "audience": "项目团队",
                    "style": "简洁",
                    "brief": "## 当前状态\n- 内容完整\n## 下一步\n- 继续推进",
                    "slide_count": 4,
                },
            )

        report = result["output"]["visual_report"]
        self.assertEqual(report["generation_attempt_count"], 1)
        self.assertEqual(report["generated_visual_count"], 0)
        self.assertEqual(report["fallback_text_only_count"], 2)
        self.assertEqual(report["generation_errors"], ["presentation_image_request_failed"])
        self.assertEqual(result["output"]["model_usage"]["cost_micros"], 0)
        self.assertEqual(len(result["files"]), 1)
        urlopen.assert_called_once()

    @patch("urllib.request.urlopen")
    def test_pptx_generates_visual_through_openrouter_image_api(self, urlopen: Mock) -> None:
        image_bytes = io.BytesIO()
        Image.new("RGB", (1600, 900), (220, 140, 65)).save(image_bytes, format="JPEG")
        response = urlopen.return_value.__enter__.return_value
        response.read.return_value = json.dumps(
            {
                "data": [{"b64_json": base64.b64encode(image_bytes.getvalue()).decode()}],
                "model": "openai/gpt-image-2",
                "provider": "OpenAI",
                "usage": {"cost": 0.08},
            }
        ).encode()
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_BASE_URL": "https://openrouter.ai/api/v1",
                "MODEL_API_KEY": "test-key",
                "MODEL_PRESENTATION_IMAGE_NAME": "openai/gpt-image-2",
                "MODEL_PRESENTATION_IMAGE_MAX_COUNT": "1",
                "WORKER_RESULT_CACHE_ENABLED": "false",
            },
        ):
            result = execute(
                "pptx_generate",
                {
                    "title": "生成配图演示",
                    "audience": "项目团队",
                    "style": "图文简洁",
                    "brief": "## 核心变化\n- 流程更加清晰\n## 实施结果\n- 信息完整呈现",
                    "slide_count": 4,
                },
            )

        report = result["output"]["visual_report"]
        self.assertEqual(report["generation_attempt_count"], 1)
        self.assertEqual(report["generated_visual_count"], 1)
        self.assertEqual(report["placed_visual_count"], 1)
        self.assertEqual(result["output"]["model_usage"]["cost_micros"], 80_000)
        request = urlopen.call_args.args[0]
        self.assertEqual(request.full_url, "https://openrouter.ai/api/v1/images")
        request_payload = json.loads(request.data)
        self.assertEqual(request_payload["aspect_ratio"], "16:9")
        self.assertEqual(request_payload["output_format"], "jpeg")
        generated = base64.b64decode(result["files"][0]["data_base64"])
        with ZipFile(io.BytesIO(generated)) as archive:
            self.assertTrue(any(name.startswith("ppt/media/") for name in archive.namelist()))

    @patch("urllib.request.urlopen")
    def test_pptx_table_and_visual_capabilities_can_be_composed(self, urlopen: Mock) -> None:
        image_bytes = io.BytesIO()
        Image.new("RGB", (1600, 900), (52, 99, 155)).save(image_bytes, format="JPEG")
        response = urlopen.return_value.__enter__.return_value
        response.read.return_value = json.dumps(
            {
                "data": [{"b64_json": base64.b64encode(image_bytes.getvalue()).decode()}],
                "model": "openai/gpt-image-2",
                "provider": "OpenAI",
                "usage": {"cost": 0.08},
            }
        ).encode()
        with patch.dict(
            os.environ,
            {
                "MODEL_PROVIDER": "openrouter",
                "MODEL_BASE_URL": "https://openrouter.ai/api/v1",
                "MODEL_API_KEY": "test-key",
                "MODEL_PRESENTATION_IMAGE_NAME": "openai/gpt-image-2",
                "MODEL_PRESENTATION_IMAGE_MAX_COUNT": "1",
                "WORKER_RESULT_CACHE_ENABLED": "false",
            },
        ):
            result = execute(
                "pptx_generate",
                {
                    "title": "智能系统比较",
                    "audience": "技术团队",
                    "style": "简洁表格与插图",
                    "brief": "比较三类智能系统",
                    "slide_count": 4,
                    "visual_mode": "auto",
                    "table": {
                        "title": "三类系统比较",
                        "columns": ["系统", "工作方式", "优势", "局限"],
                        "rows": [
                            {
                                "cells": ["规则系统", "规则推理", "可解释", "维护成本高"],
                                "source_locator": "用户输入",
                            }
                        ],
                    },
                    "task_contract": {
                        "presentation_capabilities": ["narrative", "table", "visual"],
                        "output_language": "zh-CN",
                    },
                },
            )

        report = result["output"]["visual_report"]
        self.assertEqual(report["generation_attempt_count"], 1)
        self.assertEqual(report["generated_visual_count"], 1)
        self.assertEqual(report["placed_visual_count"], 1)
        self.assertTrue(result["output"]["quality_report"]["passed"])
        self.assertIn("visual", result["output"]["outline"][-1])
        self.assertEqual(len(result["files"]), 1)

    def test_pptx_layout_balances_cover_table_and_summary(self) -> None:
        result = execute(
            "pptx_generate",
            {
                "title": "智能系统比较",
                "audience": "技术团队",
                "style": "简洁表格",
                "brief": "比较不同系统的工作方式、优势和局限",
                "slide_count": 4,
                "visual_mode": "none",
                "table": {
                    "title": "三类系统比较",
                    "columns": ["系统", "工作方式", "优势", "局限"],
                    "rows": [
                        {
                            "cells": ["规则系统", "规则推理", "透明", "维护成本高"],
                            "source_locator": "用户输入",
                        },
                        {
                            "cells": ["机器学习", "样本学习", "适应复杂数据", "依赖数据"],
                            "source_locator": "用户输入",
                        },
                        {
                            "cells": ["多智能体", "协作分工", "可并行", "协调复杂"],
                            "source_locator": "用户输入",
                        },
                    ],
                },
                "task_contract": {"presentation_capabilities": ["narrative", "table"]},
            },
        )

        deck = Presentation(io.BytesIO(base64.b64decode(result["files"][0]["data_base64"])))
        cover = deck.slides[0]
        self.assertGreaterEqual(cover.shapes.title.text_frame.paragraphs[0].font.size.pt, 40)
        self.assertLess(cover.placeholders[1].top, Inches(4.0))

        table_slide = deck.slides[1]
        self.assertEqual(
            table_slide.shapes.title.text_frame.paragraphs[0].alignment,
            PP_ALIGN.LEFT,
        )
        table_shape = next(shape for shape in table_slide.shapes if shape.has_table)
        self.assertLess(table_shape.height, Inches(5.15))
        for row in table_shape.table.rows:
            for cell in row.cells:
                self.assertEqual(cell.vertical_anchor, MSO_ANCHOR.MIDDLE)

        summary = deck.slides[-1]
        self.assertEqual(summary.placeholders[1].text_frame.vertical_anchor, MSO_ANCHOR.MIDDLE)

    def test_pptx_explicit_visual_capability_fails_closed_without_an_image(self) -> None:
        with patch.dict(os.environ, {"MODEL_PRESENTATION_IMAGE_NAME": ""}):
            result = execute(
                "pptx_generate",
                {
                    "title": "需要插图的演示",
                    "audience": "项目团队",
                    "style": "简洁",
                    "brief": "## 核心变化\n- 流程更加清晰",
                    "slide_count": 3,
                    "visual_mode": "auto",
                    "task_contract": {
                        "presentation_capabilities": ["narrative", "visual"],
                    },
                },
            )

        report = result["output"]["quality_report"]
        self.assertFalse(report["passed"])
        self.assertIn(
            "requested_visual_missing",
            {item["code"] for item in report["violations"]},
        )
        self.assertEqual(result["files"], [])

    def test_pptx_explicit_table_capability_fails_closed_without_a_table(self) -> None:
        result = execute(
            "pptx_generate",
            {
                "title": "需要表格的演示",
                "audience": "项目团队",
                "style": "简洁",
                "brief": "## 三种方案\n- 规则系统\n- 机器学习系统\n- 多智能体系统",
                "slide_count": 3,
                "visual_mode": "none",
                "task_contract": {
                    "presentation_capabilities": ["narrative", "table"],
                },
            },
        )

        report = result["output"]["quality_report"]
        self.assertFalse(report["passed"])
        self.assertIn(
            "requested_table_missing",
            {item["code"] for item in report["violations"]},
        )
        self.assertEqual(result["files"], [])

    def test_pptx_title_with_slash_is_safe_for_outline_and_generated_filename(self) -> None:
        payload = {
            "title": "Important Dates - Semester A 2026/27",
            "audience": "New Students",
            "style": "professional",
            "brief": "Key dates timeline and summary",
            "slide_count": 5,
        }

        outline = execute("pptx_outline", payload)
        self.assertEqual(outline["files"], [])
        self.assertEqual(outline["output"]["title"], "Important Dates - Semester A 2026/27")

        presentation = execute("pptx_generate", payload)
        self.assertEqual(
            presentation["files"][0]["name"],
            "Important-Dates-Semester-A-2026-27.pptx",
        )

    def test_pptx_canonical_brief_renders_one_audience_section_per_slide(self) -> None:
        brief = "\n".join(
            (
                "## CS5187：视觉计算",
                "- 聚焦图像处理与视觉应用。",
                "- 强调视觉问题的建模和实践。",
                "## CS5297：人工智能",
                "- 覆盖人工智能方法与应用。",
                "- 强调智能算法的系统理解。",
                "## CS5491：人工智能安全",
                "- 关注模型风险与可信部署。",
                "- 强调安全评估和防护能力。",
                "## 横向比较",
                "- 三门课分别侧重视觉、通用智能与安全。",
                "- 课程选择应结合基础和项目方向。",
                "## 选择建议",
                "- 视觉方向优先考虑 CS5187。",
                "- 智能或安全方向可考虑 CS5297、CS5491。",
            )
        )
        result = execute(
            "pptx_generate",
            {
                "title": "三门课程对比",
                "audience": "选课学生",
                "style": "简洁清晰",
                "brief": brief,
                "slide_count": 7,
                "task_contract": {
                    "objective": "对比 CS5187.pdf、CS5297.pdf 和 CS5491.pdf，并用中文 PPT 展示",
                    "output_language": "zh-CN",
                    "artifact_types": ["pptx"],
                },
            },
        )

        output = result["output"]
        self.assertTrue(output["quality_report"]["passed"], output["quality_report"])
        self.assertEqual(
            [slide["title"] for slide in output["outline"]],
            [
                "三门课程对比",
                "CS5187：视觉计算",
                "CS5297：人工智能",
                "CS5491：人工智能安全",
                "横向比较",
                "选择建议",
                "要点回顾",
            ],
        )
        visible = json.dumps(output["outline"], ensure_ascii=False)
        self.assertNotIn("页面结构", visible)
        self.assertNotIn("来源覆盖率", visible)
        self.assertEqual(len(result["files"]), 1)

    def test_pptx_visible_title_never_includes_file_extension(self) -> None:
        result = execute(
            "pptx_outline",
            {
                "title": "课程教学内容比较.pptx",
                "audience": "选课学生",
                "style": "简洁",
                "brief": "内容重点",
                "slide_count": 3,
                "filename": "课程教学内容比较.pptx",
            },
        )

        self.assertEqual(result["output"]["title"], "课程教学内容比较")
        self.assertEqual(result["output"]["outline"][0]["title"], "课程教学内容比较")

    def test_pptx_table_preserves_rows_fields_coverage_and_unique_pages(self) -> None:
        rows = [
            {
                "cells": [f"PED{1100 + index}", f"课程 {index}", f"周三 {9 + index:02d}:00"],
                "source_locator": f"page:{1 + index // 6}",
            }
            for index in range(1, 13)
        ]
        result = execute(
            "pptx_generate",
            {
                "title": "体育课程表",
                "audience": "选课同学",
                "style": "简洁清晰，表格为主",
                "brief": "完整展示课程代码、名称和上课时间",
                "slide_count": 5,
                "table": {
                    "title": "课程安排",
                    "columns": ["课程代码", "课程名称", "上课时间"],
                    "rows": rows,
                },
                "task_contract": {
                    "exhaustive": True,
                    "requested_fields": ["code", "name", "time"],
                },
                "source_coverage": {
                    "coverage_ratio": 1.0,
                    "truncated": False,
                    "completed_rounds": 2,
                    "round_count": 2,
                },
            },
        )
        output = result["output"]
        self.assertTrue(output["quality_report"]["passed"], output["quality_report"])
        self.assertEqual(output["quality_report"]["table_row_count"], 12)
        self.assertEqual(
            output["outline"][-1]["bullets"][-1],
            "详情见表格，来源见页脚",
        )
        self.assertEqual(
            list(
                dict.fromkeys(
                    locator
                    for slide in output["outline"][1:-1]
                    for locator in slide["table"]["source_locators"]
                )
            ),
            ["page:1", "page:2", "page:3"],
        )
        signatures = [json.dumps(item, sort_keys=True) for item in output["outline"][1:-1]]
        self.assertEqual(len(signatures), len(set(signatures)))
        deck = Presentation(io.BytesIO(base64.b64decode(result["files"][0]["data_base64"])))
        for slide in list(deck.slides)[1:-1]:
            table_shape = next(shape for shape in slide.shapes if shape.has_table)
            footer = next(
                shape
                for shape in slide.shapes
                if shape.has_text_frame and shape.text.startswith("来源：")
            )
            self.assertGreaterEqual(
                footer.top,
                table_shape.top + table_shape.height + Inches(0.25),
            )

    def test_pptx_long_schedule_uses_one_unsplit_detail_cell(self) -> None:
        schedule = "；".join(
            f"2026年{1 + index // 28}月{1 + index % 28}日 周一 09:00-09:50" for index in range(72)
        )
        result = execute(
            "pptx_generate",
            {
                "title": "体育课程表",
                "audience": "选课同学",
                "style": "简洁清晰，表格为主",
                "brief": "完整展示课程代码、名称和上课时间",
                "slide_count": 3,
                "table": {
                    "title": "课程安排",
                    "columns": ["课程代码", "课程名称", "上课时间"],
                    "rows": [
                        {
                            "cells": ["PED1305", "体能训练", schedule],
                            "source_locator": "page:2",
                            "entity_id": "course:PED1305",
                        }
                    ],
                },
                "task_contract": {
                    "exhaustive": True,
                    "requested_fields": ["code", "name", "time"],
                    "output_language": "zh-CN",
                },
                "source_coverage": {
                    "coverage_ratio": 1.0,
                    "truncated": False,
                },
            },
        )
        output = result["output"]
        self.assertTrue(output["quality_report"]["passed"], output["quality_report"])
        self.assertEqual(output["quality_report"]["table_row_count"], 1)
        self.assertEqual(output["quality_report"]["display_row_count"], 1)
        self.assertEqual(output["slide_count"], 3)
        self.assertTrue(output["outline"][1]["table"]["detail"])
        self.assertEqual(output["outline"][1]["table"]["display_row_count"], 1)
        deck = Presentation(io.BytesIO(base64.b64decode(result["files"][0]["data_base64"])))
        detail_slide = deck.slides[1]
        table_shape = next(shape for shape in detail_slide.shapes if shape.has_table)
        value_paragraph = table_shape.table.cell(0, 0).text_frame.paragraphs[1]
        self.assertLessEqual(value_paragraph.font.size.pt, 10)
        footer = next(
            shape
            for shape in detail_slide.shapes
            if shape.has_text_frame and shape.text.startswith("来源：")
        )
        self.assertGreaterEqual(
            footer.top,
            table_shape.top + table_shape.height + Inches(0.25),
        )

    def test_pptx_exhaustive_quality_rejects_partial_scope_labels(self) -> None:
        result = execute(
            "pptx_generate",
            {
                "title": "体育课程表（节选）",
                "filename": "体育课程表_摘要.pptx",
                "audience": "选课同学",
                "style": "简洁表格",
                "brief": "完整课程",
                "slide_count": 3,
                "table": {
                    "title": "课程安排示例",
                    "columns": ["课程代码", "课程名称", "上课时间"],
                    "rows": [
                        {
                            "cells": ["PED1101", "独木舟", "周三 10:00-11:50"],
                            "source_locator": "page:1",
                        }
                    ],
                },
                "task_contract": {
                    "exhaustive": True,
                    "requested_fields": ["code", "name", "time"],
                    "output_language": "zh-CN",
                },
                "source_coverage": {"coverage_ratio": 1.0, "truncated": False},
            },
        )
        report = result["output"]["quality_report"]
        self.assertFalse(report["passed"])
        self.assertEqual(result["files"], [])
        self.assertIn(
            "presentation_exhaustive_scope_mislabeled",
            {item["code"] for item in report["violations"]},
        )

    def test_pptx_structured_contract_rejects_brief_only_content(self) -> None:
        result = execute(
            "pptx_generate",
            {
                "title": "体育课课程安排总览",
                "audience": "学生",
                "style": "清晰、规范",
                "brief": "提取自课程表 PDF",
                "slide_count": 10,
                "task_contract": {
                    "exhaustive": True,
                    "requested_fields": ["code", "name", "time"],
                },
                "source_coverage": {"coverage_ratio": 1.0, "truncated": False},
            },
        )
        report = result["output"]["quality_report"]
        self.assertFalse(report["passed"])
        self.assertIn(
            "structured_table_missing",
            {item["code"] for item in report["violations"]},
        )
        self.assertEqual(report["table_row_count"], 0)

    def test_pptx_exhaustive_source_rejects_incomplete_coverage(self) -> None:
        result = execute(
            "pptx_generate",
            {
                "title": "课程表",
                "audience": "学生",
                "style": "表格",
                "brief": "所有课程",
                "slide_count": 4,
                "table": {
                    "columns": ["课程代码", "课程名称"],
                    "rows": [{"cells": ["PED1101", "独木舟"], "source_locator": "page:1"}],
                },
                "task_contract": {"exhaustive": True, "requested_fields": ["code", "name"]},
                "source_coverage": {"coverage_ratio": 0.5, "truncated": True},
            },
        )
        self.assertFalse(result["output"]["quality_report"]["passed"])
        self.assertIn(
            "source_coverage_incomplete",
            {item["code"] for item in result["output"]["quality_report"]["violations"]},
        )

    def test_pptx_quality_rejects_duplicate_entities_and_delimiter_noise(self) -> None:
        result = execute(
            "pptx_generate",
            {
                "title": "体育课程表",
                "audience": "学生",
                "style": "表格",
                "brief": "所有课程",
                "slide_count": 4,
                "table": {
                    "columns": ["课程代码", "课程名称", "上课时间"],
                    "rows": [
                        {
                            "cells": ["PED1305", "Physical Fitness", "周一；；周三"],
                            "source_locator": "page:2",
                            "entity_id": "course:PED1305",
                        },
                        {
                            "cells": ["PED1305", "Physical Fitness", "周五"],
                            "source_locator": "page:3",
                            "entity_id": "course:PED1305",
                        },
                    ],
                },
                "task_contract": {
                    "exhaustive": True,
                    "requested_fields": ["code", "name", "time"],
                },
                "source_coverage": {"coverage_ratio": 1.0, "truncated": False},
            },
        )
        report = result["output"]["quality_report"]
        self.assertFalse(report["passed"])
        self.assertEqual(result["files"], [])
        codes = {item["code"] for item in report["violations"]}
        self.assertIn("duplicate_logical_entities", codes)
        self.assertIn("presentation_cell_delimiter_noise", codes)

    def test_pptx_quality_rejects_visible_reference_guidance(self) -> None:
        result = execute(
            "pptx_generate",
            {
                "title": "体育课程表",
                "audience": "学生",
                "style": "表格",
                "brief": "所有课程",
                "slide_count": 4,
                "table": {
                    "columns": ["课程代码", "课程名称", "上课时间"],
                    "rows": [
                        {
                            "cells": ["PED1305", "Physical Fitness", "（多节，见下列来源）"],
                            "source_locator": "page:2",
                        }
                    ],
                },
                "task_contract": {
                    "exhaustive": True,
                    "requested_fields": ["code", "name", "time"],
                },
                "source_coverage": {"coverage_ratio": 1.0, "truncated": False},
            },
        )
        report = result["output"]["quality_report"]
        self.assertFalse(report["passed"])
        self.assertIn(
            "presentation_reference_guidance_visible",
            {item["code"] for item in report["violations"]},
        )

    def test_pptx_quality_rejects_off_field_metadata_in_time_column(self) -> None:
        result = execute(
            "pptx_generate",
            {
                "title": "体育课程表",
                "audience": "学生",
                "style": "表格",
                "brief": "所有课程",
                "slide_count": 4,
                "table": {
                    "columns": ["课程代码", "课程名称", "上课时间"],
                    "rows": [
                        {
                            "cells": [
                                "PED1402",
                                "Golf",
                                "周三 09:30-11:20；(部分节次标注有容量/场地信息)",
                            ],
                            "source_locator": "page:2",
                        }
                    ],
                },
                "task_contract": {
                    "exhaustive": True,
                    "requested_fields": ["code", "name", "time"],
                },
                "source_coverage": {"coverage_ratio": 1.0, "truncated": False},
            },
        )
        report = result["output"]["quality_report"]
        self.assertFalse(report["passed"])
        self.assertIn(
            "presentation_field_semantic_mismatch",
            {item["code"] for item in report["violations"]},
        )

    def test_pptx_explicit_filename_still_rejects_paths(self) -> None:
        with self.assertRaisesRegex(ValueError, "invalid_output_filename"):
            execute(
                "pptx_generate",
                {
                    "title": "Quarterly review",
                    "audience": "Leadership",
                    "style": "professional",
                    "brief": "Results and next steps",
                    "slide_count": 4,
                    "filename": "reports/quarterly-review.pptx",
                },
            )

    def test_document_extract_returns_pdf_text_without_creating_a_file(self) -> None:
        source = io.BytesIO()
        document = canvas.Canvas(source)
        document.drawString(72, 760, "CS5491 course overview")
        document.showPage()
        document.drawString(72, 760, "Assessment and project requirements")
        document.save()
        result = execute(
            "document_extract",
            {
                "source_filename": "CS5491.pdf",
                "source_base64": base64.b64encode(b"\n\n\n\n\n" + source.getvalue()).decode(),
                "media_type": "application/pdf",
            },
        )
        self.assertEqual(result["files"], [])
        self.assertEqual(result["output"]["page_count"], 2)
        self.assertEqual(result["output"]["format"], "markdown")
        self.assertIn("parser_version", result["output"])
        self.assertIn("CS5491 course overview", result["output"]["text"])
        self.assertFalse(result["output"]["source_overwritten"])

    def test_document_extract_limits_large_attachment_context_by_tokens(self) -> None:
        text = "\n\n".join(f"## Section {index}\n" + "证据" * 500 for index in range(20))
        result = execute(
            "document_extract",
            {
                "source_filename": "long.txt",
                "source_base64": base64.b64encode(text.encode()).decode(),
                "media_type": "text/plain",
            },
        )
        output = result["output"]
        self.assertTrue(output["truncated"])
        self.assertLessEqual(output["token_count"], 16_000)
        self.assertLess(output["selected_chunk_count"], output["total_chunk_count"])
        self.assertIn("[[DOCUMENT ROUND 1]]", output["text"])
        self.assertIn("[[PAGE 1]]", output["text"])
        self.assertGreater(output["round_count"], output["completed_rounds"])
        self.assertTrue(output["has_more"])
        self.assertLess(output["coverage_ratio"], 1)

    def test_csv_profile_reports_types_missing_duplicates_and_download(self) -> None:
        source = "name,amount,note\nA,10,ok\nB,20,\nB,20,\n".encode()
        result = execute(
            "tabular_profile",
            {"source_filename": "sample.csv", "source_base64": base64.b64encode(source).decode()},
        )
        output = result["output"]
        self.assertEqual(output["row_count"], 3)
        self.assertEqual(output["duplicate_rows"], 1)
        self.assertEqual(output["columns"][1]["inferred_type"], "number")
        report = json.loads(base64.b64decode(result["files"][0]["data_base64"]))
        self.assertEqual(report["columns"][2]["missing_count"], 2)

    def test_xlsx_profile_uses_active_sheet(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "数据"
        sheet.append(["项目", "数值"])
        sheet.append(["A", 12.5])
        source = io.BytesIO()
        workbook.save(source)
        result = execute(
            "tabular_profile",
            {
                "source_filename": "sample.xlsx",
                "source_base64": base64.b64encode(source.getvalue()).decode(),
            },
        )
        self.assertEqual(result["output"]["sheet_name"], "数据")
        self.assertEqual(result["output"]["columns"][1]["numeric"]["mean"], 12.5)


if __name__ == "__main__":
    unittest.main()
