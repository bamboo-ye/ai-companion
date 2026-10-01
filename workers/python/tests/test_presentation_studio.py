from __future__ import annotations

import copy
import base64
import io
import json
import os
from pathlib import Path
import zipfile
from unittest.mock import patch
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation
from PIL import Image

from ai_companion_worker.office_tools import execute
from ai_companion_worker.presentation_design import (
    LAYOUTS, THEMES, layout_deck, plan_deck, repair_layouts, validate_deck, wrap_text,
)
from ai_companion_worker.presentation_render import export_pptx, normalize_assets, render_pptx

BASE = {"design_version": "2", "title": "产品季度复盘", "audience": "项目团队", "style": "商务", "brief": "## 工作进展\n- 完成资料梳理\n- 保持内容完整\n## 下一步计划\n- 实现编辑功能\n- 验证交付效果", "slide_count": 4, "visual_mode": "none"}
CASES = [
    ("清晰列表", "## 工作进展\n- 完成资料梳理\n- 保持内容完整\n- 验证导出效果"),
    ("实施流程", "## 制作流程\n- 导入资料\n- 设计页面\n- 调整排版\n- 完成交付"),
    ("里程碑", "## 项目时间轴\n- 一月：完成调研\n- 二月：发布预览\n- 三月：收集反馈"),
    ("方案对比", "## 方案对比\n- 方案甲：集中管理与统一维护\n- 方案乙：独立部署与分批升级"),
    ("关键指标", "## 关键指标\n- 25项功能完成测试\n- 8个部门参与验证\n- 96%测试通过率"),
    ("单一观点", "## 工作原则\n- 每一项结论都应当能够追溯到原始资料"),
    ("中文长文", "## 完整保留\n- " + "这是一段必须完整保留的中文说明，用于验证分页后内容仍然连续。" * 18),
    ("混合语言", "## API 流程\n- HTTP 200 indicates success; keep request_id intact.\n- 中文与 English 共存，版本 v2.3.1 和 99.5% 保持不变"),
    ("密集要点", "## 完整记录\n" + "\n".join(f"- 记录{n:02d}：保留该条记录及相关说明" for n in range(24))),
    ("标点公式", "## 研究说明\n- Bayes: P(A|B) = P(B|A)P(A)/P(B)\n- 保留括号（字段）、百分比 12.5% 和负数 -3.2"),
]


@pytest.mark.parametrize("theme", list(THEMES))
@pytest.mark.parametrize("title,brief", CASES)
def test_thirty_design_baselines_preserve_all_source_content(theme: str, title: str, brief: str) -> None:
    deck = plan_deck({**BASE, "title": title, "brief": brief, "theme_id": theme})
    deck, _ = repair_layouts(deck)
    scenes, violations = layout_deck(deck)
    assert not violations, (title, theme, violations)
    content = "".join(item["text"] for slide in deck["slides"][1:-1] for item in slide["items"])
    expected = "".join(line[2:].strip() for line in brief.splitlines() if line.startswith("- "))
    assert content == expected
    ppt = Presentation(io.BytesIO(export_pptx(deck, scenes)))
    assert len(ppt.slides) == len(deck["slides"])
    assert all(any(shape.has_text_frame and shape.text.strip() for shape in slide.shapes) for slide in ppt.slides)


def test_round_trip_retains_notes_lock_ids_and_manual_geometry() -> None:
    deck = plan_deck(BASE)
    slide = deck["slides"][1]
    slide.update(notes="讲解时突出来源", locked=True, source="产品文档 第 3 页")
    slide["overrides"] = {"title": {"x": 88, "font_size": 27, "locked": True}}
    returned = execute("pptx_outline", {**BASE, "deck_spec": json.loads(json.dumps(deck))})["output"]
    saved = returned["deck_spec"]["slides"][1]
    assert saved["id"] == slide["id"] and saved["notes"] == slide["notes"] and saved["locked"]
    title = next(n for n in returned["scenes"][1]["nodes"] if n["id"] == "title")
    assert title["x"] == 88 and title["font_size"] == 27


def test_locked_overflow_is_reported_without_deleting_content() -> None:
    deck = plan_deck(BASE)
    slide = deck["slides"][1]
    slide["locked"] = True
    slide["overrides"] = {"title": {"x": 900, "w": 500, "locked": True}}
    result = execute("pptx_generate", {**BASE, "deck_spec": deck})
    assert not result["output"]["quality_report"]["passed"]
    assert not result["files"]
    assert result["output"]["deck_spec"]["slides"][1]["items"] == slide["items"]


@pytest.mark.parametrize("value", [float("nan"), float("inf"), -100, True])
def test_rejects_non_finite_or_invalid_geometry(value: float) -> None:
    deck = plan_deck(BASE)
    deck["slides"][1]["overrides"] = {"title": {"x": value}}
    with pytest.raises(ValueError):
        validate_deck(deck)


def test_chart_is_native_with_embedded_data_and_animation_targets_exist() -> None:
    deck = plan_deck(BASE)
    deck["animation"] = "explain"
    deck["slides"][1]["layout"] = "chart-column"
    deck["slides"][1]["chart"] = {"labels": ["甲", "乙", "丙"], "values": [12, -3, 24], "series": "营业额", "unit": "万"}
    scenes, violations = layout_deck(deck)
    assert not violations
    data = export_pptx(deck, scenes)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert "ppt/charts/chart1.xml" in archive.namelist()
        assert any(n.startswith("ppt/embeddings/") for n in archive.namelist())
        xml = ET.fromstring(archive.read("ppt/slides/slide2.xml"))
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        ids = {n.attrib["id"] for n in xml.findall(".//p:cNvPr", ns)}
        targets = {n.attrib["spid"] for n in xml.findall(".//p:spTgt", ns)}
        assert targets and targets <= ids
        assert xml.find("p:timing", ns) is not None
        timing_ids = [n.attrib["id"] for n in xml.findall(".//p:cTn", ns)]
        assert len(timing_ids) == len(set(timing_ids))


def test_no_motion_has_no_animation_xml() -> None:
    deck = plan_deck({**BASE, "animation": "none"})
    data = export_pptx(deck, layout_deck(deck)[0])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert b"<p:timing" not in archive.read("ppt/slides/slide2.xml")
        assert b"<p:transition" not in archive.read("ppt/slides/slide2.xml")


@pytest.mark.parametrize("values,number", [([12, -3, 24], "0"), ([12.5, -3, 24], "0.0"), ([12.25, -3.75, 0], "0.00")])
def test_native_chart_number_format_has_no_dangling_decimal(values: list[float], number: str) -> None:
    deck = plan_deck(BASE)
    deck["slides"][1].update(layout="chart-bar", chart={"labels": ["甲", "乙", "丙"], "values": values, "series": "变化", "unit": "项"})
    data = export_pptx(deck, layout_deck(deck)[0])
    presentation = Presentation(io.BytesIO(data))
    chart = next(shape.chart for shape in presentation.slides[1].shapes if shape.has_chart)
    assert chart.plots[0].data_labels.number_format == f'{number}"项";-{number}"项";0"项"'
    assert list(chart.series[0].values) == values


@pytest.mark.parametrize("advance", ["click", "auto"])
def test_motion_has_office_phase_containers_and_complete_build_records(advance: str) -> None:
    deck = plan_deck(BASE)
    slide = deck["slides"][1]
    slide.update(layout="steps", animation="explain", advance=advance, duration_ms=450,
                 items=[{"id": f"item-{i}", "text": str(i)} for i in range(4)])
    scenes, _ = layout_deck(deck)
    with zipfile.ZipFile(io.BytesIO(export_pptx(deck, scenes))) as archive:
        xml = ET.fromstring(archive.read("ppt/slides/slide2.xml"))
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    effects = xml.findall(".//p:cTn[@presetClass='entr']", ns)
    builds = xml.findall("p:timing/p:bldLst/p:bldP", ns)
    targets = {n.attrib["spid"] for effect in effects for n in effect.findall(".//p:spTgt", ns)}
    assert {node.attrib["spid"] for node in builds} == targets
    assert len(builds) == len(targets)
    assert all(effect.attrib["grpId"] == "0" for effect in effects)
    assert sum(effect.attrib["nodeType"] == ("clickEffect" if advance == "click" else "afterEffect") for effect in effects) == 4
    sequence = xml.find(".//p:cTn[@nodeType='mainSeq']/p:childTnLst", ns)
    assert sequence is not None
    clicks = sequence.findall("p:par/p:cTn", ns)
    assert len(clicks) == (4 if advance == "click" else 1)
    phases = [phase for click in clicks for phase in click.findall("p:childTnLst/p:par/p:cTn", ns)]
    assert len(phases) == 4
    assert [phase.find("p:stCondLst/p:cond", ns).attrib["delay"] for phase in phases] == (["0"] * 4 if advance == "click" else ["700", "1400", "2100", "2800"])


def test_assets_cannot_reference_remote_urls_or_svg() -> None:
    deck = plan_deck(BASE)
    deck["assets"] = {"bad": {"media_type": "image/svg+xml", "data_base64": "abc"}}
    with pytest.raises(ValueError):
        validate_deck(deck)
    deck["assets"] = {"bad": {"media_type": "image/png", "data_base64": "https://example.com/private"}}
    with pytest.raises(ValueError):
        normalize_assets(validate_deck(deck))


def test_renderer_unavailable_does_not_claim_visual_verification() -> None:
    with patch("ai_companion_worker.presentation_studio.render_pptx", return_value=({"status": "unavailable", "passed": False, "message": "no renderer"}, None, [])):
        result = execute("pptx_generate", BASE)
    assert len(result["files"]) == 2
    assert result["output"]["render_report"]["passed"] is False
    assert result["output"]["warnings"] == ["no renderer"]


def test_source_coverage_gate_is_preserved() -> None:
    result = execute("pptx_generate", {**BASE, "task_contract": {"exhaustive": True}, "source_coverage": {"coverage_ratio": .5, "truncated": True}})
    assert not result["files"]
    assert any(v["code"] == "source_coverage_incomplete" for v in result["output"]["quality_report"]["violations"])


def test_wrapping_keeps_formula_punctuation_and_identifiers() -> None:
    value = "Bayes P(A|B)=P(B|A)P(A)/P(B)，版本v2.3.1；准确率99.5%。"
    assert "".join(wrap_text(value, 230, 20)) == value


def test_import_rejects_duplicate_slide_ids() -> None:
    deck = plan_deck(BASE)
    deck["slides"].append(copy.deepcopy(deck["slides"][1]))
    with pytest.raises(ValueError):
        validate_deck(deck)


def pictured_deck() -> dict:
    image = io.BytesIO()
    Image.new("RGB", (600, 280), "#72DCD1").save(image, format="PNG")
    deck = plan_deck(BASE)
    deck["assets"] = {"picture": {"media_type": "image/png", "data_base64": base64.b64encode(image.getvalue()).decode(), "source": "资料 第 3 页", "kind": "source", "source_page": 3}}
    deck["slides"][1].update(image_id="picture", layout="image-right", source="资料 第 3 页")
    normalize_assets(deck)
    return deck


def test_dense_image_content_repairs_preserve_the_picture_and_all_text() -> None:
    deck = pictured_deck()
    items = [{"id": f"long-{i}", "text": "图片与文字必须同时保留。" * 7} for i in range(4)]
    deck["slides"][1]["items"] = items
    repaired, repairs = repair_layouts(deck)
    scenes, violations = layout_deck(repaired)
    assert repairs and not violations
    assert sum(n["type"] == "image" for scene in scenes for n in scene["nodes"]) == 1
    assert [item for slide in repaired["slides"] for item in slide["items"] if item["id"].startswith("long-")] == [{**item, "source": ""} for item in items]
    assert repaired["assets"]["picture"]["source_page"] == 3


def test_table_gate_and_bounded_pagination_keep_complete_records() -> None:
    deck = plan_deck(BASE)
    rows = [[f"记录{i}", "说明必须保持完整" * 5] for i in range(12)]
    deck["slides"][1].update(layout="table", items=[], table={"columns": ["名称", "说明"], "rows": rows})
    assert layout_deck(deck)[1]
    repaired, repairs = repair_layouts(deck)
    scenes, violations = layout_deck(repaired)
    assert repairs and not violations
    assert [row for slide in repaired["slides"] if "table" in slide for row in slide["table"]["rows"]] == rows
    ppt = Presentation(io.BytesIO(export_pptx(repaired, scenes)))
    assert sum(shape.has_table for slide in ppt.slides for shape in slide.shapes) > 1
    deck["slides"][1]["locked"] = True
    result = execute("pptx_generate", {**BASE, "deck_spec": deck})
    assert not result["files"]
    assert result["output"]["quality_report"]["violations"]


def test_renderer_failure_blocks_export_instead_of_claiming_success() -> None:
    with patch("ai_companion_worker.presentation_studio.render_pptx", return_value=({"status": "failed", "passed": False, "message": "timeout"}, None, [])):
        result = execute("pptx_generate", BASE)
    assert not result["files"]
    assert not result["output"]["quality_report"]["passed"]


def test_image_provenance_is_checked_after_design_round_trip() -> None:
    deck = pictured_deck()
    deck["slides"][1]["source"] = "资料 第 9 页"
    result = execute("pptx_outline", {**BASE, "deck_spec": deck})
    assert any(v["code"] == "source_visual_topic_mismatch" for v in result["output"]["quality_report"]["violations"])


@pytest.mark.parametrize("layout", list(LAYOUTS))
@pytest.mark.parametrize("theme", list(THEMES))
def test_every_catalog_layout_can_export_native_objects(layout: str, theme: str) -> None:
    deck = pictured_deck()
    deck["theme_id"] = theme
    slide = deck["slides"][1]
    slide.update(layout=layout, items=[{"id": "a", "text": "25项经过验证的功能", "source": ""}, {"id": "b", "text": "8个清晰的实施步骤", "source": ""}])
    if not layout.startswith("image-") and layout != "gallery":
        slide["image_id"] = ""
    if layout.startswith("chart-"):
        slide["chart"] = {"labels": ["第一期", "第二期", "第三期"], "values": [12, -3, 24], "series": "变化", "unit": "项"}
    if layout == "table":
        slide["table"] = {"columns": ["阶段", "状态"], "rows": [["第一期", "已完成"], ["第二期", "进行中"]]}
    if layout == "capability-map":
        from ai_companion_worker.presentation_projects import project_deck
        slide['items'] = project_deck({'project_name': '伴AI'})['slides'][0]['items']
    scenes, violations = layout_deck(deck)
    assert not violations, (layout, theme, violations)
    assert export_pptx(deck, scenes)


@pytest.mark.skipif(not os.environ.get("PRESENTATION_RENDER_AUDIT_DIR"), reason="Opt-in LibreOffice render audit; set PRESENTATION_RENDER_AUDIT_DIR")
@pytest.mark.parametrize("theme", list(THEMES))
def test_thirty_baselines_through_actual_office_renderer(theme: str) -> None:
    """Render ten content cases per theme in one deck to amortize Office startup."""
    import pymupdf

    directory = Path(os.environ["PRESENTATION_RENDER_AUDIT_DIR"])
    directory.mkdir(parents=True, exist_ok=True)
    deck = plan_deck({**BASE, "theme_id": theme})
    slides = []
    for case, (title, brief) in enumerate(CASES):
        planned, _ = repair_layouts(plan_deck({**BASE, "theme_id": theme, "title": title, "brief": brief}))
        for slide in planned["slides"]:
            slide["id"] = f"case-{case}-{slide['id']}"
            slides.append(slide)
    deck["slides"] = slides
    scenes, violations = layout_deck(deck)
    assert not violations
    data = export_pptx(deck, scenes)
    report, pdf, _ = render_pptx(data, len(slides))
    assert report["status"] == "verified" and report["passed"], report
    assert pdf
    (directory / f"{theme}-baseline.pptx").write_bytes(data)
    (directory / f"{theme}-baseline.pdf").write_bytes(pdf)
    (directory / f"{theme}-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    with pymupdf.open(stream=pdf, filetype="pdf") as document:
        for scene, page in zip(scenes, document, strict=True):
            rendered = "".join(page.get_text().split())
            for node in scene["nodes"]:
                if node["type"] == "text":
                    assert "".join(node["text"].split()) in rendered, (theme, scene["id"], node["id"])
        for index in [1, 5, 9, 13, 17]:
            document[index].get_pixmap(matrix=pymupdf.Matrix(1, 1), alpha=False).save(directory / f"{theme}-page-{index + 1}.png")


@pytest.mark.skipif(not os.environ.get("PRESENTATION_RENDER_AUDIT_DIR"), reason="Opt-in Office layout catalog rendering")
@pytest.mark.parametrize("theme", list(THEMES))
def test_catalog_through_actual_renderer(theme: str) -> None:
    import pymupdf

    directory = Path(os.environ["PRESENTATION_RENDER_AUDIT_DIR"])
    deck = pictured_deck()
    deck.update(theme_id=theme, title="页面表达方式 · 版式验收", animation="explain")
    base = deck["slides"][1]
    slides = []
    for layout, name in LAYOUTS.items():
        slide = copy.deepcopy(base)
        slide.update(id=layout, title=name, layout=layout, items=[{"id": "a", "text": "25项经过验证的功能", "source": ""}, {"id": "b", "text": "8个清晰的实施步骤", "source": ""}])
        if not layout.startswith("image-") and layout != "gallery":
            slide["image_id"] = ""
        if layout.startswith("chart-"):
            slide["chart"] = {"labels": ["第一期", "第二期", "第三期"], "values": [12, -3, 24], "series": "变化", "unit": "项"}
        if layout == "table":
            slide["table"] = {"columns": ["阶段", "状态"], "rows": [["第一期", "已完成"], ["第二期", "进行中"]]}
        slides.append(slide)
    deck["slides"] = slides
    scenes, violations = layout_deck(deck)
    assert not violations
    data = export_pptx(deck, scenes)
    report, pdf, _ = render_pptx(data, len(slides))
    assert report["status"] == "verified" and report["passed"], report
    assert pdf
    (directory / f"{theme}-catalog.pptx").write_bytes(data)
    (directory / f"{theme}-catalog.pdf").write_bytes(pdf)
    (directory / f"{theme}-catalog.deck.json").write_text(json.dumps(deck, ensure_ascii=False, indent=2))
    with pymupdf.open(stream=pdf, filetype="pdf") as document:
        for slide, page in zip(slides, document, strict=True):
            rendered = "".join(page.get_text().split())
            assert slide["title"] in rendered
            if slide.get("chart"):
                assert all(label in rendered for label in slide["chart"]["labels"])
                assert all(f"{value}项" in rendered for value in slide["chart"]["values"])
            if slide.get("table"):
                assert all(cell in rendered for row in [slide["table"]["columns"], *slide["table"]["rows"]] for cell in row)
            if slide["image_id"]:
                assert page.get_images()
            if slide["layout"] in {"table", "chart-bar", "timeline", "comparison", "quote", "gallery"}:
                page.get_pixmap(matrix=pymupdf.Matrix(1, 1), alpha=False).save(directory / f"{theme}-{slide['layout']}.png")
