import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from ai_companion_worker.ledger_export import export_workbook


class LedgerExportTest(unittest.TestCase):
    def test_generates_detail_summary_and_chart(self) -> None:
        payload = {
            "month": "2026-07",
            "currency": "CNY",
            "timezone": "Asia/Shanghai",
            "generated_at": "2026-07-13T08:00:00Z",
            "entries": [
                {
                    "id": "entry-1",
                    "direction": "expense",
                    "currency": "CNY",
                    "amount_minor": 3600,
                    "category": "transport",
                    "merchant": "出租车",
                    "note": "回家",
                    "occurred_at": "2026-07-12T12:00:00Z",
                }
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "ledger.xlsx"
            export_workbook(payload, output)
            workbook = load_workbook(output, data_only=False)

        self.assertEqual(workbook.sheetnames, ["账单明细", "月度汇总"])
        self.assertEqual(workbook["账单明细"]["G6"].value, 36)
        self.assertEqual(workbook["月度汇总"]["A11"].value, "transport")
        self.assertEqual(len(workbook["月度汇总"]._charts), 1)


if __name__ == "__main__":
    unittest.main()
